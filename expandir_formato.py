# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

C_AZUL_INST   = "0A2342"   # Azul Marino Institucional PC
C_GUINDA      = "801438"   # Vino / Guinda Medellín
C_AZUL_CLARO  = "F1F5F9"   # Fondo gris-azulado
C_GRIS_FONDO  = "F8FAFC"   # Slate 50
C_GRIS_BORDE  = "CBD5E1"   # Slate 300
C_BORDE_DARK  = "94A3B8"   # Slate 400
C_GRIS_TEXTO  = "334155"   # Slate 700
C_BLANCO      = "FFFFFF"
C_ROJO_FOLIO  = "B91C1C"   # Rojo institucional

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=35, bottom=35, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcMar'):
            tcPr.remove(child)
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcBorders'):
            tcPr.remove(child)
    xml_parts = []
    borders_dict = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for side, border in borders_dict.items():
        if border:
            val = border.get('val', 'single')
            sz = border.get('sz', '4')
            color = border.get('color', C_GRIS_BORDE)
            xml_parts.append(f'<w:{side} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>')
        else:
            xml_parts.append(f'<w:{side} w:val="none"/>')
    borders_xml = f'<w:tcBorders {nsdecls("w")}>' + "".join(xml_parts) + '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))

def apply_uniform_borders(cell, color=C_GRIS_BORDE, sz="4"):
    b = {'val': 'single', 'sz': sz, 'color': color}
    set_cell_borders(cell, top=b, bottom=b, left=b, right=b)

def add_p(cell, text="", bold=False, italic=False, size_pt=7.8, color_rgb=(10,35,66), align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=0, line_spacing=1.05):
    if len(cell.paragraphs) == 1 and cell.paragraphs[0].text == "":
        p = cell.paragraphs[0]
    else:
        p = cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(size_pt)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = RGBColor(*color_rgb)
    return p

def append_run(paragraph, text, bold=False, italic=False, size_pt=7.8, color_rgb=(10,35,66)):
    r = paragraph.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(size_pt)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor(*color_rgb)
    return r

def make_section_header(table, title, bg_color=C_AZUL_INST, text_color=(255,255,255), size_pt=8.0):
    cell = table.rows[0].cells[0]
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=25, bottom=25, left=60, right=60)
    apply_uniform_borders(cell, color=bg_color, sz="4")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    r.font.name = 'Arial'
    r.font.size = Pt(size_pt)
    r.font.bold = True
    r.font.color.rgb = RGBColor(*text_color)

def set_row_height(row, height_pt):
    trPr = row._tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(int(height_pt * 20)))
    trHeight.set(qn('w:hRule'), 'atLeast')
    trPr.append(trHeight)

def set_zero_table_indent(table):
    tblPr = table._tbl.tblPr
    tblInd = OxmlElement('w:tblInd')
    tblInd.set(qn('w:w'), '0')
    tblInd.set(qn('w:type'), 'dxa')
    tblPr.append(tblInd)

def construir_formato_expandido():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.24)
    sec.bottom_margin = Inches(0.22)
    sec.left_margin = Inches(0.48)
    sec.right_margin = Inches(0.48)
    
    W_TOTAL = Inches(7.54)
    
    # 1. Banner Superior Oficial (Más prominente e institucional)
    p_ban = doc.add_paragraph()
    p_ban.paragraph_format.space_before = Pt(0)
    p_ban.paragraph_format.space_after = Pt(2)
    p_ban.paragraph_format.line_spacing = 1.0
    r_ban = p_ban.add_run()
    r_ban.add_picture('Documentacion/banner_superior_oficial.png', width=W_TOTAL, height=Inches(1.15))
    
    # 2. Encabezado Institucional y Folio
    tbl_hdr = doc.add_table(rows=1, cols=2)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_hdr.autofit = False
    set_zero_table_indent(tbl_hdr)
    
    c_inst, c_fol = tbl_hdr.rows[0].cells
    c_inst.width = Inches(5.34)
    c_fol.width = Inches(2.20)
    set_cell_margins(c_inst, top=14, bottom=14, left=25, right=25)
    set_cell_margins(c_fol, top=20, bottom=20, left=35, right=35)
    set_cell_borders(c_inst, top=None, bottom=None, left=None, right=None)
    
    b_fol = {'val': 'single', 'sz': '8', 'color': C_BORDE_DARK}
    set_cell_borders(c_fol, top=b_fol, bottom=b_fol, left=b_fol, right=b_fol)
    set_cell_background(c_fol, C_AZUL_CLARO)
    
    add_p(c_inst, "H. AYUNTAMIENTO CONSTITUCIONAL DE MEDELLÍN DE BRAVO, VERACRUZ", bold=True, size_pt=9.0, color_rgb=(10,35,66))
    add_p(c_inst, "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS", bold=True, size_pt=8.2, color_rgb=(128,20,56), space_before=2)
    add_p(c_inst, "ACTA DE DESLINDE DE RESPONSABILIDAD Y CONSENTIMIENTO INFORMADO", bold=True, size_pt=9.0, color_rgb=(15,23,42), space_before=2)
    add_p(c_inst, "SERVICIO OPERATIVO: CONTROL, RETIRO Y/O REUBICACIÓN DE ENJAMBRES (HIMENÓPTEROS)", italic=True, bold=True, size_pt=7.5, color_rgb=(71,85,105), space_before=2)
    
    pf1 = add_p(c_fol, "FOLIO: ", bold=True, size_pt=8.5, color_rgb=(185,28,28))
    append_run(pf1, "PC-MED/FAU/26-______", bold=True, size_pt=8.5, color_rgb=(185,28,28))
    pf2 = add_p(c_fol, "FECHA: ", bold=True, size_pt=7.5, color_rgb=(51,65,85), space_before=2)
    append_run(pf2, "____ / _________ / 2026", bold=False, size_pt=7.5, color_rgb=(15,23,42))
    pf3 = add_p(c_fol, "HORA REPORTE: ", bold=True, size_pt=7.5, color_rgb=(51,65,85), space_before=2)
    append_run(pf3, "____:____ hrs", bold=False, size_pt=7.5, color_rgb=(15,23,42))
    pf4 = add_p(c_fol, "UNIDAD ASIGNADA: ", bold=True, size_pt=7.5, color_rgb=(51,65,85), space_before=2)
    append_run(pf4, "U-______", bold=True, size_pt=7.5, color_rgb=(10,35,66))
    
    # 3. Sección I: Datos del Solicitante
    t_sec1 = doc.add_table(rows=1, cols=1)
    t_sec1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_zero_table_indent(t_sec1)
    t_sec1.rows[0].cells[0].width = W_TOTAL
    make_section_header(t_sec1, "I. DATOS DEL SOLICITANTE / PROPIETARIO / REPRESENTANTE LEGAL", bg_color=C_AZUL_INST, size_pt=8.0)
    
    t_d1 = doc.add_table(rows=3, cols=2)
    t_d1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_d1.autofit = False
    set_zero_table_indent(t_d1)
    for r in t_d1.rows:
        set_row_height(r, 20)
        r.cells[0].width = Inches(5.19)
        r.cells[1].width = Inches(2.35)
        for c in r.cells:
            set_cell_margins(c, top=35, bottom=35, left=50, right=50)
            apply_uniform_borders(c, color=C_GRIS_BORDE, sz="4")
            
    p_nom = add_p(t_d1.rows[0].cells[0], "Nombre Completo del Solicitante: ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_nom, "__________________________________________________", size_pt=7.8, color_rgb=(100,116,139))
    p_ine = add_p(t_d1.rows[0].cells[1], "No. Identificación (INE): ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_ine, "____________________", size_pt=7.8, color_rgb=(100,116,139))
    
    p_inst = add_p(t_d1.rows[1].cells[0], "Institución / Razón Social (si aplica): ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_inst, "__________________________________________", size_pt=7.8, color_rgb=(100,116,139))
    p_tel = add_p(t_d1.rows[1].cells[1], "Teléfono de Contacto: ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_tel, "(_____) _____________", size_pt=7.8, color_rgb=(100,116,139))
    
    c_car = t_d1.rows[2].cells[0]
    c_car.merge(t_d1.rows[2].cells[1])
    c_car.width = W_TOTAL
    p_car = add_p(c_car, "Carácter con el que comparece:  ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_car, "☐ Propietario      ☐ Arrendatario / Inquilino      ☐ Encargado / Administrador      ☐ Representante Institucional      ☐ Vecino", size_pt=7.8, color_rgb=(51,65,85))
    
    # 4. Sección II: Ubicación
    t_sec2 = doc.add_table(rows=1, cols=1)
    t_sec2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_zero_table_indent(t_sec2)
    t_sec2.rows[0].cells[0].width = W_TOTAL
    make_section_header(t_sec2, "II. UBICACIÓN EXACTA DEL INMUEBLE O PREDIO INTERVENIDO", bg_color=C_AZUL_INST, size_pt=8.0)
    
    t_d2 = doc.add_table(rows=3, cols=2)
    t_d2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_d2.autofit = False
    set_zero_table_indent(t_d2)
    for r in t_d2.rows:
        set_row_height(r, 20)
        r.cells[0].width = Inches(4.74)
        r.cells[1].width = Inches(2.80)
        for c in r.cells:
            set_cell_margins(c, top=35, bottom=35, left=50, right=50)
            apply_uniform_borders(c, color=C_GRIS_BORDE, sz="4")
            
    p_calle = add_p(t_d2.rows[0].cells[0], "Calle y Número Exterior / Interior: ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_calle, "________________________________________", size_pt=7.8, color_rgb=(100,116,139))
    p_col = add_p(t_d2.rows[0].cells[1], "Colonia / Fracc. / Localidad: ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_col, "_____________________________", size_pt=7.8, color_rgb=(100,116,139))
    
    c_ref = t_d2.rows[1].cells[0]
    c_ref.merge(t_d2.rows[1].cells[1])
    c_ref.width = W_TOTAL
    p_ref = add_p(c_ref, "Referencias del Inmueble (entre calles, color o fachada): ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_ref, "____________________________________________________________________________", size_pt=7.8, color_rgb=(100,116,139))
    
    c_tipo = t_d2.rows[2].cells[0]
    c_tipo.merge(t_d2.rows[2].cells[1])
    c_tipo.width = W_TOTAL
    p_tipo = add_p(c_tipo, "Tipo de Inmueble:  ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_tipo, "☐ Casa Habitación      ☐ Plantel Escolar      ☐ Comercio / Negocio      ☐ Terreno / Predio Baldío      ☐ Vía Pública      ☐ Otro", size_pt=7.8, color_rgb=(51,65,85))
    
    # 5. Sección III: Diagnóstico Técnico
    t_sec3 = doc.add_table(rows=1, cols=1)
    t_sec3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_zero_table_indent(t_sec3)
    t_sec3.rows[0].cells[0].width = W_TOTAL
    make_section_header(t_sec3, "III. EVALUACIÓN TÉCNICA Y DIAGNÓSTICO DE RIESGO (USO EXCLUSIVO PROTECCIÓN CIVIL)", bg_color=C_AZUL_INST, size_pt=8.0)
    
    t_d3 = doc.add_table(rows=2, cols=2)
    t_d3.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_d3.autofit = False
    set_zero_table_indent(t_d3)
    for r in t_d3.rows:
        set_row_height(r, 20)
        r.cells[0].width = Inches(3.77)
        r.cells[1].width = Inches(3.77)
        for c in r.cells:
            set_cell_margins(c, top=35, bottom=35, left=50, right=50)
            apply_uniform_borders(c, color=C_GRIS_BORDE, sz="4")
            
    p_esp = add_p(t_d3.rows[0].cells[0], "Especie Detectada: ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_esp, "☐ Abeja Melífera     ☐ Avispa / Avispón     ☐ Otro", size_pt=7.8, color_rgb=(51,65,85))
    p_ubi = add_p(t_d3.rows[0].cells[1], "Ubicación del Nido: ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_ubi, "☐ Árbol / Ramas   ☐ Muro / Cornisa   ☐ Plafón   ☐ Tinaco", size_pt=7.8, color_rgb=(51,65,85))
    
    p_alt = add_p(t_d3.rows[1].cells[0], "Altura Aprox: ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_alt, "______ m   |   Acceso:  ☐ Libre   ☐ Escalera   ☐ Confinado", size_pt=7.8, color_rgb=(51,65,85))
    p_acc = add_p(t_d3.rows[1].cells[1], "Acción Determinada: ", bold=True, size_pt=7.8, color_rgb=(10,35,66))
    append_run(p_acc, "☐ Reubicación Sustentable     ☐ Control Físico Emergente", size_pt=7.8, color_rgb=(51,65,85))
    
    # 6. Sección IV: Declaraciones del Ciudadano
    t_sec4 = doc.add_table(rows=1, cols=1)
    t_sec4.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_zero_table_indent(t_sec4)
    t_sec4.rows[0].cells[0].width = W_TOTAL
    make_section_header(t_sec4, "IV. DECLARACIONES Y MEDIDAS DE SEGURIDAD (BAJO PROTESTA DE DECIR VERDAD)", bg_color=C_AZUL_INST, size_pt=8.0)
    
    t_d4 = doc.add_table(rows=1, cols=1)
    t_d4.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_zero_table_indent(t_d4)
    c_dec = t_d4.rows[0].cells[0]
    c_dec.width = W_TOTAL
    set_cell_background(c_dec, C_GRIS_FONDO)
    set_cell_margins(c_dec, top=30, bottom=30, left=50, right=50)
    apply_uniform_borders(c_dec, color=C_GRIS_BORDE, sz="4")
    
    txt_dec = [
        ("1. Autorización de Acceso: ", "Autorizo el libre ingreso del personal operativo y unidades de auxilio de Protección Civil y Bomberos al inmueble señalado para realizar las maniobras pertinentes."),
        ("2. Personas Vulnerables y Alergias: ", "Declaro bajo protesta de decir verdad si habitan personas alérgicas a picaduras (anafilaxia): ☐ SÍ   ☐ NO. Me obligo a evacuar preventivamente a menores, adultos mayores y resguardar animales domésticos."),
        ("3. Perímetro de Seguridad: ", "Me obligo a acatar las instrucciones de seguridad, manteniendo puertas y ventanas cerradas, luces apagadas y a los ocupantes a más de 30 metros de distancia del área de operación."),
        ("4. Conducta Biológica de Pecoreo: ", "Quedo debidamente enterado(a) de que abejas pecoreadoras retornarán naturalmente al sitio durante 24 a 48 horas, debiendo mantener precauciones sin que ello sea negligencia oficial."),
        ("5. Intervención en Estructuras: ", "En caso de requerirse aperturas mecánicas en falso plafón, muros, techumbres o ramas para extraer el nido, autorizo la maniobra asumiendo reparaciones sin reclamo a la corporación.")
    ]
    for idx, (tit, cuerpo) in enumerate(txt_dec):
        p_item = add_p(c_dec, tit, bold=True, size_pt=7.2, color_rgb=(10,35,66), space_before=(1 if idx>0 else 0), line_spacing=1.04)
        append_run(p_item, cuerpo, bold=False, size_pt=7.2, color_rgb=(51,65,85))
        
    # 7. Sección V: Deslinde Legal
    t_sec5 = doc.add_table(rows=1, cols=1)
    t_sec5.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_zero_table_indent(t_sec5)
    t_sec5.rows[0].cells[0].width = W_TOTAL
    make_section_header(t_sec5, "V. FUNDAMENTO JURÍDICO Y CLÁUSULA DE DESLINDE LEGAL DE RESPONSABILIDAD", bg_color=C_GUINDA, size_pt=8.0)
    
    t_d5 = doc.add_table(rows=1, cols=1)
    t_d5.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_zero_table_indent(t_d5)
    c_leg = t_d5.rows[0].cells[0]
    c_leg.width = W_TOTAL
    set_cell_background(c_leg, C_AZUL_CLARO)
    set_cell_margins(c_leg, top=30, bottom=30, left=50, right=50)
    apply_uniform_borders(c_leg, color=C_BORDE_DARK, sz="6")
    
    p_leg = add_p(c_leg, "FUNDAMENTO LEGAL: ", bold=True, size_pt=7.1, color_rgb=(128,20,56), align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_spacing=1.04)
    append_run(p_leg, "Con fundamento en los Artículos 8° y 115 Constitucionales; 1°, 2°, 3°, 18, 38 y 41 de la Ley No. 856 de Protección Civil y la Reducción del Riesgo de Desastres para el Estado de Veracruz de Ignacio de la Llave; Ley de Fomento Apícola del Estado de Veracruz; y Reglamento de Protección Civil de Medellín de Bravo: El(la) suscrito(a) ", size_pt=7.1, color_rgb=(15,23,42))
    append_run(p_leg, "DESLINDA DE TODA RESPONSABILIDAD LEGAL, CIVIL, PENAL, ADMINISTRATIVA Y PATRIMONIAL ", bold=True, size_pt=7.1, color_rgb=(128,20,56))
    append_run(p_leg, "a la Dirección de Protección Civil y Bomberos de Medellín de Bravo, a su personal operativo y al H. Ayuntamiento de Medellín de Bravo por cualquier picadura fortuita, daño fortuito o consecuencia derivada antes, durante o después del procedimiento, solicitando el servicio de manera voluntaria y de buena fe.", size_pt=7.1, color_rgb=(15,23,42))
    
    # 8. Sección VI: Firmas
    t_sec6 = doc.add_table(rows=1, cols=1)
    t_sec6.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_zero_table_indent(t_sec6)
    t_sec6.rows[0].cells[0].width = W_TOTAL
    make_section_header(t_sec6, "VI. CONSTANCIA DE CONFORMIDAD Y FIRMAS DE VALIDACIÓN", bg_color=C_AZUL_INST, size_pt=8.0)
    
    t_sig = doc.add_table(rows=1, cols=3)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sig.autofit = False
    set_zero_table_indent(t_sig)
    set_row_height(t_sig.rows[0], 55)
    
    c_s1, c_s2, c_s3 = t_sig.rows[0].cells
    c_s1.width = Inches(2.67)
    c_s2.width = Inches(2.67)
    c_s3.width = Inches(2.20)
    
    for c in (c_s1, c_s2, c_s3):
        set_cell_margins(c, top=20, bottom=20, left=35, right=35)
        apply_uniform_borders(c, color=C_GRIS_BORDE, sz="4")
        
    add_p(c_s1, "SOLICITANTE / PROPIETARIO", bold=True, size_pt=7.2, color_rgb=(10,35,66), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(c_s1, "\n________________________________________", bold=False, size_pt=7.2, color_rgb=(100,116,139), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(c_s1, "Nombre: _______________________________", bold=True, size_pt=7.0, color_rgb=(51,65,85), space_before=2)
    add_p(c_s1, "INE / Pasaporte: ________________________", bold=True, size_pt=7.0, color_rgb=(51,65,85), space_before=2)
    
    add_p(c_s2, "ELEMENTO OPERATIVO A CARGO", bold=True, size_pt=7.2, color_rgb=(10,35,66), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(c_s2, "\n________________________________________", bold=False, size_pt=7.2, color_rgb=(100,116,139), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(c_s2, "Nombre: _______________________________", bold=True, size_pt=7.0, color_rgb=(51,65,85), space_before=2)
    add_p(c_s2, "Rango y No. Unidad: ___________________", bold=True, size_pt=7.0, color_rgb=(51,65,85), space_before=2)
    
    add_p(c_s3, "SELLO OFICIAL DE VALIDACIÓN", bold=True, size_pt=7.2, color_rgb=(128,20,56), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(c_s3, "\n[ Espacio reservado para sello ]", italic=True, size_pt=6.8, color_rgb=(148,163,184), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(c_s3, "Protección Civil y Bomberos Medellín", bold=True, size_pt=6.5, color_rgb=(10,35,66), align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4)
    
    # 9. Banner Inferior Oficial
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(2)
    p_foot.paragraph_format.space_after = Pt(0)
    p_foot.paragraph_format.line_spacing = 1.0
    r_foot = p_foot.add_run()
    r_foot.add_picture('Documentacion/banner_inferior_oficial.png', width=W_TOTAL, height=Inches(0.48))
    
    out_docx = 'Documentacion/Formato_Deslinde_Responsabilidad_Abejas.docx'
    doc.save(out_docx)
    print(f'Documento Word generado en: {out_docx}')

if __name__ == '__main__':
    construir_formato_expandido()
