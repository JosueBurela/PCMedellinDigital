import sqlite3
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def crear_excel_profesional():
    print("Iniciando generación de Excel Maestro Anual...")
    conn = sqlite3.connect('whatsapp_messages.db')
    df_todo = pd.read_sql_query('SELECT * FROM servicios_anuales_2026 ORDER BY mes_num ASC, fecha ASC, hora ASC', conn)
    conn.close()

    # Crear Workbook
    wb = openpyxl.Workbook()
    # Eliminar hoja por defecto
    wb.remove(wb.active)

    # Estilos Institucionales
    fuente_titulo = Font(name='Calibri', size=16, bold=True, color='0A2342')
    fuente_subtitulo = Font(name='Calibri', size=11, bold=True, color='555555')
    fuente_director = Font(name='Calibri', size=11, italic=True, color='D62828')
    
    fuente_encabezado = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    fill_encabezado_navy = PatternFill(start_color='0A2342', end_color='0A2342', fill_type='solid')
    fill_encabezado_rojo = PatternFill(start_color='D62828', end_color='D62828', fill_type='solid')
    
    fill_zebra = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
    fill_blanco = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    fill_total = PatternFill(start_color='E9ECEF', end_color='E9ECEF', fill_type='solid')
    
    fuente_datos = Font(name='Calibri', size=10)
    fuente_total = Font(name='Calibri', size=10, bold=True, color='0A2342')
    
    borde_delgado = Border(
        left=Side(style='thin', color='D0D5DD'),
        right=Side(style='thin', color='D0D5DD'),
        top=Side(style='thin', color='D0D5DD'),
        bottom=Side(style='thin', color='D0D5DD')
    )
    borde_total = Border(
        top=Side(style='thin', color='0A2342'),
        bottom=Side(style='double', color='0A2342')
    )
    
    align_centro = Alignment(horizontal='center', vertical='center')
    align_izq = Alignment(horizontal='left', vertical='center')
    align_der = Alignment(horizontal='right', vertical='center')

    # =========================================================================
    # 1. HOJA: RESUMEN EJECUTIVO (DASHBOARD)
    # =========================================================================
    print("Creando hoja: 00_Resumen_Ejecutivo...")
    ws_resumen = wb.create_sheet(title='00_Resumen_Ejecutivo')
    ws_resumen.views.sheetView[0].showGridLines = True
    
    # Encabezado Institucional
    ws_resumen['B2'] = 'H. AYUNTAMIENTO DE MEDELLÍN DE BRAVO, VERACRUZ'
    ws_resumen['B2'].font = Font(name='Calibri', size=13, bold=True, color='D62828')
    
    ws_resumen['B3'] = 'DIRECCIÓN DE PROTECCIÓN CIVIL Y BOMBEROS MUNICIPALES'
    ws_resumen['B3'].font = fuente_titulo
    
    ws_resumen['B4'] = 'TABLERO CONSOLIDADO ANUAL DE SERVICIOS Y EMERGENCIAS 2026'
    ws_resumen['B4'].font = fuente_subtitulo
    
    ws_resumen['B5'] = 'Director Titular: Lic. Daniel Eduardo Romero Pilar | Periodo: 01 de Enero al 10 de Septiembre de 2026'
    ws_resumen['B5'].font = fuente_director
    
    # Tabla 1: Resumen por Mes
    ws_resumen['B7'] = '1. CONCENTRADO MENSUAL OPERATIVO'
    ws_resumen['B7'].font = Font(name='Calibri', size=11, bold=True, color='0A2342')
    
    headers_mes = ['Mes', 'Servicios Totales', 'Efectivos', 'No Efectivos / Cancelados', '% Efectividad', '% del Total Anual']
    for col_idx, h in enumerate(headers_mes, start=2):
        cell = ws_resumen.cell(row=8, column=col_idx, value=h)
        cell.font = fuente_encabezado
        cell.fill = fill_encabezado_navy
        cell.alignment = align_centro
        cell.border = borde_delgado

    meses_info = [
        ('Enero', 1, 31), ('Febrero', 2, 28), ('Marzo', 3, 31),
        ('Abril', 4, 30), ('Mayo', 5, 31), ('Junio', 6, 30),
        ('Julio', 7, 31), ('Agosto', 8, 31), ('Septiembre (al 10)', 9, 10)
    ]
    
    total_gral = len(df_todo)
    fila_actual = 9
    for m_nom, m_num, d_dias in meses_info:
        sub = df_todo[df_todo['mes_num'] == m_num]
        tot = len(sub)
        efec = len(sub[sub['efectividad'] == 'Efectivo'])
        no_efec = len(sub[sub['efectividad'] != 'Efectivo'])
        pct_efec = (efec / tot * 100) if tot > 0 else 0
        pct_anual = (tot / total_gral * 100)
        
        ws_resumen.cell(row=fila_actual, column=2, value=m_nom).alignment = align_izq
        ws_resumen.cell(row=fila_actual, column=3, value=tot).alignment = align_centro
        ws_resumen.cell(row=fila_actual, column=4, value=efec).alignment = align_centro
        ws_resumen.cell(row=fila_actual, column=5, value=no_efec).alignment = align_centro
        ws_resumen.cell(row=fila_actual, column=6, value=f"{pct_efec:.1f}%").alignment = align_centro
        ws_resumen.cell(row=fila_actual, column=7, value=f"{pct_anual:.1f}%").alignment = align_centro
        
        for c in range(2, 8):
            ws_resumen.cell(row=fila_actual, column=c).font = fuente_datos
            ws_resumen.cell(row=fila_actual, column=c).border = borde_delgado
            if fila_actual % 2 == 1:
                ws_resumen.cell(row=fila_actual, column=c).fill = fill_zebra
        fila_actual += 1
        
    # Fila de Totales
    ws_resumen.cell(row=fila_actual, column=2, value='TOTAL GENERAL').alignment = align_izq
    ws_resumen.cell(row=fila_actual, column=3, value=total_gral).alignment = align_centro
    tot_efec = len(df_todo[df_todo['efectividad'] == 'Efectivo'])
    tot_no_efec = len(df_todo[df_todo['efectividad'] != 'Efectivo'])
    ws_resumen.cell(row=fila_actual, column=4, value=tot_efec).alignment = align_centro
    ws_resumen.cell(row=fila_actual, column=5, value=tot_no_efec).alignment = align_centro
    ws_resumen.cell(row=fila_actual, column=6, value=f"{tot_efec/total_gral*100:.1f}%").alignment = align_centro
    ws_resumen.cell(row=fila_actual, column=7, value="100.0%").alignment = align_centro
    
    for c in range(2, 8):
        cell = ws_resumen.cell(row=fila_actual, column=c)
        cell.font = fuente_total
        cell.fill = fill_total
        cell.border = borde_total
        
    # Tabla 2: Por Categoría Macro (Lado Derecho)
    ws_resumen['I7'] = '2. DISTRIBUCIÓN POR CATEGORÍA OFICIAL'
    ws_resumen['I7'].font = Font(name='Calibri', size=11, bold=True, color='0A2342')
    
    headers_cat = ['Categoría Macro', 'Servicios', '% del Total']
    for col_idx, h in enumerate(headers_cat, start=9):
        cell = ws_resumen.cell(row=8, column=col_idx, value=h)
        cell.font = fuente_encabezado
        cell.fill = fill_encabezado_navy
        cell.alignment = align_centro
        cell.border = borde_delgado
        
    cat_counts = df_todo['categoria_macro'].value_counts()
    fila_cat = 9
    for cat_name, cnt in cat_counts.items():
        ws_resumen.cell(row=fila_cat, column=9, value=cat_name).alignment = align_izq
        ws_resumen.cell(row=fila_cat, column=10, value=cnt).alignment = align_centro
        ws_resumen.cell(row=fila_cat, column=11, value=f"{cnt/total_gral*100:.1f}%").alignment = align_centro
        
        for c in range(9, 12):
            ws_resumen.cell(row=fila_cat, column=c).font = fuente_datos
            ws_resumen.cell(row=fila_cat, column=c).border = borde_delgado
            if fila_cat % 2 == 1:
                ws_resumen.cell(row=fila_cat, column=c).fill = fill_zebra
        fila_cat += 1
        
    # Total Categoría
    ws_resumen.cell(row=fila_cat, column=9, value='TOTAL').alignment = align_izq
    ws_resumen.cell(row=fila_cat, column=10, value=total_gral).alignment = align_centro
    ws_resumen.cell(row=fila_cat, column=11, value="100.0%").alignment = align_centro
    for c in range(9, 12):
        cell = ws_resumen.cell(row=fila_cat, column=c)
        cell.font = fuente_total
        cell.fill = fill_total
        cell.border = borde_total

    # Autoajuste de columnas en Resumen
    for col in ws_resumen.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws_resumen.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # =========================================================================
    # FUNCIÓN HELPER PARA CREAR HOJAS DE DETALLE
    # =========================================================================
    def agregar_hoja_datos(nombre_hoja, df_datos, titulo_hoja):
        print(f"Creando hoja: {nombre_hoja} ({len(df_datos)} registros)...")
        ws = wb.create_sheet(title=nombre_hoja)
        ws.views.sheetView[0].showGridLines = True
        
        # Título de Hoja
        ws['A1'] = 'PROTECCIÓN CIVIL Y BOMBEROS MEDELLÍN DE BRAVO'
        ws['A1'].font = Font(name='Calibri', size=11, bold=True, color='D62828')
        
        ws['A2'] = titulo_hoja
        ws['A2'].font = Font(name='Calibri', size=14, bold=True, color='0A2342')
        
        ws['A3'] = f'Total de Registros: {len(df_datos)} servicios | Titular: Lic. Daniel Eduardo Romero Pilar'
        ws['A3'].font = Font(name='Calibri', size=10, italic=True, color='555555')
        
        headers = [
            'ID', 'Mes', 'Fecha', 'Hora', 'Categoría Macro',
            'Subtipo de Servicio', 'Localidad / Cuadrante', 'Efectividad',
            'Fuente de Información', 'Descripción / Resumen Institucional'
        ]
        
        fila_h = 5
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=fila_h, column=col_idx, value=h)
            cell.font = fuente_encabezado
            cell.fill = fill_encabezado_navy if 'Anual' in nombre_hoja else fill_encabezado_rojo
            cell.alignment = align_centro
            cell.border = borde_delgado
            
        r_idx = fila_h + 1
        for _, row in df_datos.iterrows():
            ws.cell(row=r_idx, column=1, value=int(row['id'])).alignment = align_centro
            ws.cell(row=r_idx, column=2, value=str(row['mes'])).alignment = align_centro
            ws.cell(row=r_idx, column=3, value=str(row['fecha'])).alignment = align_centro
            ws.cell(row=r_idx, column=4, value=str(row['hora'])).alignment = align_centro
            ws.cell(row=r_idx, column=5, value=str(row['categoria_macro'])).alignment = align_izq
            ws.cell(row=r_idx, column=6, value=str(row['subtipo_servicio'])).alignment = align_izq
            ws.cell(row=r_idx, column=7, value=str(row['localidad'])).alignment = align_izq
            ws.cell(row=r_idx, column=8, value=str(row['efectividad'])).alignment = align_centro
            ws.cell(row=r_idx, column=9, value=str(row['fuente_datos'])).alignment = align_izq
            ws.cell(row=r_idx, column=10, value=str(row['resumen_servicio'])).alignment = align_izq
            
            # Formato y zebra
            for c in range(1, 11):
                cell = ws.cell(row=r_idx, column=c)
                cell.font = fuente_datos
                cell.border = borde_delgado
                if r_idx % 2 == 1:
                    cell.fill = fill_zebra
            r_idx += 1
            
        # Filtros automáticos
        ws.auto_filter.ref = f"A{fila_h}:J{r_idx-1}"
        
        # Autoajuste de columnas
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            if col[0].column == 10:  # Descripción más ancha
                ws.column_dimensions[col_letter].width = 65
            elif col[0].column in [5, 6, 7]:
                ws.column_dimensions[col_letter].width = 32
            elif col[0].column in [8, 9]:
                ws.column_dimensions[col_letter].width = 24
            elif col[0].column in [3, 4]:
                ws.column_dimensions[col_letter].width = 14
            else:
                ws.column_dimensions[col_letter].width = 10

    # =========================================================================
    # 2. HOJA: ANUAL CONSOLIDADO (TODO EL AÑO)
    # =========================================================================
    agregar_hoja_datos('Anual_Consolidado', df_todo, 'Padrón Maestro Anual de Servicios y Despachos 2026')

    # =========================================================================
    # 3. HOJAS INDIVIDUALES POR MES
    # =========================================================================
    nombres_meses = [
        (1, '01_Enero', 'Enero'),
        (2, '02_Febrero', 'Febrero'),
        (3, '03_Marzo', 'Marzo'),
        (4, '04_Abril', 'Abril'),
        (5, '05_Mayo', 'Mayo'),
        (6, '06_Junio', 'Junio'),
        (7, '07_Julio', 'Julio'),
        (8, '08_Agosto', 'Agosto'),
        (9, '09_Septiembre', 'Septiembre (01 al 10 de Septiembre)')
    ]
    
    for m_num, sheet_name, mes_label in nombres_meses:
        df_mes = df_todo[df_todo['mes_num'] == m_num]
        agregar_hoja_datos(sheet_name, df_mes, f'Bitácora Operativa Mensual: {mes_label} 2026')

    # Guardar en Documentacion
    archivo_salida = 'Documentacion/Registro_Anual_Operativo_2026_PC_Medellin.xlsx'
    wb.save(archivo_salida)
    print(f"\n=======================================================")
    print(f"  EXCEL MAESTRO ANUAL GENERADO CON ÉXITO")
    print(f"  Total de Hojas creadas: {len(wb.sheetnames)}")
    print(f"  Hojas: {wb.sheetnames}")
    print(f"  Guardado en: {archivo_salida}")
    print(f"=======================================================")

if __name__ == '__main__':
    crear_excel_profesional()
