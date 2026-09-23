import psycopg2
import re
import os
import sys
from datetime import datetime, timedelta
from collections import defaultdict, Counter
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList

sys.stdout.reconfigure(encoding='utf-8')

print("=" * 80)
print("GENERADOR DEL EXPEDIENTE EJECUTIVO ANUAL Y CATÁLOGO MAESTRO DE INCENDIOS 2026")
print("Cuerpo de Bomberos y Protección Civil - Medellín de Bravo, Veracruz")
print("=" * 80)

# 1. CONEXIÓN A BASE DE DATOS POSTGRESQL
conn = psycopg2.connect("postgresql://postgres:postgres@localhost:5432/evolution")
cur = conn.cursor()

GROUP_JID = '120363042493725288@g.us'

cur.execute("""
    SELECT 
        "id",
        to_timestamp("messageTimestamp") as fecha,
        "pushName",
        "messageType",
        COALESCE(
            "message"->>'conversation',
            "message"->'extendedTextMessage'->>'text',
            "message"->'imageMessage'->>'caption',
            "message"->'videoMessage'->>'caption',
            ''
        ) as texto,
        "message"->'imageMessage' IS NOT NULL as tiene_imagen,
        "key"->>'id' as wa_id,
        "key"->>'participant' as remitente
    FROM "Message"
    WHERE "key"->>'remoteJid' = %s
      AND to_timestamp("messageTimestamp") >= '2026-01-01 00:00:00'
    ORDER BY "messageTimestamp" ASC
""", (GROUP_JID,))

rows = cur.fetchall()
print(f"Total mensajes históricos 2026 recuperados de BD: {len(rows)}")

# 2. DEFINICIÓN DE FILTROS Y REGLAS DE DETECCIÓN
regex_fuego = re.compile(
    r'\b(incendio|incendios|quema|quemas|pastizal|pastizales|fuego|conato|conatos|humo|sofoca|sofocad[oa]s?|sofocaci[oó]n|bomberos?|llantas?|basurero|basura|chispas?|transformador|fuga.*gas|gas\s+lp|tanque.*fuga|olor\s+a\s+gas)\b',
    re.IGNORECASE
)

def es_falso_positivo(t_low):
    if any(k in t_low for k in ['darte fuego', 'echar fuego', 'fuego amigo', 'darle fuego']):
        return True
    if 'hacer fuego' in t_low and not any(k in t_low for k in ['pastizal', 'basura', 'casa', 'lote']):
        return True
    if ('quemadura de' in t_low or 'quemaduras' in t_low or 'quemadura por' in t_low) and not any(k in t_low for k in ['pastizal', 'casa', 'basura', 'fuego', 'sofoc', 'pipa', 'tanque', 'gas']):
        return True
    if 'lavar piso' in t_low or 'lavar hasta el domo' in t_low:
        if not any(k in t_low for k in ['incendio', 'fuego', 'pastizal', 'conato', 'sofoc']):
            return True
    if 'llevar agua' in t_low and not any(k in t_low for k in ['incendio', 'fuego', 'pastizal', 'conato', 'sofoc', 'pipa']):
        return True
    if 'fuerza de tarea' in t_low and not any(k in t_low for k in ['incendio', 'fuego', 'pastizal', 'quema', 'conato', 'sofoc', 'fuga']):
        return True
    if 'gasolina' in t_low and not any(k in t_low for k in ['incendio', 'fuego', 'pastizal', 'quema', 'conato', 'sofoc', 'fuga', 'tanque']):
        return True
    return False

def extraer_localidad(texto):
    t = texto.lower()
    if 'puente moreno' in t or 'lagos de puente' in t:
        return "Fracc. Lagos de Puente Moreno"
    elif 'arboledas san ram' in t or 'san ramon' in t or 'san ramón' in t:
        return "Fracc. Arboledas San Ramón"
    elif 'arboledas san miguel' in t or 'san miguel' in t:
        return "Fracc. Arboledas San Miguel"
    elif 'el tejar' in t or 'tejar' in t or 'atrancon' in t or 'atrancón' in t:
        return "Localidad El Tejar"
    elif 'paso del toro' in t or 'paso de toro' in t:
        return "Localidad Paso del Toro"
    elif 'los robles' in t:
        return "Localidad Los Robles"
    elif 'ixcualco' in t or 'ixcoalco' in t:
        return "Localidad Ixcualco"
    elif 'rancho del padre' in t or 'rancho de el padre' in t:
        return "Rancho del Padre"
    elif 'playa de vacas' in t or 'playaca' in t:
        return "Localidad Playa de Vacas"
    elif 'moralillo' in t:
        return "Localidad El Moralillo"
    elif 'la candelaria' in t or 'candelaria' in t:
        return "Localidad La Candelaria"
    elif 'las palmas' in t:
        return "Fracc. Las Palmas"
    elif 'las vegas' in t:
        return "Fracc. Las Vegas"
    elif 'san jose novillero' in t or 'novillero' in t:
        return "San José Novillero"
    elif 'juan de alfaro' in t or 'juan alfaro' in t:
        return "Localidad Juan de Alfaro"
    elif 'mangal' in t:
        return "Localidad El Mangal"
    elif 'la posta' in t or 'posta' in t:
        return "Zona La Posta"
    elif 'casa blanca' in t:
        return "Localidad Casa Blanca"
    elif 'sal si puedes' in t:
        return "Localidad Sal si Puedes"
    elif 'dos bocas' in t:
        return "Localidad Dos Bocas"
    elif 'alvaradito' in t:
        return "Localidad Alvaradito"
    elif 'rancho nuevo' in t or 'el 12' in t or 'km 12' in t or 'km-12' in t:
        return "Rancho Nuevo / Km 12"
    elif 'la bascula' in t or 'la báscula' in t:
        return "Localidad La Báscula"
    elif 'pichones' in t:
        return "Zona Puente los Pichones"
    elif 'santa fe' in t:
        return "Carretera Santa Fe - Paso del Toro"
    elif 'la gloria' in t:
        return "Localidad La Gloria"
    elif 'las balsas' in t or 'balsas' in t:
        return "Localidad Las Balsas"
    elif 'arrieros' in t or 'los arrieron' in t:
        return "Localidad Los Arrieros"
    elif 'aaron proal' in t or 'aarón proal' in t:
        return "Colonia Aarón Proal"
    elif 'crucero de medell' in t:
        return "Crucero de Medellín"
    elif 'centro de transferencia' in t or 'centro de acopio' in t:
        return "Centro de Transferencia / Acopio"
    elif 'aurrera' in t or 'bodega aurrera' in t:
        return "Zona Bodega Aurrerá"
    elif 'costera' in t or 'carretera veracruz alvarado' in t or 'carretera del golfo' in t:
        return "Carretera Federal Veracruz - Alvarado"
    elif 'boca del río' in t or 'conurbados' in t:
        return "Zona Limítrofe Medellín - Boca del Río"
    elif 'medellin de bravo' in t or 'cabecera' in t:
        return "Cabecera Municipal Medellín"
    
    m = re.search(r'(?:en\s+|colonia\s+|fracc\w*\s+|localidad\s+|carretera\s+|calle\s+|camino\s+)([A-Za-zÁÉÍÓÚáéíóúñÑ0-9\s]{4,30}?)(?:,|\.|\s+sale|\s+se|\s+con|\s+por|\s+entre|\n|$)', texto, re.IGNORECASE)
    if m:
        cand = m.group(1).strip()
        if len(cand) > 3 and not cand.lower().startswith(('el punto', 'apoyo', 'un pastizal', 'un incendio', 'camino a')):
            return cand.title()
    return "Medellín de Bravo (Municipio)"

def clasificar_siniestro(texto):
    t = texto.lower()
    if any(k in t for k in ['fuga de gas', 'fuga en cilindro', 'fuga en tanque', 'olor a gas', 'tanke de 20', 'tanque de 20', 'tanque de 30', 'tanque estacionario', 'estufa']):
        return "Fuga de Gas LP / Manejo de Tanque"
    elif any(k in t for k in ['transformador', 'chispas', 'corto circuito', 'cortocircuito', 'cable', 'mufa', 'alcantarilla de cfe', 'bfa']):
        return "Conato Eléctrico / Cortocircuito / CFE"
    elif any(k in t for k in ['casa habitaci', 'casa aviaci', 'habitaci', 'domicilio', 'vivienda', 'departamento']):
        return "Incendio en Casa Habitación / Vivienda"
    elif any(k in t for k in ['veh[ií]culo', 'auto\b', 'camioneta', 'trailer', 'tráiler', 'suv', 'automóvil', 'automovil', 'motocicleta']) or (re.search(r'\b(carro|moto)\b', t) and 'motobomba' not in t):
        return "Incendio de Vehículo / Automotriz"
    elif re.search(r'\b(comercio|comercial|taller|bodega|tienda|palapa|restaurante)\b', t) and 'localidad' not in t:
        return "Incendio en Comercio / Taller / Bodega"
    elif any(k in t for k in ['basura', 'basurero', 'llanta', 'centro de acopio', 'desechos', 'centro de transferencia']):
        return "Quema de Basura / Residuos Sólidos"
    elif any(k in t for k in ['lote bald', 'terreno bald', 'baldío']):
        return "Incendio de Lote Baldío"
    else:
        return "Incendio de Pastizal / Maleza"

def extraer_unidades(texto):
    t = texto.lower()
    unidades = []
    if '072' in t or 'u-072' in t or 'u072' in t or '72' in t:
        unidades.append("Unidad 072 (Bomberos)")
    if '073' in t or 'u-073' in t or 'u073' in t or '73' in t:
        unidades.append("Unidad 073 (Cisterna)")
    if 'pipa' in t or '20 mil litros' in t or '20,000' in t or 'cisterna' in t:
        if "Unidad 073 (Cisterna)" not in unidades and "Pipa de Bomberos (20k Lts)" not in unidades:
            unidades.append("Pipa de Bomberos (20k Lts)")
    if '096' in t or 'u-096' in t or 'colorado' in t:
        unidades.append("Unidad 096 (Rescate / Apoyo)")
    if '208' in t or 'u_208' in t or 'u-208' in t:
        unidades.append("Unidad 208 (Ambulancia)")
    if '097' in t or 'u-097' in t:
        unidades.append("Unidad 097 (Ambulancia)")
    if '098' in t or 'u-098' in t:
        unidades.append("Unidad 098 (Ambulancia)")
    if 'motorizada' in t or re.search(r'\bmoto\b', t) and 'motobomba' not in t:
        unidades.append("Unidad Motorizada PC")
    if not unidades:
        unidades.append("Cuerpo de Bomberos PC (Unidad 072)")
    return ", ".join(unidades)

def extraer_recurso_hidrico(texto, tipo):
    t = texto.lower()
    if 'pipa' in t or 'cisterna' in t or '20 mil' in t or '20,000' in t:
        return "Pipa Cisterna (20,000 Lts) y Línea de Mangueras"
    elif '072' in t or '073' in t or 'ataque' in t or 'motobomba' in t:
        return "Línea Presión U-072 / Ataque Rápido"
    elif 'manual' in t or 'aspersor' in t or 'abatefuego' in t or 'a mano' in t:
        return "Sofocación Manual (Mochilas / Herramienta)"
    elif 'extintor' in t or 'pqs' in t:
        return "Extintor PQS / Maniobra Técnica"
    elif tipo in ["Fuga de Gas LP / Manejo de Tanque", "Conato Eléctrico / Cortocircuito / CFE"]:
        return "Aseguramiento y Ventilación Preventiva"
    elif 'falsa alarma' in t or 'sin novedad' in t or 'negativo de incendio' in t:
        return "Sin Requerimiento Hídrico (Inspección)"
    else:
        return "Línea de Mangueras / Ataque Rápido"

def extraer_estatus(texto):
    t = texto.lower()
    if 'falsa alarma' in t or 'no se visualiza nada' in t or 'no hay ning' in t or 'sin novedad' in t or 'negativo de incendio' in t:
        return "Falsa Alarma / Sin Riesgo al Arribo"
    elif any(k in t for k in ['sofocado y liquidado', 'liquidado 100%', 'sofocado 100%', 'keda sofocado', 'queda sofocado', 'termina servicio', 'termino']):
        return "Sofocado y Liquidado 100%"
    elif any(k in t for k in ['controlado', 'keda controlado', 'queda controlado']):
        return "Controlado y Enfriado"
    elif any(k in t for k in ['tanque asegurado', 'retira tanque', 'se traslada a base']):
        return "Riesgo Mitigado / Tanque Asegurado"
    elif 'fuga' in t and any(k in t for k in ['controlada', 'cerrada', 'revisada', 'checa']):
        return "Fuga Controlada / Sin Riesgo"
    return "Servicio Atendido y Concluido"

def extraer_turno(hora_str):
    try:
        h = int(hora_str.split(':')[0])
        if 7 <= h < 15:
            return "Matutino (07:00 - 14:59)"
        elif 15 <= h < 23:
            return "Vespertino (15:00 - 22:59)"
        else:
            return "Nocturno (23:00 - 06:59)"
    except:
        return "Matutino (07:00 - 14:59)"

def extraer_despachador(texto):
    t = texto.lower()
    if 'c5' in t or '9-1-1' in t or '911' in t:
        return "C5 / 911 Estatal"
    elif any(k in t for k in ['polic', 'seguridad pública', 'mando único']):
        return "Policía Municipal / Estatal"
    elif any(k in t for k in ['reporte ciudadano', 'llama a base', 'vecinos', 'reportante:']):
        return "Reporte Ciudadano a Base"
    return "Base de Protección Civil"

def extraer_coordinacion(texto):
    t = texto.lower()
    coords = []
    if 'conurbado' in t:
        coords.append("Bomberos Conurbados Boca del Río")
    if 'spc' in t or 'protección civil estatal' in t or 'enlace de la spc' in t:
        coords.append("Protección Civil Estatal (SPC)")
    if 'c5' in t or '911' in t:
        coords.append("Centro de Comando C5")
    if 'polic' in t:
        coords.append("Policía Municipal / Estatal")
    if 'cfe' in t:
        coords.append("Comisión Federal de Electricidad (CFE)")
    if not coords:
        return "Operación Directa Bomberos Medellín"
    return ", ".join(coords)

def extraer_afectacion(texto, estatus):
    t = texto.lower()
    if estatus == "Falsa Alarma / Sin Riesgo al Arribo":
        return "Sin Afectaciones (Falsa Alarma)"
    if any(k in t for k in ['crisis', 'inhalaci', 'paramédico', 'u-208', 'u-097', 'ambulancia']):
        return "Atención Médica por Crisis / Inhalación"
    if any(k in t for k in ['calcinad', 'fallecid', 'código 14', 'era 14']):
        return "Víctima Mortal en Sitio (Código 14)"
    if any(k in t for k in ['pérdida total', 'afectación material', 'daño estructural']):
        return "Afectación Material Estructural"
    return "Sin Víctimas ni Lesionados (Saldo Blanco)"

def extraer_personal(texto, push):
    t = texto.lower()
    if 'director' in t or 'toño' in t or 'burela' in t:
        return "Comandante Toño Burela (Director)"
    elif 'aurelio' in t:
        return "Oficial Aurelio (Bomberos)"
    elif 'chayanne' in t:
        return "Operador Chayanne (Bomberos)"
    elif 'chapo' in t or 'luis alberto' in t:
        return "Oficial Luis Alberto Chapo (PC)"
    elif 'rodolfo' in t or 'gato' in t:
        return "Oficial Rodolfo Gato (PC)"
    elif 'santos' in t:
        return "Oficial Santos Álvarez (Bomberos)"
    elif 'charly' in t:
        return "Operador Charly (Bomberos)"
    elif push and push.strip() and push not in ['Você', 'PC&B Medellín De Bravo']:
        return f"Oficial {push.strip()}"
    return "Guardia de Bomberos PC Medellín"

# 3. EXTRACCIÓN MULTI-MODO DE LA BASE DE DATOS
raw_items = []
seen_dedup = set()

for r in rows:
    mid, dt, push, mtype, txt, tiene_img, wa_id, rem = r
    if not txt or len(txt.strip()) < 8:
        continue
    t_clean = ' '.join(txt.split())
    dedup_key = (dt.strftime('%Y-%m-%d %H:%M'), t_clean[:50].lower())
    if dedup_key in seen_dedup:
        continue
    seen_dedup.add(dedup_key)
    
    t_low = txt.lower()
    if es_falso_positivo(t_low):
        continue
        
    is_guard = (
        any(k in t_low for k in ['novedades del d', 'reportes de servicios pc', 'novedades de pc', 'reporte de guardia', 'novedades de la guardia'])
        and any(s in t_low for s in ['sale unidad', 'sale el compañero', 'salida', '1_', '2_', '1.-', '1.-'])
    )
    
    if is_guard:
        lines = [l.strip() for l in txt.split('\n') if l.strip()]
        header = lines[0] if lines else ""
        m_shift_date = re.search(r'(\d{1,2})[-/](\d{1,2})[-/](20\d{2}|\d{2})', header)
        d_str = dt.strftime('%Y-%m-%d')
        if m_shift_date:
            d, m, y = m_shift_date.groups()
            if len(y) == 2: y = f"20{y}"
            if y == "2025" and dt.year == 2026: y = "2026"
            d_str = f"{y}-{int(m):02d}-{int(d):02d}"
        for l in lines[1:]:
            l_low = l.lower()
            if regex_fuego.search(l_low) and not es_falso_positivo(l_low):
                m_t = re.search(r'\b([012]?\d:[0-5]\d)\b', l)
                t_str = m_t.group(1) if m_t else dt.strftime('%H:%M')
                raw_items.append({
                    'origen': 'GUARDIA',
                    'fecha_str': d_str,
                    'hora_str': t_str,
                    'dt': dt,
                    'texto': l,
                    'push': push or "Guardia PC",
                    'tiene_img': tiene_img
                })
    else:
        if regex_fuego.search(t_low):
            if any(k in t_low for k in ['sale', 'salida', 'acude', 'sofoc', 'control', 'apoy', 'en ruta', 'fuego', 'incendio', 'quema', 'kema', 'fuga', 'tanque', 'pipa', '072', '073', '096', 'c5', '911', 'reportan', 'arrib', 'humo', 'alarma']):
                raw_items.append({
                    'origen': 'STANDALONE',
                    'fecha_str': dt.strftime('%Y-%m-%d'),
                    'hora_str': dt.strftime('%H:%M'),
                    'dt': dt,
                    'texto': txt.strip(),
                    'push': push or "Personal Operativo",
                    'tiene_img': tiene_img
                })

print(f"Total registros raw recopilados: {len(raw_items)}")

# 4. ALGORITMO DE CORRELACIÓN ESPACIAL Y TEMPORAL POR DÍA
by_date = defaultdict(list)
for it in raw_items:
    by_date[it['fecha_str']].append(it)

incidentes = []

for f_str, day_items in sorted(by_date.items()):
    def get_time_val(item):
        try:
            return datetime.strptime(item['hora_str'], "%H:%M")
        except:
            return datetime.strptime("12:00", "%H:%M")
    day_items_sorted = sorted(day_items, key=get_time_val)
    
    current_clusters = []
    for it in day_items_sorted:
        loc = extraer_localidad(it['texto'])
        tipo = clasificar_siniestro(it['texto'])
        t_it = get_time_val(it)
        
        is_status_only = any(k in it['texto'].lower() for k in ['retorna a base', 'retorno a base', 'cargando agua', 'en el punto', 'queda sofocado', 'keda sofocado', 'controlado']) and len(it['texto']) < 80
        
        best_cluster = None
        for cl in reversed(current_clusters):
            t_cl = get_time_val(cl['anchor_item'])
            diff_hours = abs((t_it - t_cl).total_seconds()) / 3600.0
            
            if loc and loc == cl['localidad'] and diff_hours <= 4.0:
                best_cluster = cl
                break
            if (not loc or is_status_only) and diff_hours <= 3.0:
                best_cluster = cl
                break
            if (it['origen'] != cl['anchor_item']['origen']) and (tipo == cl['tipo']) and diff_hours <= 2.5:
                best_cluster = cl
                break
                
        if best_cluster:
            best_cluster['items'].append(it)
            if best_cluster['localidad'] == "Medellín de Bravo (Municipio)" and loc != "Medellín de Bravo (Municipio)":
                best_cluster['localidad'] = loc
            if it['tiene_img']:
                best_cluster['tiene_img'] = True
        else:
            current_clusters.append({
                'fecha_str': f_str,
                'anchor_item': it,
                'hora_str': it['hora_str'],
                'localidad': loc,
                'tipo': tipo,
                'tiene_img': it['tiene_img'],
                'items': [it]
            })
            
    incidentes.extend(current_clusters)

records = []
for idx, inc in enumerate(incidentes, 1):
    folio = f"INC-2026-{idx:04d}"
    all_text = " || ".join(x['texto'].replace('\n', ' ') for x in inc['items'])
    unidades = extraer_unidades(all_text)
    tipo = inc['tipo']
    loc = inc['localidad']
    estatus = extraer_estatus(all_text)
    recurso = extraer_recurso_hidrico(all_text, tipo)
    afectacion = extraer_afectacion(all_text, estatus)
    despachador = extraer_despachador(all_text)
    coordinacion = extraer_coordinacion(all_text)
    personal = extraer_personal(all_text, inc['items'][0]['push'])
    turno = extraer_turno(inc['hora_str'])
    
    # Resumen operativo oficial detallado
    resumen = all_text
    
    records.append({
        'No_Folio': folio,
        'Fecha': inc['fecha_str'],
        'Hora': inc['hora_str'],
        'Turno_Operativo': turno,
        'Unidad_Bomberos': unidades,
        'Personal_A_Cargo': personal,
        'Tipo_Siniestro': tipo,
        'Estatus_Servicio': estatus,
        'Recursos_Despliegue': recurso,
        'Afectacion_Personas': afectacion,
        'Ubicacion_Localidad': loc,
        'Origen_Despacho': despachador,
        'Evidencia_Foto': 'SÍ' if inc['tiene_img'] else 'NO',
        'Coordinacion_Externa': coordinacion,
        'Resumen_Operativo': resumen
    })

df_anual = pd.DataFrame(records)
print(f"Total servicios de incendios y siniestros consolidados: {len(df_anual)}")

# 5. CREACIÓN DEL LIBRO EXCEL OFICIAL CON OPENPYXL
wb = openpyxl.Workbook()

# Colores y Fuentes Institucionales de Bomberos
RED_FIRE_DARK = "7F1D1D"     # Borgoña Fuego Oficial
RED_FIRE_MED = "991B1B"      # Rojo Carmesí Institucional
NAVY_TITLE = "1B365D"        # Azul Marino Encabezados
GRAY_BORDER = "D1D5DB"
ZEBRA_FILL = "FFF5F5"        # Fondo suave rojizo/cálido
WHITE = "FFFFFF"

fill_table_hdr = PatternFill(start_color=RED_FIRE_DARK, end_color=RED_FIRE_DARK, fill_type="solid")
font_table_hdr = Font(name='Arial', size=10, bold=True, color=WHITE)
font_data = Font(name='Arial', size=9)
font_data_bold = Font(name='Arial', size=9, bold=True)
font_title = Font(name='Arial', size=14, bold=True, color=NAVY_TITLE)
font_subtitle = Font(name='Arial', size=10, italic=True, color="4B5563")
font_sec_header = Font(name='Arial', size=11, bold=True, color=RED_FIRE_MED)

fill_zebra = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
fill_white = PatternFill(start_color=WHITE, end_color=WHITE, fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color=GRAY_BORDER),
    right=Side(style='thin', color=GRAY_BORDER),
    top=Side(style='thin', color=GRAY_BORDER),
    bottom=Side(style='thin', color=GRAY_BORDER)
)

def crear_kpi_card(ws, r_start, c_start, r_end, c_end, titulo, valor, subtitulo, bg_color, text_color):
    ws.merge_cells(start_row=r_start, start_column=c_start, end_row=r_start, end_column=c_end)
    ws.merge_cells(start_row=r_start+1, start_column=c_start, end_row=r_start+1, end_column=c_end)
    ws.merge_cells(start_row=r_start+2, start_column=c_start, end_row=r_start+2, end_column=c_end)
    
    c_tit = ws.cell(r_start, c_start, titulo)
    c_tit.font = Font(name='Arial', size=8, bold=True, color=text_color)
    c_tit.alignment = Alignment(horizontal='center', vertical='center')
    
    c_val = ws.cell(r_start+1, c_start, valor)
    c_val.font = Font(name='Arial', size=16, bold=True, color=text_color)
    c_val.alignment = Alignment(horizontal='center', vertical='center')
    
    c_sub = ws.cell(r_start+2, c_start, subtitulo)
    c_sub.font = Font(name='Arial', size=7, italic=True, color=text_color)
    c_sub.alignment = Alignment(horizontal='center', vertical='center')
    
    fill_card = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    for r in range(r_start, r_start+3):
        for c in range(c_start, c_end+1):
            ws.cell(r, c).fill = fill_card
            ws.cell(r, c).border = thin_border

# ======================================================================
# HOJA 1: 📊 DASHBOARD EJECUTIVO BOMBEROS
# ======================================================================
ws1 = wb.active
ws1.title = "📊 Dashboard Bomberos"
ws1.views.sheetView[0].showGridLines = True

ws1['A1'] = "H. AYUNTAMIENTO DE MEDELLÍN DE BRAVO, VERACRUZ"
ws1['A1'].font = font_title
ws1['A2'] = "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS"
ws1['A2'].font = Font(name='Arial', size=11, bold=True, color=RED_FIRE_DARK)
ws1['A3'] = "DASHBOARD EJECUTIVO ANUAL DE COMBATE DE INCENDIOS Y SINIESTROS (ENERO - SEPTIEMBRE 2026)"
ws1['A3'].font = font_subtitle

tot_sin = len(df_anual)
cnt_past = len(df_anual[df_anual['Tipo_Siniestro'] == 'Incendio de Pastizal / Maleza'])
cnt_casa = len(df_anual[df_anual['Tipo_Siniestro'] == 'Incendio en Casa Habitación / Vivienda'])
cnt_gas = len(df_anual[df_anual['Tipo_Siniestro'] == 'Fuga de Gas LP / Manejo de Tanque'])
cnt_veh_com = len(df_anual[df_anual['Tipo_Siniestro'].isin(['Incendio de Vehículo / Automotriz', 'Incendio en Comercio / Taller / Bodega'])])

turno_counts = df_anual['Turno_Operativo'].value_counts()
top_turno = turno_counts.index[0] if len(turno_counts) > 0 else "Vespertino"
top_turno_cnt = turno_counts.iloc[0] if len(turno_counts) > 0 else 0
top_turno_nom = top_turno.split()[0].upper()
top_turno_pct = (top_turno_cnt / tot_sin) * 100 if tot_sin > 0 else 0

# 6 Tarjetas KPI (Filas 5 a 7)
crear_kpi_card(ws1, 5, 1, 7, 2, "TOTAL DE SINIESTROS", str(tot_sin), "Período Ene - Sep 2026", "FEF2F2", "7F1D1D")
crear_kpi_card(ws1, 5, 3, 7, 4, "INCENDIOS PASTIZAL", str(cnt_past), f"{(cnt_past/tot_sin)*100:.1f}% Mayor riesgo estiaje", "FEF3C7", "92400E")
crear_kpi_card(ws1, 5, 5, 7, 6, "CASA HABITACIÓN", str(cnt_casa), f"{(cnt_casa/tot_sin)*100:.1f}% Siniestros en vivienda", "FEE2E2", "B91C1C")
crear_kpi_card(ws1, 5, 7, 7, 8, "FUGAS DE GAS LP", str(cnt_gas), f"{(cnt_gas/tot_sin)*100:.1f}% Riesgo químico/cilindros", "EFF6FF", "1E40AF")
crear_kpi_card(ws1, 5, 9, 7, 10, "VEHÍCULOS Y COMERCIOS", str(cnt_veh_com), f"{(cnt_veh_com/tot_sin)*100:.1f}% Impacto patrimonial", "FFEDD5", "C2410C")
crear_kpi_card(ws1, 5, 11, 7, 12, "TURNO CON MÁS ACTIVIDAD", top_turno_nom, f"{top_turno_cnt} salidas ({top_turno_pct:.1f}%)", "F3E8FF", "6B21A8")

ws1.row_dimensions[5].height = 16
ws1.row_dimensions[6].height = 24
ws1.row_dimensions[7].height = 14

# 1. TABLA DE EVOLUCIÓN MENSUAL (Filas 9 a 20)
ws1['A9'] = "1. EVOLUCIÓN MENSUAL DE INCENDIOS Y SINIESTROS (ENERO A SEPTIEMBRE 2026)"
ws1['A9'].font = font_sec_header

cols_t1 = [("Mes de Registro", 18), ("Incendios Pastizal", 16), ("Otros Siniestros y Fugas", 18), ("Total Siniestros", 15)]
for idx, (nom, w) in enumerate(cols_t1, 1):
    c = ws1.cell(10, idx, nom)
    c.fill = fill_table_hdr
    c.font = font_table_hdr
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = thin_border
    ws1.column_dimensions[get_column_letter(idx)].width = w

nombres_meses = [
    ('2026-01', 'Enero 2026'),
    ('2026-02', 'Febrero 2026'),
    ('2026-03', 'Marzo 2026'),
    ('2026-04', 'Abril 2026'),
    ('2026-05', 'Mayo 2026'),
    ('2026-06', 'Junio 2026'),
    ('2026-07', 'Julio 2026'),
    ('2026-08', 'Agosto 2026'),
    ('2026-09', 'Septiembre 2026')
]

r_t1 = 11
for ym, nom_mes in nombres_meses:
    df_m = df_anual[df_anual['Fecha'].astype(str).str.startswith(ym)]
    tot_m = len(df_m)
    past_m = len(df_m[df_m['Tipo_Siniestro'] == 'Incendio de Pastizal / Maleza'])
    otros_m = tot_m - past_m
    
    ws1.cell(r_t1, 1, nom_mes).alignment = Alignment(horizontal='left')
    ws1.cell(r_t1, 2, past_m).alignment = Alignment(horizontal='center')
    ws1.cell(r_t1, 3, otros_m).alignment = Alignment(horizontal='center')
    ws1.cell(r_t1, 4, tot_m).alignment = Alignment(horizontal='center')
    for c_idx in range(1, 5):
        cell_obj = ws1.cell(r_t1, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_t1 % 2 == 0 else fill_zebra
    r_t1 += 1

# Fila Total Anual
c1 = ws1.cell(r_t1, 1, "TOTAL ANUAL CONSOLIDADO"); c1.font = font_data_bold; c1.alignment = Alignment(horizontal='left'); c1.border = thin_border
c2 = ws1.cell(r_t1, 2, cnt_past); c2.font = font_data_bold; c2.alignment = Alignment(horizontal='center'); c2.border = thin_border
c3 = ws1.cell(r_t1, 3, tot_sin - cnt_past); c3.font = font_data_bold; c3.alignment = Alignment(horizontal='center'); c3.border = thin_border
c4 = ws1.cell(r_t1, 4, tot_sin); c4.font = font_data_bold; c4.alignment = Alignment(horizontal='center'); c4.border = thin_border
for c_idx in range(1, 5): ws1.cell(r_t1, c_idx).fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

# Gráfica 1: Evolución Mensual
chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "Evolución Mensual: Incendios de Pastizal vs Otros Siniestros (2026)"
chart1.y_axis.title = "Número de Siniestros"
chart1.x_axis.title = "Mes del Año 2026"
chart1.width = 18
chart1.height = 10
chart1.dataLabels = DataLabelList()
chart1.dataLabels.showVal = True
data1 = Reference(ws1, min_col=2, min_row=10, max_col=3, max_row=19)
cats1 = Reference(ws1, min_col=1, min_row=11, max_row=19)
chart1.add_data(data1, titles_from_data=True)
chart1.set_categories(cats1)
ws1.add_chart(chart1, "F9")

# 2. MATRIZ MENSUAL POR UNIDAD MÓVIL (Fila 23 en adelante)
ws1['A23'] = "2. MATRIZ MENSUAL POR UNIDAD MÓVIL DE BOMBEROS (ENERO - SEPTIEMBRE 2026)"
ws1['A23'].font = font_sec_header

cols_matriz = [
    ("Unidad de Bomberos", 28), ("Ene", 8), ("Feb", 8), ("Mar", 8), ("Abr", 8),
    ("May", 8), ("Jun", 8), ("Jul", 8), ("Ago", 8), ("Sep", 8), ("Total 2026", 12), ("% Part.", 10)
]
for idx, (nom, w) in enumerate(cols_matriz, 1):
    c = ws1.cell(24, idx, nom)
    c.fill = fill_table_hdr
    c.font = font_table_hdr
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = thin_border
    if idx > 4:
        ws1.column_dimensions[get_column_letter(idx)].width = w

r_mat = 25
unidades_lista = [
    ("Unidad 072 (Bomberos)", "072"),
    ("Unidad 096 (Rescate / Apoyo)", "096"),
    ("Unidad 073 (Cisterna)", "073"),
    ("Pipa de Bomberos (20k Lts)", "Pipa"),
    ("Unidad 208 (Ambulancia)", "208"),
    ("Unidades 097 / 098 (Apoyo)", "097"),
    ("Unidad Motorizada PC", "Motorizada")
]

for u_nom, u_tag in unidades_lista:
    df_u = df_anual[df_anual['Unidad_Bomberos'].str.contains(u_tag, case=False, na=False)]
    u_tot = len(df_u)
    ws1.cell(r_mat, 1, u_nom).alignment = Alignment(horizontal='left')
    for m_idx, (ym, _) in enumerate(nombres_meses, 2):
        cant_m = len(df_u[df_u['Fecha'].astype(str).str.startswith(ym)])
        c_m = ws1.cell(r_mat, m_idx, cant_m)
        c_m.alignment = Alignment(horizontal='center')
    c_tot = ws1.cell(r_mat, 11, u_tot)
    c_tot.alignment = Alignment(horizontal='center')
    c_tot.font = font_data_bold
    c_pct = ws1.cell(r_mat, 12, f"{(u_tot/tot_sin)*100:.1f}%")
    c_pct.alignment = Alignment(horizontal='center')
    
    for c_idx in range(1, 13):
        cell_obj = ws1.cell(r_mat, c_idx)
        cell_obj.font = font_data if c_idx != 11 else font_data_bold
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_mat % 2 == 0 else fill_zebra
    r_mat += 1

# Fila Total Matriz
c_tot_lbl = ws1.cell(r_mat, 1, "TOTAL DESPACHOS CONSOLIDADOS")
c_tot_lbl.font = font_data_bold
c_tot_lbl.alignment = Alignment(horizontal='left')
for m_idx, (ym, _) in enumerate(nombres_meses, 2):
    tot_col = len(df_anual[df_anual['Fecha'].astype(str).str.startswith(ym)])
    c_col = ws1.cell(r_mat, m_idx, tot_col)
    c_col.alignment = Alignment(horizontal='center')
    c_col.font = font_data_bold
c_gran_tot = ws1.cell(r_mat, 11, tot_sin)
c_gran_tot.alignment = Alignment(horizontal='center')
c_gran_tot.font = font_data_bold
c_100 = ws1.cell(r_mat, 12, "100.0%")
c_100.alignment = Alignment(horizontal='center')
c_100.font = font_data_bold
for c_idx in range(1, 13):
    ws1.cell(r_mat, c_idx).border = thin_border
    ws1.cell(r_mat, c_idx).fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

# Gráfica 2: Barras Horizontales de Unidades
chart2 = BarChart()
chart2.type = "bar"
chart2.style = 13
chart2.title = "Servicios Totales por Unidad Móvil de Bomberos (2026)"
chart2.width = 17
chart2.height = 10
chart2.y_axis.title = "Unidad Móvil de Bomberos"
chart2.x_axis.title = "Total de Siniestros Atendidos"
chart2.legend = None
chart2.dataLabels = DataLabelList()
chart2.dataLabels.showVal = True
data2 = Reference(ws1, min_col=11, min_row=24, max_row=r_mat-1)
cats2 = Reference(ws1, min_col=1, min_row=25, max_row=r_mat-1)
chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)
ws1.add_chart(chart2, "A34")

# Gráfica 3: Participación Porcentual
chart3 = DoughnutChart()
chart3.title = "Participación Operativa por Unidad de Bomberos (%)"
chart3.width = 15
chart3.height = 10
chart3.dataLabels = DataLabelList()
chart3.dataLabels.showPercent = True
chart3.dataLabels.showVal = False
data3 = Reference(ws1, min_col=11, min_row=24, max_row=r_mat-1)
cats3 = Reference(ws1, min_col=1, min_row=25, max_row=r_mat-1)
chart3.add_data(data3, titles_from_data=True)
chart3.set_categories(cats3)
ws1.add_chart(chart3, "H34")

# ======================================================================
# HOJA 2: 🔥 TIPOLOGÍA Y RECURSOS
# ======================================================================
ws2 = wb.create_sheet(title="🔥 Tipología y Recursos")
ws2.views.sheetView[0].showGridLines = True

ws2['A1'] = "CLASIFICACIÓN DE INCENDIOS Y DESPLIEGUE DE RECURSOS (ENERO - SEPTIEMBRE 2026)"
ws2['A1'].font = font_title
ws2['A2'] = "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS DE MEDELLÍN DE BRAVO"
ws2['A2'].font = font_subtitle

# Sección 1: Clasificación de Siniestros
ws2['A4'] = "1. DISTRIBUCIÓN POR TIPO DE INCENDIO Y SINIESTRO (AÑO 2026)"
ws2['A4'].font = font_sec_header
ws2.cell(5, 1, "Tipo de Siniestro").fill = fill_table_hdr
ws2.cell(5, 1).font = font_table_hdr
ws2.cell(5, 2, "Casos").fill = fill_table_hdr
ws2.cell(5, 2).font = font_table_hdr
ws2.cell(5, 3, "%").fill = fill_table_hdr
ws2.cell(5, 3).font = font_table_hdr
for c_idx in range(1, 4): ws2.cell(5, c_idx).border = thin_border
ws2.column_dimensions['A'].width = 38
ws2.column_dimensions['B'].width = 14
ws2.column_dimensions['C'].width = 12

r_p1 = 6
for cat_nom, cant in df_anual['Tipo_Siniestro'].value_counts().items():
    ws2.cell(r_p1, 1, cat_nom).alignment = Alignment(horizontal='left')
    ws2.cell(r_p1, 2, cant).alignment = Alignment(horizontal='center')
    ws2.cell(r_p1, 3, f"{(cant/tot_sin)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 4):
        cell_obj = ws2.cell(r_p1, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_p1 % 2 == 0 else fill_zebra
    r_p1 += 1

chart4 = BarChart()
chart4.type = "bar"
chart4.style = 10
chart4.title = "Siniestros Atendidos por Clasificación (Ene - Sep 2026)"
chart4.width = 17
chart4.height = 10
chart4.legend = None
chart4.dataLabels = DataLabelList()
chart4.dataLabels.showVal = True
chart4.x_axis.title = "Número de Servicios"
chart4.y_axis.title = "Tipo de Incendio / Siniestro"
data4 = Reference(ws2, min_col=2, min_row=5, max_row=r_p1-1)
cats4 = Reference(ws2, min_col=1, min_row=6, max_row=r_p1-1)
chart4.add_data(data4, titles_from_data=True)
chart4.set_categories(cats4)
ws2.add_chart(chart4, "E4")

# Sección 2: Despliegue Hídrico y Métodos (Fila 17)
ws2['A17'] = "2. DESPLIEGUE HÍDRICO Y MÉTODOS DE EXTINCIÓN (AÑO 2026)"
ws2['A17'].font = font_sec_header
ws2.cell(18, 1, "Recurso / Método de Combate").fill = fill_table_hdr
ws2.cell(18, 1).font = font_table_hdr
ws2.cell(18, 2, "Servicios").fill = fill_table_hdr
ws2.cell(18, 2).font = font_table_hdr
ws2.cell(18, 3, "%").fill = fill_table_hdr
ws2.cell(18, 3).font = font_table_hdr
for c_idx in range(1, 4): ws2.cell(18, c_idx).border = thin_border

r_p2 = 19
for rec_nom, cant in df_anual['Recursos_Despliegue'].value_counts().items():
    ws2.cell(r_p2, 1, rec_nom).alignment = Alignment(horizontal='left')
    ws2.cell(r_p2, 2, cant).alignment = Alignment(horizontal='center')
    ws2.cell(r_p2, 3, f"{(cant/tot_sin)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 4):
        cell_obj = ws2.cell(r_p2, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_p2 % 2 == 0 else fill_zebra
    r_p2 += 1

chart5 = PieChart()
chart5.title = "Distribución de Despliegue Hídrico y Recursos (2026)"
chart5.width = 15
chart5.height = 9.5
chart5.dataLabels = DataLabelList()
chart5.dataLabels.showPercent = True
chart5.dataLabels.showVal = False
data5 = Reference(ws2, min_col=2, min_row=18, max_row=r_p2-1)
cats5 = Reference(ws2, min_col=1, min_row=19, max_row=r_p2-1)
chart5.add_data(data5, titles_from_data=True)
chart5.set_categories(cats5)
ws2.add_chart(chart5, "E17")

# Sección 3: Evaluación de Daños y Personas (Fila 29)
ws2['A29'] = "3. EVALUACIÓN DE DAÑOS Y AFECTACIÓN A PERSONAS (AÑO 2026)"
ws2['A29'].font = font_sec_header
ws2.cell(30, 1, "Afectación / Saldo").fill = fill_table_hdr
ws2.cell(30, 1).font = font_table_hdr
ws2.cell(30, 2, "Incidentes").fill = fill_table_hdr
ws2.cell(30, 2).font = font_table_hdr
ws2.cell(30, 3, "%").fill = fill_table_hdr
ws2.cell(30, 3).font = font_table_hdr
for c_idx in range(1, 4): ws2.cell(30, c_idx).border = thin_border

r_p3 = 31
for af_nom, cant in df_anual['Afectacion_Personas'].value_counts().items():
    ws2.cell(r_p3, 1, af_nom).alignment = Alignment(horizontal='left')
    ws2.cell(r_p3, 2, cant).alignment = Alignment(horizontal='center')
    ws2.cell(r_p3, 3, f"{(cant/tot_sin)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 4):
        cell_obj = ws2.cell(r_p3, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_p3 % 2 == 0 else fill_zebra
    r_p3 += 1

chart6 = BarChart()
chart6.type = "col"
chart6.style = 10
chart6.title = "Evaluación de Saldo y Afectación en Siniestros (2026)"
chart6.width = 16
chart6.height = 9
chart6.legend = None
chart6.dataLabels = DataLabelList()
chart6.dataLabels.showVal = True
chart6.x_axis.title = "Tipo de Afectación"
chart6.y_axis.title = "Número de Incidentes"
data6 = Reference(ws2, min_col=2, min_row=30, max_row=r_p3-1)
cats6 = Reference(ws2, min_col=1, min_row=31, max_row=r_p3-1)
chart6.add_data(data6, titles_from_data=True)
chart6.set_categories(cats6)
ws2.add_chart(chart6, "E29")

# ======================================================================
# HOJA 3: 📍 COBERTURA Y TURNOS
# ======================================================================
ws3 = wb.create_sheet(title="📍 Cobertura y Turnos")
ws3.views.sheetView[0].showGridLines = True

ws3['A1'] = "COBERTURA TERRITORIAL Y OPERATIVIDAD POR TURNOS DE BOMBEROS (2026)"
ws3['A1'].font = font_title
ws3['A2'] = "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS DE MEDELLÍN DE BRAVO"
ws3['A2'].font = font_subtitle

# Sección 1: Top 10 Zonas de Mayor Incidencia
ws3['A4'] = "1. TOP 10 ZONAS CON MAYOR INCIDENCIA DE INCENDIOS Y SINIESTROS (2026)"
ws3['A4'].font = font_sec_header
ws3.cell(5, 1, "Zona / Localidad / Fraccionamiento").fill = fill_table_hdr
ws3.cell(5, 1).font = font_table_hdr
ws3.cell(5, 2, "Siniestros").fill = fill_table_hdr
ws3.cell(5, 2).font = font_table_hdr
ws3.cell(5, 3, "%").fill = fill_table_hdr
ws3.cell(5, 3).font = font_table_hdr
for c_idx in range(1, 4): ws3.cell(5, c_idx).border = thin_border
ws3.column_dimensions['A'].width = 38
ws3.column_dimensions['B'].width = 14
ws3.column_dimensions['C'].width = 12

top10_zonas = df_anual['Ubicacion_Localidad'].value_counts().head(10)
r_g1 = 6
for z_nom, cant in top10_zonas.items():
    ws3.cell(r_g1, 1, z_nom).alignment = Alignment(horizontal='left')
    ws3.cell(r_g1, 2, cant).alignment = Alignment(horizontal='center')
    ws3.cell(r_g1, 3, f"{(cant/tot_sin)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 4):
        cell_obj = ws3.cell(r_g1, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_g1 % 2 == 0 else fill_zebra
    r_g1 += 1

chart7 = BarChart()
chart7.type = "bar"
chart7.style = 10
chart7.title = "Top 10 Zonas con Mayor Incidencia de Fuego (2026)"
chart7.width = 17
chart7.height = 11
chart7.legend = None
chart7.dataLabels = DataLabelList()
chart7.dataLabels.showVal = True
chart7.x_axis.title = "Siniestros Atendidos"
chart7.y_axis.title = "Zona / Fraccionamiento"
data7 = Reference(ws3, min_col=2, min_row=5, max_row=r_g1-1)
cats7 = Reference(ws3, min_col=1, min_row=6, max_row=r_g1-1)
chart7.add_data(data7, titles_from_data=True)
chart7.set_categories(cats7)
ws3.add_chart(chart7, "E4")

# Sección 2: Carga Horaria por Turnos (Fila 19)
ws3['A19'] = "2. CARGA OPERATIVA POR TURNO DE GUARDIA (ENERO - SEPTIEMBRE 2026)"
ws3['A19'].font = font_sec_header
ws3.cell(20, 1, "Turno de Guardia").fill = fill_table_hdr
ws3.cell(20, 1).font = font_table_hdr
ws3.cell(20, 2, "Horario Operativo").fill = fill_table_hdr
ws3.cell(20, 2).font = font_table_hdr
ws3.cell(20, 3, "Siniestros").fill = fill_table_hdr
ws3.cell(20, 3).font = font_table_hdr
ws3.cell(20, 4, "%").fill = fill_table_hdr
ws3.cell(20, 4).font = font_table_hdr
for c_idx in range(1, 5): ws3.cell(20, c_idx).border = thin_border

turnos_meta = [
    ('Matutino (07:00 - 14:59)', '07:00 a 14:59 hrs'),
    ('Vespertino (15:00 - 22:59)', '15:00 a 22:59 hrs'),
    ('Nocturno (23:00 - 06:59)', '23:00 a 06:59 hrs')
]

r_g2 = 21
for t_nom, t_horario in turnos_meta:
    cant = len(df_anual[df_anual['Turno_Operativo'] == t_nom])
    ws3.cell(r_g2, 1, t_nom.split(' (')[0]).alignment = Alignment(horizontal='left')
    ws3.cell(r_g2, 2, t_horario).alignment = Alignment(horizontal='center')
    ws3.cell(r_g2, 3, cant).alignment = Alignment(horizontal='center')
    ws3.cell(r_g2, 4, f"{(cant/tot_sin)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 5):
        cell_obj = ws3.cell(r_g2, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_g2 % 2 == 0 else fill_zebra
    r_g2 += 1

chart8 = DoughnutChart()
chart8.title = "Carga de Servicios por Turno de Bomberos (Ene - Sep 2026)"
chart8.width = 14
chart8.height = 9
chart8.dataLabels = DataLabelList()
chart8.dataLabels.showPercent = True
chart8.dataLabels.showVal = False
data8 = Reference(ws3, min_col=3, min_row=20, max_row=r_g2-1)
cats8 = Reference(ws3, min_col=1, min_row=21, max_row=r_g2-1)
chart8.add_data(data8, titles_from_data=True)
chart8.set_categories(cats8)
ws3.add_chart(chart8, "E19")

# ======================================================================
# HOJA 4: 🚒 PARQUE VEHICULAR Y COORDINACIÓN
# ======================================================================
ws4 = wb.create_sheet(title="🚒 Unidades y Coordinación")
ws4.views.sheetView[0].showGridLines = True

ws4['A1'] = "DESEMPEÑO DEL PARQUE MÓVIL Y ARTICULACIÓN INSTITUCIONAL (ENERO - SEPTIEMBRE 2026)"
ws4['A1'].font = font_title
ws4['A2'] = "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS DE MEDELLÍN DE BRAVO"
ws4['A2'].font = font_subtitle

# Sección 1: Desempeño del Parque Vehicular
ws4['A4'] = "1. DESEMPEÑO Y DESPACHOS DEL PARQUE MÓVIL DE BOMBEROS (2026)"
ws4['A4'].font = font_sec_header
ws4.cell(5, 1, "Unidad Móvil").fill = fill_table_hdr
ws4.cell(5, 1).font = font_table_hdr
ws4.cell(5, 2, "Función Operativa Principal").fill = fill_table_hdr
ws4.cell(5, 2).font = font_table_hdr
ws4.cell(5, 3, "Despachos").fill = fill_table_hdr
ws4.cell(5, 3).font = font_table_hdr
ws4.cell(5, 4, "%").fill = fill_table_hdr
ws4.cell(5, 4).font = font_table_hdr
for c_idx in range(1, 5): ws4.cell(5, c_idx).border = thin_border
ws4.column_dimensions['A'].width = 30
ws4.column_dimensions['B'].width = 36
ws4.column_dimensions['C'].width = 14
ws4.column_dimensions['D'].width = 12

parque_meta = [
    ("Unidad 072 (Bomberos)", "Ataque Rápido / Primera Respuesta", "072"),
    ("Unidad 096 (Rescate / Apoyo)", "Rescate, Logística y Fugas de Gas LP", "096"),
    ("Unidad 073 (Cisterna)", "Abastecimiento Hídrico y Bomba", "073"),
    ("Pipa de Bomberos (20k Lts)", "Cisterna Pesada de 20,000 Litros", "Pipa"),
    ("Unidades Médicas (208/097/098)", "Cobertura Preventiva y Atención a Crisis", "Ambulancia"),
    ("Unidad Motorizada PC", "Inspección Rápida de Campo y Conatos", "Motorizada")
]

r_v1 = 6
for u_nom, u_func, u_tag in parque_meta:
    cant = len(df_anual[df_anual['Unidad_Bomberos'].str.contains(u_tag, case=False, na=False)])
    ws4.cell(r_v1, 1, u_nom).alignment = Alignment(horizontal='left')
    ws4.cell(r_v1, 2, u_func).alignment = Alignment(horizontal='left')
    ws4.cell(r_v1, 3, cant).alignment = Alignment(horizontal='center')
    ws4.cell(r_v1, 4, f"{(cant/tot_sin)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 5):
        cell_obj = ws4.cell(r_v1, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_v1 % 2 == 0 else fill_zebra
    r_v1 += 1

chart9 = BarChart()
chart9.type = "col"
chart9.style = 10
chart9.title = "Despachos Operativos por Unidad Móvil de Bomberos (2026)"
chart9.width = 16
chart9.height = 9.5
chart9.legend = None
chart9.dataLabels = DataLabelList()
chart9.dataLabels.showVal = True
chart9.x_axis.title = "Unidad Móvil"
chart9.y_axis.title = "Despachos Operativos"
data9 = Reference(ws4, min_col=3, min_row=5, max_row=r_v1-1)
cats9 = Reference(ws4, min_col=1, min_row=6, max_row=r_v1-1)
chart9.add_data(data9, titles_from_data=True)
chart9.set_categories(cats9)
ws4.add_chart(chart9, "E4")

# Sección 2: Coordinación Interinstitucional (Fila 16)
ws4['A16'] = "2. COORDINACIÓN INTERINSTITUCIONAL Y AUXILIO MUTUO (2026)"
ws4['A16'].font = font_sec_header
ws4.cell(17, 1, "Instancia de Coordinación").fill = fill_table_hdr
ws4.cell(17, 1).font = font_table_hdr
ws4.cell(17, 2, "Servicios").fill = fill_table_hdr
ws4.cell(17, 2).font = font_table_hdr
ws4.cell(17, 3, "%").fill = fill_table_hdr
ws4.cell(17, 3).font = font_table_hdr
for c_idx in range(1, 4): ws4.cell(17, c_idx).border = thin_border

r_v2 = 18
coord_counts = df_anual['Coordinacion_Externa'].value_counts()
for c_nom, cant in coord_counts.items():
    ws4.cell(r_v2, 1, c_nom).alignment = Alignment(horizontal='left')
    ws4.cell(r_v2, 2, cant).alignment = Alignment(horizontal='center')
    ws4.cell(r_v2, 3, f"{(cant/tot_sin)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 4):
        cell_obj = ws4.cell(r_v2, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_v2 % 2 == 0 else fill_zebra
    r_v2 += 1

chart10 = PieChart()
chart10.title = "Articulación Interinstitucional y Auxilio Mutuo (2026)"
chart10.width = 15
chart10.height = 9.5
chart10.dataLabels = DataLabelList()
chart10.dataLabels.showPercent = True
chart10.dataLabels.showVal = False
data10 = Reference(ws4, min_col=2, min_row=17, max_row=r_v2-1)
cats10 = Reference(ws4, min_col=1, min_row=18, max_row=r_v2-1)
chart10.add_data(data10, titles_from_data=True)
chart10.set_categories(cats10)
ws4.add_chart(chart10, "E16")

# Sección 3: Origen del Reporte / Despacho (Fila 28)
ws4['A28'] = "3. CANAL DE ACTIVACIÓN Y ORIGEN DEL REPORTE (2026)"
ws4['A28'].font = font_sec_header
ws4.cell(29, 1, "Vía de Notificación").fill = fill_table_hdr
ws4.cell(29, 1).font = font_table_hdr
ws4.cell(29, 2, "Reportes").fill = fill_table_hdr
ws4.cell(29, 2).font = font_table_hdr
ws4.cell(29, 3, "%").fill = fill_table_hdr
ws4.cell(29, 3).font = font_table_hdr
for c_idx in range(1, 4): ws4.cell(29, c_idx).border = thin_border

r_v3 = 30
desp_counts = df_anual['Origen_Despacho'].value_counts()
for d_nom, cant in desp_counts.items():
    ws4.cell(r_v3, 1, d_nom).alignment = Alignment(horizontal='left')
    ws4.cell(r_v3, 2, cant).alignment = Alignment(horizontal='center')
    ws4.cell(r_v3, 3, f"{(cant/tot_sin)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 4):
        cell_obj = ws4.cell(r_v3, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_v3 % 2 == 0 else fill_zebra
    r_v3 += 1

chart11 = BarChart()
chart11.type = "bar"
chart11.style = 13
chart11.title = "Canales de Activación para el Cuerpo de Bomberos (2026)"
chart11.width = 16
chart11.height = 8.5
chart11.legend = None
chart11.dataLabels = DataLabelList()
chart11.dataLabels.showVal = True
chart11.x_axis.title = "Número de Reportes"
chart11.y_axis.title = "Canal de Activación"
data11 = Reference(ws4, min_col=2, min_row=29, max_row=r_v3-1)
cats11 = Reference(ws4, min_col=1, min_row=30, max_row=r_v3-1)
chart11.add_data(data11, titles_from_data=True)
chart11.set_categories(cats11)
ws4.add_chart(chart11, "E28")

# ======================================================================
# HOJA 5: 📋 CATÁLOGO MAESTRO DE INCENDIOS 2026
# ======================================================================
ws5 = wb.create_sheet(title="📋 Catálogo Maestro 2026")
ws5.views.sheetView[0].showGridLines = True

ws5.merge_cells('A1:O1')
ws5['A1'] = "H. AYUNTAMIENTO DE MEDELLÍN DE BRAVO, VERACRUZ"
ws5['A1'].font = Font(name='Arial', size=14, bold=True, color=NAVY_TITLE)
ws5['A1'].alignment = Alignment(horizontal='center', vertical='center')

ws5.merge_cells('A2:O2')
ws5['A2'] = "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS"
ws5['A2'].font = Font(name='Arial', size=11, bold=True, color=RED_FIRE_DARK)
ws5['A2'].alignment = Alignment(horizontal='center', vertical='center')

ws5.merge_cells('A3:O3')
ws5['A3'] = "CATÁLOGO MAESTRO Y CONTROL OFICIAL DE SALIDAS Y COMBATE DE INCENDIOS 2026"
ws5['A3'].font = font_subtitle
ws5['A3'].alignment = Alignment(horizontal='center', vertical='center')

ws5.row_dimensions[1].height = 24
ws5.row_dimensions[2].height = 18
ws5.row_dimensions[3].height = 18
ws5.row_dimensions[4].height = 8
ws5.row_dimensions[5].height = 28

columnas_maestro = [
    ("No. Folio", 14, 'center'),
    ("Fecha", 13, 'center'),
    ("Hora", 10, 'center'),
    ("Turno Operativo", 24, 'center'),
    ("Unidad de Bomberos", 26, 'center'),
    ("Personal a Cargo", 30, 'left'),
    ("Tipo de Incendio / Siniestro", 34, 'left'),
    ("Estatus / Resultado", 30, 'center'),
    ("Recursos / Despliegue Hídrico", 34, 'left'),
    ("Afectación / Víctimas", 28, 'center'),
    ("Ubicación / Localidad", 34, 'left'),
    ("Origen del Despacho", 26, 'left'),
    ("Evidencia Foto", 14, 'center'),
    ("Coordinación Externa", 32, 'left'),
    ("Resumen Operativo Oficial", 85, 'left')
]

for col_idx, (col_name, col_width, col_align) in enumerate(columnas_maestro, 1):
    cell = ws5.cell(row=5, column=col_idx)
    cell.value = col_name
    cell.fill = fill_table_hdr
    cell.font = font_table_hdr
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border
    col_letter = get_column_letter(col_idx)
    ws5.column_dimensions[col_letter].width = col_width

# Formatos condicionales para el Catálogo
fill_pastizal = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
font_pastizal = Font(name='Arial', size=9, bold=True, color="92400E")

fill_casa = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
font_casa = Font(name='Arial', size=9, bold=True, color="991B1B")

fill_gas = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
font_gas = Font(name='Arial', size=9, bold=True, color="1E40AF")

fill_sofocado = PatternFill(start_color="DEF7EC", end_color="DEF7EC", fill_type="solid")
font_sofocado = Font(name='Arial', size=9, bold=True, color="03543F")

fill_controlado = PatternFill(start_color="FEFCBF", end_color="FEFCBF", fill_type="solid")
font_controlado = Font(name='Arial', size=9, bold=True, color="975A16")

fill_falsa = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
font_falsa = Font(name='Arial', size=9, bold=True, color="475569")

for row_idx, r in df_anual.iterrows():
    current_row = row_idx + 6
    ws5.row_dimensions[current_row].height = 20
    is_even = (row_idx % 2 == 0)
    row_fill = fill_white if is_even else fill_zebra
    
    valores = [
        r['No_Folio'],
        r['Fecha'],
        r['Hora'],
        r['Turno_Operativo'],
        r['Unidad_Bomberos'],
        r['Personal_A_Cargo'],
        r['Tipo_Siniestro'],
        r['Estatus_Servicio'],
        r['Recursos_Despliegue'],
        r['Afectacion_Personas'],
        r['Ubicacion_Localidad'],
        r['Origen_Despacho'],
        r['Evidencia_Foto'],
        r['Coordinacion_Externa'],
        r['Resumen_Operativo']
    ]
    
    for col_idx, val in enumerate(valores, 1):
        cell = ws5.cell(row=current_row, column=col_idx)
        cell.value = val
        cell.font = font_data
        cell.border = thin_border
        cell.fill = row_fill
        
        _, _, col_align = columnas_maestro[col_idx - 1]
        cell.alignment = Alignment(horizontal=col_align, vertical='center', wrap_text=(col_idx == 15))
        
        # Coloreado temático de Tipo de Siniestro (Col 7)
        if col_idx == 7:
            if "Pastizal" in str(val):
                cell.fill = fill_pastizal
                cell.font = font_pastizal
            elif "Casa Habitación" in str(val):
                cell.fill = fill_casa
                cell.font = font_casa
            elif "Gas LP" in str(val):
                cell.fill = fill_gas
                cell.font = font_gas
                
        # Coloreado temático de Estatus (Col 8)
        elif col_idx == 8:
            if "Sofocado" in str(val) or "Liquidado" in str(val):
                cell.fill = fill_sofocado
                cell.font = font_sofocado
            elif "Controlado" in str(val):
                cell.fill = fill_controlado
                cell.font = font_controlado
            elif "Falsa Alarma" in str(val):
                cell.fill = fill_falsa
                cell.font = font_falsa

ws5.auto_filter.ref = f"A5:O{len(df_anual) + 5}"
ws5.freeze_panes = "A6"

# 6. GUARDADO DEL LIBRO Y MANEJO DE BLOQUEO DE ARCHIVO (LOCKS DE WINDOWS)
output_path = "c:/Users/burel/OneDrive/Documentos/PCivil Digital/Registro_Oficial_Incendios_Medellin_2026.xlsx"

try:
    wb.save(output_path)
    print(f"\n¡Éxito! Libro Ejecutivo Oficial guardado en: {output_path}")
except PermissionError:
    alt_path = "c:/Users/burel/OneDrive/Documentos/PCivil Digital/Registro_Oficial_Incendios_Medellin_2026_V2.xlsx"
    wb.save(alt_path)
    print(f"\nArchivo en uso por Excel. Guardado exitosamente en ruta alternativa: {alt_path}")
    output_path = alt_path

# Guardar réplica V2 para sincronización inmediata
replica_path = "c:/Users/burel/OneDrive/Documentos/PCivil Digital/Registro_Oficial_Incendios_Medellin_2026_V2.xlsx"
if output_path != replica_path:
    try:
        wb.save(replica_path)
        print(f"Copia réplica sincronizada en: {replica_path}")
    except Exception as e:
        print(f"Nota de réplica: {e}")

conn.close()
print("=" * 80)
print(f"PROCESO CONCLUIDO CON ÉXITO: 5 HOJAS Y 11 GRÁFICAS GENERADAS")
print(f"Total Siniestros e Incendios Auditados: {len(df_anual)}")
print("=" * 80)
