import docx

path = r'C:\Users\burel\Downloads\CEDULA DE EVALUACION SIMU MEDELLIN.docx'
doc = docx.Document(path)

print("=== TABLE 0 ===")
for r_idx, r in enumerate(doc.tables[0].rows):
    row_txt = [c.text.replace("\n", " ").strip() for c in r.cells]
    print(f"R{r_idx}: {row_txt}")

print("\n=== TABLE 1 (First 15 rows) ===")
for r_idx, r in enumerate(doc.tables[1].rows[:15]):
    row_txt = [c.text.replace("\n", " ").strip() for c in r.cells]
    print(f"R{r_idx}: {row_txt}")

print("\n=== TABLE 1 (Remaining rows) ===")
for r_idx, r in enumerate(doc.tables[1].rows[15:]):
    row_txt = [c.text.replace("\n", " ").strip() for c in r.cells]
    print(f"R{r_idx+15}: {row_txt}")
