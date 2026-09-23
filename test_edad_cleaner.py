import openpyxl
import re

wb = openpyxl.load_workbook('Catalogo_Oficial_Traslados_Medellin_2026.xlsx')
ws = wb['Catálogo General de Traslados']

def extraer_edad_valida(texto, edad_orig):
    texto_str = str(texto or '')
    edad_orig_str = str(edad_orig or '')

    # 1. Intentar limpiar edad_orig si ya tiene un número
    # Pero primero verificar que no sea texto basura
    palabras_basura = ['novedades', 'pc medellin', 'patolog', 'es del', 'dx:', 'columna', 'morena', 'salinas', 'estable', 'cenando']
    es_basura = any(b in edad_orig_str.lower() for b in palabras_basura)
    
    if not es_basura and edad_orig_str:
        m_num = re.search(r'\b(\d{1,3})\s*(?:a[ñn]os?|m(?:eses)?|d[ií]as?)?\b', edad_orig_str, re.IGNORECASE)
        if m_num:
            val = int(m_num.group(1))
            if 0 <= val <= 115:
                if 'mes' in edad_orig_str.lower(): return f"{val} meses"
                if 'd\xeda' in edad_orig_str.lower() or 'dia' in edad_orig_str.lower(): return f"{val} días"
                return f"{val} años"

    # 2. Buscar en el texto con palabra clave EDAD delimitada por palabra (\bEDAD\b)
    m_edad = re.search(r'\bEDAD\b[:\s\*]*\n?\s*(\d{1,3})\s*(?:a[ñn]os?|m(?:eses)?|d[ií]as?)?', texto_str, re.IGNORECASE)
    if m_edad:
        val = int(m_edad.group(1))
        if 0 <= val <= 115:
            return f"{val} años"

    # 3. Buscar patrones comunes tipo "XX años de edad" o "XX años"
    m_pat = re.search(r'\b(\d{1,3})\s*(?:a[ñn]os(?:\s+de\s+edad)?|de\s+edad)\b', texto_str, re.IGNORECASE)
    if m_pat:
        val = int(m_pat.group(1))
        if 0 <= val <= 115:
            return f"{val} años"

    # 4. Bebés / Pediatría en meses o días
    m_bebe = re.search(r'\b(\d{1,2})\s*(?:meses|mes|d[ií]as)\b', texto_str, re.IGNORECASE)
    if m_bebe:
        return m_bebe.group(0).strip()

    # 5. Desconocida
    if re.search(r'\b(?:edad\s+desconocida|se\s+desconoce\s+edad|desconoce\s+edad|desconocida)\b', texto_str, re.IGNORECASE):
        return "Desconocida"

    return "N/D"

print("Probando corrección de edad en las 258 filas...")
filas_corregidas = 0
for r in range(6, ws.max_row + 1):
    folio = ws.cell(row=r, column=1).value
    edad_actual = ws.cell(row=r, column=7).value
    resumen = ws.cell(row=r, column=14).value
    
    nueva_edad = extraer_edad_valida(resumen, edad_actual)
    if nueva_edad != edad_actual:
        filas_corregidas += 1
        print(f"Fila {r:3d} (Folio {folio:3d}) | Antes: {repr(str(edad_actual))} ---> Ahora: {repr(nueva_edad)}")

print(f"\nTotal filas corregidas: {filas_corregidas}")
