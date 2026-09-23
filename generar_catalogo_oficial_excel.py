import pandas as pd
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

df_orig = pd.read_excel('Catalogo_Traslados_Medellin_2026.xlsx')
print(f"Cargadas {len(df_orig)} filas del catálogo original.")

MAPA_REMITENTES = {
    '128733256110143': 'Oficialía / Guardia PC',
    '28067678416994': 'TUM Paramédico PC',
    '233453719150672': 'Paramédico Jorge Luna / U-208',
    '60314259308792': 'Paramédico Rut Sánchez / U-098',
    '244006101479643': 'Operador PC Medellín',
    '43465723384010': 'Paramédico PC Medellín',
    '170188213383234': 'Operador Unidad 096',
    '198930738487413': 'Personal Operativo PC',
    '72104464666678': 'Operador Carlos Peña / U-208',
    '214482211094536': 'Base / Despacho PC',
    '226585563119766': 'Operador PC',
    '44367045730369': 'Operador PC',
    '194171226824727': 'Operador PC',
    '65915584544854': 'Operador PC',
}

def normalizar_fecha(f_rep, f_msg):
    if not f_rep and f_msg:
        f_rep = str(f_msg).split(' ')[0]
    f_rep = str(f_rep).strip()
    
    m_iso = re.match(r'^(20\d{2})[-/](\d{1,2})[-/](\d{1,2})', f_rep)
    if m_iso:
        y, m, d = m_iso.groups()
        return f"{y}-{int(m):02d}-{int(d):02d}"
        
    m_lat = re.match(r'^(\d{1,2})[-/](\d{1,2})[-/](20\d{2}|\d{2})', f_rep)
    if m_lat:
        d, m, y = m_lat.groups()
        if len(y) == 2: y = f"20{y}"
        return f"{y}-{int(m):02d}-{int(d):02d}"
        
    if f_msg:
        return str(f_msg).split(' ')[0]
    return f_rep

def extraer_multilinea(patron_nombre, texto):
    pat = rf'\*?{patron_nombre}\*?[:\s]*\n?\s*([^\n\r\*]+)'
    m = re.search(pat, texto, re.IGNORECASE)
    if m:
        v = m.group(1).strip()
        v = re.sub(r'^\*+|\*+$', '', v).strip()
        if v and v.lower() not in ['nan', 'none', 'n/a', 'na', 'null', 'ninguno', 'no']:
            return v
    return ""

def limpiar_cadena(val):
    if pd.isna(val) or val is None:
        return ""
    v = str(val).strip()
    v = re.sub(r'^\*+|\*+$', '', v).strip()
    if v.lower() in ['nan', 'none', 'n/a', 'na', 'null', 'ninguno']:
        return ""
    return v

def extraer_edad_valida(texto_str, edad_orig_str):
    texto = str(texto_str or '')
    edad_raw = str(edad_orig_str or '').strip()

    # 1. Descartar si el valor original contiene frases o palabras basura
    palabras_basura = [
        'novedades', 'pc medellin', 'patolog', 'es del', 'dx:', 'columna',
        'morena', 'salinas', 'estable', 'cenando', 'servicio', 'general',
        'crónica', 'inconveniente', 'vidaña', 'atención', 'paciente'
    ]
    es_basura = any(b in edad_raw.lower() for b in palabras_basura)
    
    if not es_basura and edad_raw:
        m_num = re.search(r'\b(\d{1,3})\s*(?:a[ñn]os?|m(?:eses)?|d[ií]as?)?\b', edad_raw, re.IGNORECASE)
        if m_num:
            val = int(m_num.group(1))
            if 0 <= val <= 115:
                if 'mes' in edad_raw.lower(): return f"{val} meses"
                if 'd\xeda' in edad_raw.lower() or 'dia' in edad_raw.lower(): return f"{val} días"
                return f"{val} años"

    # 2. Buscar en el texto con palabra clave EDAD explícita (\bEDAD\b)
    m_edad = re.search(r'\bEDAD\b[:\s\*]*\n?\s*(\d{1,3})\s*(?:a[ñn]os?|m(?:eses)?|d[ií]as?)?', texto, re.IGNORECASE)
    if m_edad:
        val = int(m_edad.group(1))
        if 0 <= val <= 115:
            return f"{val} años"

    # 3. Buscar patrones tipo "XX años de edad" o "XX años"
    m_pat = re.search(r'\b(\d{1,3})\s*(?:a[ñn]os(?:\s+de\s+edad)?|de\s+edad)\b', texto, re.IGNORECASE)
    if m_pat:
        val = int(m_pat.group(1))
        if 0 <= val <= 115:
            return f"{val} años"

    # 4. Bebés en meses o días
    m_bebe = re.search(r'\b(\d{1,2})\s*(?:meses|mes|d[ií]as)\b', texto, re.IGNORECASE)
    if m_bebe:
        return m_bebe.group(0).strip()

    # 5. Desconocida
    if re.search(r'\b(?:edad\s+desconocida|se\s+desconoce\s+edad|desconoce\s+edad|desconocida)\b', texto, re.IGNORECASE):
        return "Desconocida"

    return "N/D"

def limpiar_nombre_paciente(nombre):
    if not nombre: return ""
    n = nombre.strip().rstrip('.')
    n_l = n.lower()
    
    palabras_invalidas = [
        'inconveniente', 'a base', 'que ', 'se traslada', 'en el ', 'para ', 'por ',
        'apoyo', 'tercera edad', 'recomienda', 'estable', 'consciente', 'orientado',
        'femenina', 'masculino', 'px', 'paciente', 'solicita', 'reporte', 'negado',
        'no ameritó', 'a domicilio'
    ]
    for p in palabras_invalidas:
        if n_l.startswith(p):
            return ""
            
    partes = n.split()
    if len(partes) < 2 and len(n) < 6:
        return ""
        
    return n.title()

def estandarizar_hospital(hosp_raw, texto_lower):
    h = (hosp_raw or '').strip().rstrip('.')
    h_l = (h + " " + texto_lower).lower()
    
    if 'se niega' in h_l or 'niega traslado' in h_l:
        return "No aplica (Negativa de Traslado)"
    if 'no amerit' in h_l:
        return "Atención en Sitio (No ameritó)"
    if 'alta' in h_l and ('domicilio' in h_l or 'atrancon' in h_l):
        return "Domicilio Particular (Alta Médica)"
    if '20 de noviembre' in h_l:
        return "Hospital General 20 de Noviembre"
    if '71' in h_l or 'díaz mirón' in h_l or 'dias miron' in h_l or 'imss 71' in h_l:
        if 'issste' in h_l: return "Clínica Hospital ISSSTE Díaz Mirón"
        return "Hospital General de Zona 71 IMSS Díaz Mirón"
    if 'issste' in h_l:
        return "Clínica Hospital ISSSTE Díaz Mirón"
    if 'boca del r' in h_l:
        return "Hospital General de Boca del Río"
    if 'haev' in h_l or 'regional' in h_l:
        return "Hospital Regional de Alta Especialidad (HAEV)"
    if 'torre pediatrica' in h_l or 'torre pediátrica' in h_l:
        return "Torre Pediátrica de Veracruz"
    if 'imss' in h_l or 'montesinos' in h_l or 'cuauhtemoc' in h_l:
        return "Hospital General de Zona IMSS"
    if 'criver' in h_l:
        return "CRIVER Veracruz"
    if 'tarimoya' in h_l:
        return "Hospital de Tarimoya"
    if 'covadonga' in h_l:
        return "Hospital Español / Covadonga"
    if h and h not in ['.', 'ARIO', 'Atención en Sitio']:
        return h.title()
    return "Atención en Sitio"

registros = []

for idx, row in df_orig.iterrows():
    texto = str(row.get('Texto_Original', ''))
    t_lower = texto.lower()
    
    # 1. FECHA Y HORA
    f_msg = str(row.get('Fecha_Mensaje', ''))
    f_rep_raw = limpiar_cadena(row.get('Fecha_Reporte')) or extraer_multilinea(r'FECHA', texto)
    f_rep = normalizar_fecha(f_rep_raw, f_msg)
    
    h_rep = limpiar_cadena(row.get('Hora')) or extraer_multilinea(r'HORA', texto)
    if not h_rep and f_msg and ' ' in f_msg:
        h_rep = f_msg.split(' ')[1]

    # 2. UNIDAD QUE ATIENDE
    amb_txt = extraer_multilinea(r'AMBULANCIA', texto) or extraer_multilinea(r'UNIDAD', texto) or limpiar_cadena(row.get('Unidad'))
    u_comb = (amb_txt + " " + t_lower).lower()
    
    if '097' in u_comb: unidad = "Unidad 097 (Ambulancia)"
    elif '098' in u_comb: unidad = "Unidad 098 (Ambulancia)"
    elif '208' in u_comb: unidad = "Unidad 208 (Ambulancia)"
    elif '096' in u_comb or 'o96' in u_comb: unidad = "Unidad 096 (Rescate/Apoyo)"
    elif '047' in u_comb: unidad = "Unidad 047"
    elif '072' in u_comb: unidad = "Unidad 072"
    elif '073' in u_comb: unidad = "Unidad 073"
    elif '041' in u_comb: unidad = "Unidad 041"
    elif 'ambulancia' in u_comb: unidad = "Ambulancia PC"
    else: unidad = "Guardia PC Medellín"

    # 3. PERSONAL QUE REALIZÓ EL TRASLADO
    operador = extraer_multilinea(r'OPERADOR', texto) or limpiar_cadena(row.get('Operador'))
    paramedico = extraer_multilinea(r'PARAM[EÉ]DICO', texto) or limpiar_cadena(row.get('Paramédico'))
    tercero = extraer_multilinea(r'TERCERO\s*ABORDO', texto) or limpiar_cadena(row.get('Tercero_Abordo'))
    
    personal_parts = []
    if paramedico: personal_parts.append(f"Paramédico: {paramedico}")
    if operador: personal_parts.append(f"Operador: {operador}")
    if tercero: personal_parts.append(f"3ro: {tercero}")
        
    rem_wa = str(row.get('Remitente_WhatsApp', ''))
    rem_alias = MAPA_REMITENTES.get(rem_wa, f"ID {rem_wa[:6]}...")
    if not personal_parts:
        personal_cargo = f"Reportado por {rem_alias}"
    else:
        personal_cargo = " | ".join(personal_parts)

    # 4. NOMBRE DEL PACIENTE
    raw_pac = extraer_multilinea(r'NOMBRE\s*DEL\s*PACIENTE', texto) or limpiar_cadena(row.get('Nombre_Paciente'))
    paciente = limpiar_nombre_paciente(raw_pac)
    
    if not paciente:
        m_p1 = re.search(r'(?:se atiende al joven\s*:?|se atiende a masculino\s*:?|se atiende a femenina\s*:?|se atiende a paciente\s*:?|px\s*:?|paciente\s*:?)\s*([A-Za-zÁÉÍÓÚáéíóúñÑ\s]+?)(?:\s+\d+\s*a[ñn]os|\s+de\s+\d+|\s+edad|\s+dx|\s+padece|\s+presenta|\s+noruega|\s+circuito|\n|\r|$)', texto, re.IGNORECASE)
        if m_p1: paciente = limpiar_nombre_paciente(m_p1.group(1))
            
        if not paciente:
            m_p2 = re.search(r'^([A-ZÁÉÍÓÚÑ][a-záéíóúñ]+(?:\s+[A-ZÁÉÍÓÚÑ][a-záéíóúñ]+){1,3}),?\s+(\d{1,2})\s*a[ñn]os', texto, re.MULTILINE)
            if m_p2: paciente = limpiar_nombre_paciente(m_p2.group(1))

    if not paciente:
        if 'persona de la tercera edad' in t_lower:
            paciente = "Persona de la tercera edad (No identificada)"
        elif 'paciente inconveniente' in t_lower:
            paciente = "Persona reportada inconveniente"
        elif 'atender reporte de ciudadanos' in t_lower:
            paciente = "Reporte ciudadano en vía pública"
        elif 'fuga de gas' in t_lower:
            paciente = "Servicio de Bomberos (Fuga de gas)"
        elif 'panal' in t_lower:
            paciente = "Servicio de Rescate (Retiro de panal)"
        elif 'gasolina' in t_lower:
            paciente = "Servicio Operativo (Carga de Gasolina)"
        elif 'apoyo de agua' in t_lower or 'riego' in t_lower:
            paciente = "Apoyo Social (Pipa de Agua / Riego)"
        else:
            paciente = "Paciente en valoración prehospitalaria"

    # 5. EDAD (CORRECCIÓN Y VALIDACIÓN TOTAL)
    edad_orig_raw = limpiar_cadena(row.get('Edad'))
    edad = extraer_edad_valida(texto, edad_orig_raw)

    # 6. DIAGNÓSTICO / MOTIVO
    diagnostico = extraer_multilinea(r'DIAGN[OÓ]STICO', texto) or extraer_multilinea(r'DESCRIPCI[OÓ]N\s*DE\s*LO\s*OCURRIDO', texto) or limpiar_cadena(row.get('Diagnóstico_Motivo'))
    if not diagnostico:
        m_dx = re.search(r'(?:dx\s*:?|diagn[oó]stico\s*:?|padecimiento\s*:?|motivo\s*:?)\s*([^\n\r\*]+)', texto, re.IGNORECASE)
        if m_dx: diagnostico = m_dx.group(1).strip()
        elif 'hipoglucemia' in t_lower: diagnostico = "Cuadro de Hipoglucemia"
        elif 'epilexia' in t_lower or 'epilepsia' in t_lower: diagnostico = "Crisis epiléptica"
        elif 'columna' in t_lower: diagnostico = "Problemas de columna / Dificultad motriz"
        elif 'disnea' in t_lower or 'dificultad respiratoria' in t_lower: diagnostico = "Dificultad respiratoria / Disnea"
        elif 'hipertens' in t_lower: diagnostico = "Crisis hipertensiva"
        elif 'policontundido' in t_lower: diagnostico = "Paciente policontundido"
        elif 'fractura' in t_lower: diagnostico = "Probable fractura"
        elif 'accidente' in t_lower or 'choque' in t_lower: diagnostico = "Accidente vehicular"
        elif 'parto' in t_lower: diagnostico = "Trabajo de parto"
        elif 'traslado programado' in t_lower: diagnostico = "Traslado programado / Cita médica"
        elif 'gasolina' in t_lower: diagnostico = "Abastecimiento de combustible"
        elif 'fuga de gas' in t_lower: diagnostico = "Fuga de gas LP"
        elif 'panal' in t_lower: diagnostico = "Retiro de enjambre de avispas/abejas"
        else: diagnostico = "Valoración médica prehospitalaria"

    # 7. HOSPITAL DE TRASLADO Y ESTATUS
    raw_hosp = extraer_multilinea(r'HOSPITAL\s*(?:DE\s*TRASLADO)?', texto) or limpiar_cadena(row.get('Hospital_Destino'))
    hospital = estandarizar_hospital(raw_hosp, t_lower)
    
    # Estatus según hospital y texto
    if "Negativa" in hospital:
        estatus = "Negativa de Traslado (Firma Deslinde)"
    elif "No ameritó" in hospital:
        estatus = "No Ameritó Traslado (Estabilizado)"
    elif "Alta Médica" in hospital:
        estatus = "Traslado de Alta a Domicilio"
    elif "Traslado Programado" in hospital or "traslado programado" in t_lower:
        estatus = "Traslado Programado"
    elif hospital not in ["Atención en Sitio", "No aplica (Negativa de Traslado)"]:
        estatus = "Traslado Efectivo a Hospital"
    else:
        estatus = "Atención en Sitio / Sin Traslado"

    # 8. MÉDICO QUE RECIBE
    recibe = extraer_multilinea(r'RECIBE', texto) or limpiar_cadena(row.get('Médico_Recibe'))
    if not recibe:
        recibe = "N/A (Atención en sitio)" if estatus != "Traslado Efectivo a Hospital" else "Personal de Guardia Hospitalaria"
    else:
        recibe = recibe.rstrip('.')

    # 9. DIRECCIÓN / UBICACIÓN
    direccion = extraer_multilinea(r'DIRECCI[OÓ]N\s*(?:DEL\s*SERVICIO)?', texto) or limpiar_cadena(row.get('Dirección'))
    if not direccion:
        if 'puente moreno' in t_lower: direccion = "Fracc. Lagos de Puente Moreno"
        elif 'el tejar' in t_lower: direccion = "Localidad El Tejar"
        elif 'playa de vacas' in t_lower: direccion = "Localidad Playa de Vacas"
        elif 'moralillo' in t_lower: direccion = "Localidad El Moralillo"
        elif 'heron proal' in t_lower or 'herón proal' in t_lower: direccion = "Colonia Herón Proal"
        elif 'arboledas' in t_lower: direccion = "Fracc. Arboledas San Ramón"
        else: direccion = "Municipio de Medellín de Bravo"

    # 10. CÓDIGO DE TRIAGE
    codigo = extraer_multilinea(r'C[OÓ]DIGO', texto) or limpiar_cadena(row.get('Código_Prioridad'))
    if not codigo:
        if 'código rojo' in t_lower or 'codigo rojo' in t_lower or 'rojo' in t_lower: codigo = "Rojo (Prioridad Alta)"
        elif 'código amarillo' in t_lower or 'codigo amarillo' in t_lower or 'amarillo' in t_lower: codigo = "Amarillo (Moderado)"
        else: codigo = "Verde (Estable)"
    else:
        c_l = codigo.lower()
        if 'rojo' in c_l: codigo = "Rojo (Prioridad Alta)"
        elif 'amarillo' in c_l: codigo = "Amarillo (Moderado)"
        else: codigo = "Verde (Estable)"

    resumen = texto.replace('\n', ' ').strip()
    resumen = re.sub(r'\s+', ' ', resumen)[:300]

    registros.append({
        'No_Folio': idx + 1,
        'Fecha': f_rep,
        'Hora': h_rep,
        'Unidad_Atiende': unidad,
        'Personal_Que_Realizo_Traslado': personal_cargo,
        'Persona_Trasladada_Paciente': paciente,
        'Edad': edad,
        'Diagnostico_Motivo': diagnostico,
        'Estatus_Traslado': estatus,
        'Hospital_Destino': hospital,
        'Medico_Recibe': recibe,
        'Direccion_Ubicacion': direccion,
        'Codigo_Triage': codigo,
        'Resumen_Operativo': resumen
    })

df_clean = pd.DataFrame(registros)
print(f"Total registros procesados: {len(df_clean)}")

output_path = "c:/Users/burel/OneDrive/Documentos/PCivil Digital/Catalogo_Oficial_Traslados_Medellin_2026.xlsx"
wb = openpyxl.Workbook()

# ==========================================
# HOJA 1: CATÁLOGO GENERAL FORMAL
# ==========================================
ws = wb.active
ws.title = "Catálogo General de Traslados"
ws.views.sheetView[0].showGridLines = True

NAVY_HEADER = "1B365D"
WHITE = "FFFFFF"
GRAY_BORDER = "D1D5DB"
ZEBRA_FILL = "F8FAFC"
TITLE_COLOR = "0F294A"

ws.merge_cells('A1:N1')
ws['A1'] = "H. AYUNTAMIENTO DE MEDELLÍN DE BRAVO, VERACRUZ"
ws['A1'].font = Font(name='Arial', size=14, bold=True, color=TITLE_COLOR)
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

ws.merge_cells('A2:N2')
ws['A2'] = "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS"
ws['A2'].font = Font(name='Arial', size=11, bold=True, color="374151")
ws['A2'].alignment = Alignment(horizontal='center', vertical='center')

ws.merge_cells('A3:N3')
ws['A3'] = "BITÁCORA Y REGISTRO OFICIAL DE TRASLADOS Y ATENCIONES DE AMBULANCIA 2026"
ws['A3'].font = Font(name='Arial', size=10, italic=True, color="4B5563")
ws['A3'].alignment = Alignment(horizontal='center', vertical='center')

ws.row_dimensions[1].height = 24
ws.row_dimensions[2].height = 18
ws.row_dimensions[3].height = 18
ws.row_dimensions[4].height = 8
ws.row_dimensions[5].height = 28

columnas = [
    ("No. Folio", 10, 'center'),
    ("Fecha", 13, 'center'),
    ("Hora", 10, 'center'),
    ("Unidad que Atiende", 24, 'center'),
    ("Personal a Cargo del Traslado", 38, 'left'),
    ("Nombre de la Persona Trasladada (Paciente)", 36, 'left'),
    ("Edad", 12, 'center'),
    ("Diagnóstico / Motivo del Traslado", 36, 'left'),
    ("Estatus del Traslado", 28, 'center'),
    ("Hospital de Destino / Receptor", 38, 'left'),
    ("Médico / Personal que Recibe", 28, 'left'),
    ("Dirección / Ubicación del Servicio", 36, 'left'),
    ("Código Triage", 20, 'center'),
    ("Resumen del Reporte Oficial", 45, 'left')
]

thin_border = Border(
    left=Side(style='thin', color=GRAY_BORDER),
    right=Side(style='thin', color=GRAY_BORDER),
    top=Side(style='thin', color=GRAY_BORDER),
    bottom=Side(style='thin', color=GRAY_BORDER)
)

header_fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
header_font = Font(name='Arial', size=10, bold=True, color=WHITE)

for col_idx, (col_name, col_width, col_align) in enumerate(columnas, 1):
    cell = ws.cell(row=5, column=col_idx)
    cell.value = col_name
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = col_width

data_font = Font(name='Arial', size=9)
zebra_pattern = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
white_pattern = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

fill_verde_traslado = PatternFill(start_color="DEF7EC", end_color="DEF7EC", fill_type="solid")
font_verde_traslado = Font(name='Arial', size=9, bold=True, color="03543F")

fill_naranja_negada = PatternFill(start_color="FEEBC8", end_color="FEEBC8", fill_type="solid")
font_naranja_negada = Font(name='Arial', size=9, bold=True, color="C05621")

fill_azul_sitio = PatternFill(start_color="EBF8FF", end_color="EBF8FF", fill_type="solid")
font_azul_sitio = Font(name='Arial', size=9, color="2B6CB0")

fill_rojo_triage = PatternFill(start_color="FED7D7", end_color="FED7D7", fill_type="solid")
font_rojo_triage = Font(name='Arial', size=9, bold=True, color="9B2C2C")

fill_amarillo_triage = PatternFill(start_color="FEFCBF", end_color="FEFCBF", fill_type="solid")
font_amarillo_triage = Font(name='Arial', size=9, bold=True, color="975A16")

fill_verde_triage = PatternFill(start_color="C6F6D5", end_color="C6F6D5", fill_type="solid")
font_verde_triage = Font(name='Arial', size=9, bold=True, color="22543D")

for row_idx, r in df_clean.iterrows():
    current_row = row_idx + 6
    ws.row_dimensions[current_row].height = 20
    is_even = (row_idx % 2 == 0)
    row_fill = white_pattern if is_even else zebra_pattern
    
    valores_fila = [
        r['No_Folio'],
        r['Fecha'],
        r['Hora'],
        r['Unidad_Atiende'],
        r['Personal_Que_Realizo_Traslado'],
        r['Persona_Trasladada_Paciente'],
        r['Edad'],
        r['Diagnostico_Motivo'],
        r['Estatus_Traslado'],
        r['Hospital_Destino'],
        r['Medico_Recibe'],
        r['Direccion_Ubicacion'],
        r['Codigo_Triage'],
        r['Resumen_Operativo']
    ]
    
    for col_idx, val in enumerate(valores_fila, 1):
        cell = ws.cell(row=current_row, column=col_idx, value=val)
        cell.font = data_font
        cell.border = thin_border
        cell.fill = row_fill
        
        align_type = columnas[col_idx - 1][2]
        cell.alignment = Alignment(horizontal=align_type, vertical='center')
        
        if col_idx == 9:
            if "Traslado Efectivo" in str(val) or "Traslado de Alta" in str(val) or "Traslado Programado" in str(val):
                cell.fill = fill_verde_traslado
                cell.font = font_verde_traslado
            elif "Negativa" in str(val):
                cell.fill = fill_naranja_negada
                cell.font = font_naranja_negada
            elif "No Ameritó" in str(val) or "Atención en Sitio" in str(val):
                cell.fill = fill_azul_sitio
                cell.font = font_azul_sitio
                
        if col_idx == 13:
            if "Rojo" in str(val):
                cell.fill = fill_rojo_triage
                cell.font = font_rojo_triage
            elif "Amarillo" in str(val):
                cell.fill = fill_amarillo_triage
                cell.font = font_amarillo_triage
            elif "Verde" in str(val):
                cell.fill = fill_verde_triage
                cell.font = font_verde_triage

ws.auto_filter.ref = f"A5:N{len(df_clean) + 5}"
ws.freeze_panes = "A6"

# ==========================================
# HOJA 2: RESUMEN EJECUTIVO Y MÉTRICAS
# ==========================================
ws2 = wb.create_sheet(title="Resumen Ejecutivo y Métricas")
ws2.views.sheetView[0].showGridLines = True

ws2['A1'] = "RESUMEN EJECUTIVO DE SERVICIOS Y TRASLADOS - PROTECCIÓN CIVIL MEDELLÍN"
ws2['A1'].font = Font(name='Arial', size=13, bold=True, color=TITLE_COLOR)

ws2['A3'] = "ESTATUS GENERAL DEL SERVICIO"
ws2['A3'].font = Font(name='Arial', size=11, bold=True, color=NAVY_HEADER)
ws2['A4'] = "Tipo de Estatus"
ws2['B4'] = "Total Servicios"
ws2['C4'] = "Porcentaje"
for c in ['A4', 'B4', 'C4']:
    ws2[c].fill = header_fill
    ws2[c].font = header_font
    ws2[c].alignment = Alignment(horizontal='center')

total_serv = len(df_clean)
r_idx = 5
for estatus, cant in df_clean['Estatus_Traslado'].value_counts().items():
    ws2.cell(row=r_idx, column=1, value=estatus).font = data_font
    ws2.cell(row=r_idx, column=2, value=cant).alignment = Alignment(horizontal='center')
    ws2.cell(row=r_idx, column=3, value=f"{(cant/total_serv)*100:.1f}%").alignment = Alignment(horizontal='center')
    r_idx += 1

ws2['E3'] = "DISTRIBUCIÓN POR UNIDAD MÓVIL"
ws2['E3'].font = Font(name='Arial', size=11, bold=True, color=NAVY_HEADER)
ws2['E4'] = "Unidad"
ws2['F4'] = "Total Servicios"
ws2['G4'] = "Porcentaje"
for c in ['E4', 'F4', 'G4']:
    ws2[c].fill = header_fill
    ws2[c].font = header_font
    ws2[c].alignment = Alignment(horizontal='center')

r_idx2 = 5
for unidad_name, cant in df_clean['Unidad_Atiende'].value_counts().items():
    ws2.cell(row=r_idx2, column=5, value=unidad_name).font = data_font
    ws2.cell(row=r_idx2, column=6, value=cant).alignment = Alignment(horizontal='center')
    ws2.cell(row=r_idx2, column=7, value=f"{(cant/total_serv)*100:.1f}%").alignment = Alignment(horizontal='center')
    r_idx2 += 1

ws2['A14'] = "PRINCIPALES HOSPITALES RECEPTORES DE TRASLADOS"
ws2['A14'].font = Font(name='Arial', size=11, bold=True, color=NAVY_HEADER)
ws2['A15'] = "Hospital / Nosocomio"
ws2['B15'] = "Pacientes Recibidos"
for c in ['A15', 'B15']:
    ws2[c].fill = header_fill
    ws2[c].font = header_font
    ws2[c].alignment = Alignment(horizontal='center')

r_idx3 = 16
hosp_counts = df_clean[df_clean['Estatus_Traslado'] == 'Traslado Efectivo a Hospital']['Hospital_Destino'].value_counts()
for h_name, cant in hosp_counts.head(10).items():
    ws2.cell(row=r_idx3, column=1, value=h_name).font = data_font
    ws2.cell(row=r_idx3, column=2, value=cant).alignment = Alignment(horizontal='center')
    r_idx3 += 1

for col in ['A', 'B', 'C', 'D', 'E', 'F', 'G']:
    ws2.column_dimensions[col].width = 32 if col in ['A', 'E'] else 18

try:
    wb.save(output_path)
    print(f"\nArchivo guardado exitosamente en:\n  {output_path}")
except PermissionError:
    alt_path = "c:/Users/burel/OneDrive/Documentos/PCivil Digital/Catalogo_Oficial_Traslados_Medellin_2026_Corregido.xlsx"
    wb.save(alt_path)
    print(f"\nEl archivo original estaba abierto en Excel. Se guardó exitosamente como versión corregida en:\n  {alt_path}")
    output_path = alt_path
