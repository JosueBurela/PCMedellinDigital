import docx

path = r'C:\Users\burel\Downloads\CEDULA DE EVALUACION SIMU MEDELLIN.docx'
doc = docx.Document(path)

print('=== SECTIONS ===')
for i, s in enumerate(doc.sections):
    print(f'Section {i}: start_type={s.start_type}, top={s.top_margin.pt}, bottom={s.bottom_margin.pt}, left={s.left_margin.pt}, right={s.right_margin.pt}, page_w={s.page_width.pt}, page_h={s.page_height.pt}')

print('\n=== PARAGRAPHS ===')
for i, p in enumerate(doc.paragraphs):
    xml = p._p.xml
    has_break = '<w:br w:type="page"' in xml or '<w:pageBreak' in xml
    has_drawing = '<w:drawing' in xml or '<v:shape' in xml
    text = p.text.strip()
    if has_break or has_drawing or text or len(xml) > 200:
        print(f'P{i}: [len_text={len(text)}] [text: {text[:60]}] [page_break={has_break}] [drawing={has_drawing}]')

print('\n=== TABLES ===')
for i, t in enumerate(doc.tables):
    print(f'Table {i}: rows={len(t.rows)}, cols={len(t.columns)}')
