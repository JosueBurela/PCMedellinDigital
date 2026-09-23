import psycopg2
import re
import os
import sys
from datetime import datetime, timedelta
from collections import defaultdict, Counter
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList

sys.stdout.reconfigure(encoding='utf-8')

print("=" * 80)
print("GENERADOR DEL EXPEDIENTE EJECUTIVO ANUAL Y CATÁLOGO MAESTRO 2026")
print("Protección Civil y Bomberos - Medellín de Bravo, Veracruz")
print("=" * 80)

conn = psycopg2.connect("postgresql://postgres:postgres@localhost:5432/evolution")
cur = conn.cursor()

GROUP_JID = '120363042493725288@g.us'

cur.execute("""
    SELECT 
        "id",
        to_timestamp("messageTimestamp") as fecha,
        "pushName",
        "messageType",
        COALESCE(
            "message"->>'conversation',
            "message"->'extendedTextMessage'->>'text',
            "message"->'imageMessage'->>'caption',
            "message"->'videoMessage'->>'caption',
            ''
        ) as texto,
        "message"->'imageMessage' IS NOT NULL as tiene_imagen,
        "key"->>'id' as wa_id,
        "key"->>'participant' as remitente
    FROM "Message"
    WHERE "key"->>'remoteJid' = %s
      AND to_timestamp("messageTimestamp") >= '2026-01-01 00:00:00'
    ORDER BY "messageTimestamp" ASC
""", (GROUP_JID,))

todos_2026 = cur.fetchall()
print(f"Total mensajes históricos 2026 recuperados de BD: {len(todos_2026)}")

MAPA_REMITENTES = {
    '128733256110143': 'Oficialía / Guardia PC',
    '28067678416994': 'TUM Paramédico PC',
    '233453719150672': 'Paramédico Jorge Luna / U-208',
    '60314259308792': 'Paramédico Rut Sánchez / U-098',
    '244006101479643': 'Operador PC Medellín',
    '43465723384010': 'Paramédico PC Medellín',
    '170188213383234': 'Operador Unidad 096',
    '198930738487413': 'Personal Operativo PC',
    '72104464666678': 'Operador Carlos Peña / U-208',
    '214482211094536': 'Base / Despacho PC',
    '177206105997437': 'Médico / Paramédico Evaluador PC',
    '140355437649994': 'TUM Paramédico de Guardia',
    '219442864750741': 'Enlace C5 / Despacho 911'
}

def normalizar_fecha(f_rep, f_msg):
    if not f_rep and f_msg:
        f_rep = str(f_msg).split(' ')[0]
    f_rep = str(f_rep).strip()
    m_iso = re.match(r'^(20\d{2})[-/](\d{1,2})[-/](\d{1,2})', f_rep)
    if m_iso:
        y, m, d = m_iso.groups()
        return f"{y}-{int(m):02d}-{int(d):02d}"
    m_lat = re.match(r'^(\d{1,2})[-/](\d{1,2})[-/](20\d{2}|\d{2})', f_rep)
    if m_lat:
        d, m, y = m_lat.groups()
        if len(y) == 2: y = f"20{y}"
        return f"{y}-{int(m):02d}-{int(d):02d}"
    if f_msg:
        return str(f_msg).split(' ')[0]
    return f_rep

def extraer_multilinea(patron_nombre, texto):
    pat = rf'\*?{patron_nombre}\*?[:\s]*\n?\s*([^\n\r\*]+)'
    m = re.search(pat, texto, re.IGNORECASE)
    if m:
        v = m.group(1).strip()
        v = re.sub(r'^\*+|\*+$', '', v).strip()
        if v and v.lower() not in ['nan', 'none', 'n/a', 'na', 'null', 'ninguno', 'no']:
            return v
    return ""

def extraer_edad_valida(texto_str):
    texto = str(texto_str or '')
    m_pat = re.search(r'\b(\d{1,3})\s*(?:a[ñn]os?(?:\s+de\s+edad)?|de\s+edad)\b', texto, re.IGNORECASE)
    if m_pat:
        val = int(m_pat.group(1))
        if 0 <= val <= 115:
            return f"{val} años"
    m_bebe = re.search(r'\b(\d{1,2})\s*(?:meses|mes|d[ií]as)\b', texto, re.IGNORECASE)
    if m_bebe:
        return m_bebe.group(0).strip()
    if re.search(r'\b(?:edad\s+desconocida|se\s+desconoce\s+edad|desconoce\s+edad|desconocida)\b', texto, re.IGNORECASE):
        return "Desconocida"
    return "N/D"

PAT_MUERTE = re.compile(
    r'\b(?:c[oó]digo\s+(?:14|negro)|ya\s+era\s+14|era\s+14|es\s+14(?!\s+de|\s+del)|px\s*14|px_14|fallecid[oa]s?|[oó]bito|sin\s+signos\s+vitales|sin\s+vida|cad[aá]ver|post\s*mortem|rigor\s*mortis|livideses|muerte\s+patol[oó]gica|muerte\s+por\s+asfixia)\b', 
    re.IGNORECASE
)

def limpiar_nombre_paciente(nombre, full_text=""):
    if not nombre:
        n = ""
    else:
        n = str(nombre).strip().rstrip('.')
    
    n = re.sub(r'^(?:de\s+nombre|nombre\s*:?|px\s*:?|paciente\s*:?|al\s+joven|a\s+la\s+paciente|se\s+atiende\s+a(?:l)?|masculino\s*:?|femenin[oa]\s*:?)\s*', '', n, flags=re.IGNORECASE).strip()

    cortes = [
        r'\bde\s+\d{1,3}\s*a[ñn]os\b',
        r'\bcon\b',
        r'\bheridas\b',
        r'\btraslad',
        r'\bfirma\b',
        r'\bsolicit',
        r'\bse\s+queda\b',
        r'\ba\s+la\s+cl[íi]nica\b',
        r'\bal\s+hospital\b',
        r'\bes\s+el\s+14\b',
        r'[\r\n]+'
    ]
    for c in cortes:
        m = re.search(c, n, flags=re.IGNORECASE)
        if m:
            n = n[:m.start()].strip()

    n = n.strip(' ".,;:-_')
    n = re.sub(r'\s+de$', '', n, flags=re.IGNORECASE).strip()

    basura_nombre = [
        'clínico', 'clinico', 'oncol', 'máximo beneficio', 'maximo beneficio',
        'inconveniente', 'laceracion', 'sin signos vitales', 'derrape', 'inconsciente',
        'desvanecid', 'cirugía', 'cirugia', 'fractura', 'traslado', 'se queda', 'se atiende',
        'aliento alcohólico', 'aliento alcoholico', 'beneficio', 'fallecid', 'choque',
        'atropellad', 'accidente', 'crítico', 'critico', 'código', 'codigo', 'star médica',
        'star medica', 'star medic', 'retorno a su domicilio', 'lo que es una localidad', 'se dirige',
        'a base', 'que ', 'en el ', 'para ', 'por ', 'apoyo', 'tercera edad', 'recomienda',
        'estable', 'consciente', 'orientado', 'conciente', 'femenina', 'masculino', 'reporte', 'negado',
        'no ameritó', 'a petición', 'de bace', 'hemodiálisis', 'hemodialisis', 'domicilio', 'en base',
        'pediatrico', 'pediátrico', 'geriátrico', 'geriatrico', 'somnoliento', 'reactivo', 'cuenta'
    ]
    if any(b in n.lower() for b in basura_nombre):
        n = ""

    if any(h in n.lower() for h in ['hospital', 'clínica', 'clinica', 'imss', 'issste', 'naval', 'pemex', 'criver']):
        n = ""

    partes = n.split()
    if len(partes) < 2:
        n = ""

    if n:
        return n.title()

    if full_text:
        pat_prefijo = r'(?:nombre\s+del\s+px(?:\s*14)?|px(?:\s*14)?\s*nombre|px\s*:|segundo\s+px|nombre\s+de\s+la\s+femenina(?:\s+y\s+sus\s+signos\s+vitales)?|nombre\s+del\s+masculino|px\s+gediatrica\s+de\s+nombre|px\s+de\s+nombre|se\s+atiende\s+a\s+femenina\s+de\s+nombre|se\s+atiende\s+a\s+masculino\s+de\s+nombre|paciente\s+de\s+nombre|de\s+nombre)\s*:?\s*([A-Za-zÁÉÍÓÚáéíóúñÑ\s]+?)(?:\s+de\s+\d|\s+\d{1,3}\s*a[ñn]os|\s+\d{2}\b|\s+con\s+los|\s+quien|\s+t/a|\s+ta:|\s+fc|\s+presentando|\s+se\s+traslada|\.|\n|\||$)'
        m = re.search(pat_prefijo, full_text, re.IGNORECASE)
        if m:
            cand = m.group(1).strip()
            cand = re.sub(r'^(?:y\s+sus\s+signos\s+vitales|de\s+la\s+femenina|del\s+px)\s*', '', cand, flags=re.IGNORECASE).strip()
            cand = re.sub(r'\s+de$', '', cand, flags=re.IGNORECASE).strip()
            pts = cand.split()
            if 2 <= len(pts) <= 5 and not any(w in cand.lower() for w in ['signos', 'vitales', 'hospital', 'clínica', 'unidad', 'reporta', 'traslado']):
                return cand.title()

        m_ini_guion = re.match(r'^([A-ZÁÉÍÓÚ][a-záéíóúñ]+(?:\s+[A-ZÁÉÍÓÚa-záéíóúñ]+){1,3})(?:\s+de\s+[A-ZÁÉÍÓÚa-z\s]+)?\s*--\s*(?:Femenina|Masculino)', full_text, re.IGNORECASE)
        if m_ini_guion:
            cand = m_ini_guion.group(1).strip()
            cand = re.sub(r'\s+de\s+Boca.*$', '', cand, flags=re.IGNORECASE)
            cand = re.sub(r'\s+de$', '', cand, flags=re.IGNORECASE).strip()
            pts = cand.split()
            if 2 <= len(pts) <= 5 and not any(w in cand.lower() for w in ['buen', 'reporte', 'novedades', 'guardia', 'unidad', 'medellin', 'salida', 'persona']):
                return cand.title()

        m_ini_edad = re.match(r'^([A-ZÁÉÍÓÚ][a-záéíóúñ]+(?:\s+[A-ZÁÉÍÓÚa-záéíóúñ]+){1,3})\s+(?:femenina|masculino)?\s*(?:de\s+\d{1,3}|\d{1,3}\s*a[ñn]os)', full_text, re.IGNORECASE)
        if m_ini_edad:
            cand = m_ini_edad.group(1).strip()
            cand = re.sub(r'\s+de$', '', cand, flags=re.IGNORECASE).strip()
            pts = cand.split()
            if 2 <= len(pts) <= 5 and not any(w in cand.lower() for w in ['buen', 'reporte', 'novedades', 'guardia', 'unidad', 'medellin', 'salida', 'persona', 'se atiende']):
                return cand.title()

    return ""

def estandarizar_hospital(hosp_raw, texto_lower, es_fallecido=False):
    if es_fallecido:
        return "No aplica (Persona Fallecida en Sitio)"
    h = (hosp_raw or '').strip().rstrip('.')
    h_l = (h + " " + texto_lower).lower()
    
    if PAT_MUERTE.search(h_l) and not any(k in h_l for k in ['traslada', 'traslado al']):
        return "No aplica (Persona Fallecida en Sitio)"
    
    if any(k in h_l for k in ['se niega', 'niega traslado', 'nego traslado', 'firma deslinde', 'firma responsiva']):
        return "No aplica (Negativa de Traslado)"
    if 'no amerit' in h_l or 'no amierita' in h_l:
        return "Atención en Sitio (No ameritó traslado)"
    if any(k in h_l for k in ['a su domicilio', 'a domicilio', 'domicilio particular', 'retorno a su domicilio', 'al domicilio del paciente']):
        return "Domicilio Particular (Alta Médica)"
        
    # Hospital General de Boca del Río
    if any(k in h_l for k in ['hg de boca', 'hg boca', 'hospital general de boca', 'hospital de boca', 'traslada al hg', 'traslado al hg', 'traslada a boca', 'traslado a boca', 'apoyo a el hg', 'al hg de boca']):
        return "Hospital General de Boca del Río"
    if 'boca del r' in h.lower() or 'hg boca' in h.lower():
        return "Hospital General de Boca del Río"

    # Red Hospitalaria General (aplica tanto si viene en campo hospital como en la bitácora de guardia/nota)
    if '71' in h_l or 'díaz mirón' in h_l or 'dias miron' in h_l or 'imss 71' in h_l:
        if 'issste' in h_l: return "Clínica Hospital ISSSTE Díaz Mirón"
        return "Hospital General de Zona 71 IMSS Díaz Mirón"
    if '61' in h_l or 'clínica 61' in h_l:
        return "Unidad de Medicina Familiar 61 IMSS"
    if '20 de noviembre' in h_l or 'general 20' in h_l:
        return "Hospital General 20 de Noviembre"
    if 'pemex' in h_l:
        return "Hospital de PEMEX"
    if 'santa lucia' in h_l or 'santa lucía' in h_l:
        return "Clínica Santa Lucía"
    if 'issste' in h_l:
        return "Clínica Hospital ISSSTE Díaz Mirón"
    if 'haev' in h_l or 'regional' in h_l or 'alta especialidad' in h_l:
        return "Hospital Regional de Alta Especialidad (HAEV)"
    if 'torre pediatrica' in h_l or 'torre pediátrica' in h_l:
        return "Torre Pediátrica de Veracruz"
    if 'naval' in h_l or 'hosnaver' in h_l:
        return "Hospital Naval de Especialidades (HOSNAVER)"
    if 'militar' in h_l:
        return "Hospital Militar de Zona Veracruz"
    if 'star medic' in h_l or 'star medica' in h_l:
        return "Hospital Star Médica Veracruz"
    if 'san francisco' in h_l:
        return "Sanatorio San Francisco"
    if 'gainza' in h_l or 'gaínza' in h_l:
        return "Clínica Gaínza"
    if 'criver' in h_l:
        return "CRIVER Veracruz"
    if 'covadonga' in h_l or 'español' in h_l:
        return "Hospital Español / Covadonga"
    if 'imss' in h_l or 'cuauhtemoc' in h_l or '14 de cuauhtemoc' in h_l:
        return "Hospital General de Zona IMSS"
        
    if h and len(h) <= 40 and not any(w in h.lower() for w in ['paciente', 'retorno', 'compañero', 'familiar', 'n/a', 'no']):
        return h.title()
    return "Atención en Sitio"

def estandarizar_ubicacion(txt):
    t_low = txt.lower()
    for loc in [
        ('lagos de puente moreno', 'Fracc. Lagos de Puente Moreno'),
        ('puente moreno', 'Fracc. Puente Moreno'),
        ('arboledas san ramón', 'Fracc. Arboledas San Ramón'),
        ('arboledas san ramon', 'Fracc. Arboledas San Ramón'),
        ('san ramón', 'Fracc. San Ramón'),
        ('san ramon', 'Fracc. San Ramón'),
        ('arboledas san miguel', 'Fracc. Arboledas San Miguel'),
        ('san miguel', 'Fracc. San Miguel'),
        ('el tejar', 'Localidad El Tejar'),
        ('tejar', 'Localidad El Tejar'),
        ('los robles', 'Localidad Los Robles'),
        ('paso del toro', 'Localidad Paso del Toro'),
        ('paso de toro', 'Localidad Paso del Toro'),
        ('rancho del padre', 'Localidad Rancho del Padre'),
        ('dos bocas', 'Localidad Dos Bocas'),
        ('playa de vacas', 'Localidad Playa de Vacas'),
        ('la laguna', 'Localidad La Laguna'),
        ('ixcualco', 'Localidad Ixcualco'),
        ('el 12', 'Localidad El 12'),
        ('la bocana', 'Localidad La Bocana'),
        ('la gloria', 'Localidad La Gloria'),
        ('moralillo', 'Localidad El Moralillo'),
        ('las palmas', 'Fracc. Las Palmas')
    ]:
        if loc[0] in t_low:
            return f"{loc[1]}, Medellín de Bravo, Ver."
    return "Medellín de Bravo, Ver."

def detectar_unidad(txt):
    t_low = txt.lower()
    if '097' in t_low or 'u-97' in t_low or 'unidad 97' in t_low or 'u097' in t_low: return "Unidad 097 (Ambulancia)"
    if '098' in t_low or 'u-98' in t_low or 'unidad 98' in t_low or 'u098' in t_low: return "Unidad 098 (Ambulancia)"
    if '208' in t_low or 'u-208' in t_low or 'u208' in t_low: return "Unidad 208 (Ambulancia)"
    if 'moto' in t_low or 'motorizada' in t_low: return "Unidad Motorizada PC"
    if '096' in t_low: return "Unidad 096 (Rescate / Apoyo)"
    if 'ambulancia' in t_low: return "Ambulancia PC Medellín"
    return "Ambulancia PC Medellín"

def clasificar_turno(hora_str):
    if not hora_str: return "Matutino (07:00 - 14:59)"
    try:
        h = int(str(hora_str).split(':')[0])
        if 7 <= h < 15: return "Matutino (07:00 - 14:59)"
        elif 15 <= h < 23: return "Vespertino (15:00 - 22:59)"
        else: return "Nocturno (23:00 - 06:59)"
    except:
        return "Matutino (07:00 - 14:59)"

def clasificar_categoria_clinica(motivo_str, full_text):
    t = (str(motivo_str) + " " + str(full_text)).lower()
    if any(k in t for k in ['moto', 'derrape', 'choque', 'choke', 'volcadura', 'atropellad', 'accidente', 'vehicular']):
        return "Accidentes Viales y Traumatismos"
    if any(k in t for k in ['iam', 'infart', 'acb', 'evc', 'cardíac', 'cardiac', 'pecho', 'hemodiálisis', 'hemodialisis', 'diabetes', 'hipoglucemia', 'hipertens', 'presión', 'presion']):
        return "Urgencias Cardiovasculares y Metabólicas"
    if any(k in t for k in ['caída', 'caida', 'escaleras', 'sustentación', 'propia altura', 'policontundid', 'fractura', 'esguince']):
        return "Caídas y Traumatismos Menores"
    if any(k in t for k in ['respirar', 'respiratori', 'disnea', 'asma', 'dengue', 'fiebre', 'intoxic', 'covid', 'pulmon']):
        return "Urgencias Respiratorias e Infecciosas"
    if any(k in t for k in ['violencia', 'golpeó', 'golpeo', 'marido', 'riña', 'arma blanca', 'agresión', 'asalt']):
        return "Violencia Familiar y Agresiones"
    if any(k in t for k in ['parto', 'embaraz', 'aborto', 'pediátr', 'pediatr', 'menor de edad', 'recién nacido']):
        return "Urgencias Gineco-Obstétricas y Pediátricas"
    return "Otras Urgencias Médicas y Valoraciones"

def clasificar_sector_salud(hosp_str):
    h = str(hosp_str).lower()
    if any(k in h for k in ['boca del río', 'boca del rio', 'haev', 'regional', 'torre pediátrica', 'torre pediatrica', 'tarimoya']):
        return "Sector SESVER (Salud Estatal)"
    if any(k in h for k in ['61', '71', 'imss', 'cuauhtemoc']):
        return "Sector IMSS (Seguro Social)"
    if any(k in h for k in ['issste', 'pemex', 'militar', 'naval', 'hosnaver']):
        return "Sector Federal y Militar"
    if any(k in h for k in ['gainza', 'gaínza', 'santa lucia', 'santa lucía', 'covadonga', 'español', 'star médic', 'star medic', 'san francisco', 'criver']):
        return "Sector Privado y Asistencial"
    if 'domicilio' in h:
        return "Traslado de Alta a Domicilio"
    return "Atención en Sitio (No Hospitalario)"

eventos = []

# MODO A: Resúmenes de guardia
for mid, dt, push, mtype, txt, tiene_img, wa_id, remitente in todos_2026:
    if not txt or len(txt) < 15: continue
    t_low = txt.lower()
    
    if any(k in t_low for k in ['novedades del d', 'reportes de servicios pc', 'novedades de pc', 'reporte de guardia']) and any(s in t_low for s in ['sale unidad', 'sale el compañero', 'salida', '1_']):
        lines = txt.split('\n')
        header = lines[0].strip()
        m_fecha = re.search(r'(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})', header)
        if m_fecha:
            d, m, y = m_fecha.groups()
            if len(y) == 2: y = f"20{y}"
            if y == "2025" and dt.year == 2026: y = "2026"
            fecha_guardia = f"{y}-{int(m):02d}-{int(d):02d}"
        else:
            fecha_guardia = dt.strftime('%Y-%m-%d')
            
        pat_inicio = re.compile(r'^\s*(?:(\d{1,2}[:;]\d{2})|(\d{1,2})[_\.\-\)])\s*(.*)', re.IGNORECASE)
        items = []
        cur_it = []
        for line in lines[1:]:
            l_str = line.strip()
            if not l_str: continue
            if pat_inicio.match(l_str):
                if cur_it:
                    items.append("\n".join(cur_it))
                    cur_it = []
                cur_it.append(l_str)
            else:
                if cur_it: cur_it.append(l_str)
        if cur_it: items.append("\n".join(cur_it))
        
        for it in items:
            it_low = it.lower()
            es_med = False
            if any(k in it_low for k in ['traslado', 'clínica', 'clinica', 'hospital', 'px', 'paciente', 'lesionad', 'herid',
                                         'atropellad', 'choque', 'choke', 'derrape', 'volcadura', 'inconsciente', 'inconciente',
                                         'combulci', 'convulsi', 'infart', 'iam', 'acb', 'sonda', 'enfermedad general',
                                         'desvanecid', 'caída', 'caida', 'embaraz', 'parto', 'dificultad para respirar',
                                         'respiratorio', 'suicidio', 'intoxicaci', 'hemodiálisis', 'hemodialisis', 'silla de ruedas']):
                es_med = True
            elif any(u in it_low for u in ['208', '097', 'u-97', 'ambulancia']) and not any(ex in it_low for ex in ['combustible', 'gasolina', 'taller', 'lavar', 'agua']):
                es_med = True
                
            if any(ex in it_low for ex in ['incendio de pastizal', 'incendio de lote', 'incendio de basura', 'fuga de gas', 'kema de basura']):
                if not any(k in it_low for k in ['lesionad', 'quemad', 'kemad', 'traslad', 'atend']):
                    es_med = False
            if any(ex in it_low for ex in ['carga de combustible', 'taller mecánico', 'rotulación', 'capturar un mapache', 'capturar una culebra']):
                es_med = False

            if es_med:
                h_sal = ""
                m_sal = re.search(r'salida\s*[:;]?\s*(\d{1,2}[:;]\d{2})', it, re.IGNORECASE)
                if m_sal: h_sal = m_sal.group(1).replace(';', ':')
                else:
                    m_ini_h = re.match(r'^\s*(\d{1,2}[:;]\d{2})\b', it)
                    if m_ini_h: h_sal = m_ini_h.group(1).replace(';', ':')
                if not h_sal:
                    all_h = re.findall(r'\b\d{1,2}[:;]\d{2}\b', it)
                    if all_h: h_sal = all_h[0].replace(';', ':')
                if not h_sal: h_sal = dt.strftime('%H:%M')
                
                parts_h = h_sal.split(':')
                if len(parts_h) == 2:
                    h_sal = f"{int(parts_h[0]):02d}:{parts_h[1]}"

                u = detectar_unidad(it)
                ub = estandarizar_ubicacion(it)
                es_muerte = bool(PAT_MUERTE.search(it))

                # Caso especial Sal Si Puedes (accidente con 1 lesionado trasladado y 1 fallecido en sitio)
                if 'sal si puedes' in it_low and 'traslada un masculino' in it_low and 'otro ya estaba fallecido' in it_low:
                    # Paciente 1: Trasladado
                    eventos.append({
                        'origen': 'resumen_guardia',
                        'fecha': fecha_guardia,
                        'hora': h_sal,
                        'unidad': u,
                        'paciente': "Masculino lesionado",
                        'edad': "N/D",
                        'sexo': "Masculino",
                        'diagnostico': "Politraumatizado en accidente vehicular",
                        'hospital': "Hospital General de Boca del Río",
                        'ubicacion': ub,
                        'personal': "Guardia PC Medellín",
                        'resumen': it.replace('\n', ' | ').strip(),
                        'wa_id': wa_id,
                        'timestamp': dt
                    })
                    # Paciente 2: Fallecido en sitio
                    eventos.append({
                        'origen': 'resumen_guardia',
                        'fecha': fecha_guardia,
                        'hora': h_sal,
                        'unidad': u,
                        'paciente': "Persona fallecida en sitio",
                        'edad': "N/D",
                        'sexo': "No especificado",
                        'diagnostico': "Defunción en sitio por accidente vehicular",
                        'hospital': "No aplica (Persona Fallecida en Sitio)",
                        'ubicacion': ub,
                        'personal': "Guardia PC Medellín",
                        'resumen': it.replace('\n', ' | ').strip(),
                        'wa_id': wa_id,
                        'timestamp': dt
                    })
                    continue

                hosp = estandarizar_hospital("", it, es_fallecido=es_muerte)
                
                motivo = ""
                for kw in ['posible iam y posible acb', 'posible iam', 'persona desvanecida', 'cambio de sonda',
                           'accidente con múltiples lesionados', 'accidente de moto', 'accidente',
                           'derrape de moto', 'volcadura', 'persona inconsciente', 'persona infartada',
                           'enfermedad general', 'caída de escaleras', 'caída de su propia altura',
                           'caída', 'traslado programado', 'traslado de menor', 'traslado de paciente', 'hemodiálisis',
                           'dificultad para respirar', 'convulsiones', 'combulciones', 'herida por arma blanca',
                           'intoxicación', 'intoxicacion', 'violencia familiar', 'crisis nerviosa']:
                    if kw in it_low:
                        motivo = kw.title()
                        break
                if not motivo: motivo = "Atención Prehospitalaria de Urgencia"

                pac = limpiar_nombre_paciente("", it)
                if not pac:
                    if es_muerte: pac = "Persona fallecida en sitio"
                    elif 'tercera edad' in it_low: pac = "Persona de la tercera edad"
                    elif 'menor' in it_low: pac = "Menor de edad"
                    elif 'femenina' in it_low: pac = "Femenina en valoración"
                    elif 'masculino' in it_low: pac = "Masculino en valoración"
                    else: pac = "Paciente en valoración prehospitalaria"

                res_clean = it.replace('\n', ' | ').strip()
                res_clean = re.sub(r'\*+', '', res_clean)

                eventos.append({
                    'origen': 'resumen_guardia',
                    'fecha': fecha_guardia,
                    'hora': h_sal,
                    'unidad': u,
                    'paciente': pac,
                    'edad': extraer_edad_valida(it),
                    'sexo': "Femenino" if 'femenina' in it_low else ("Masculino" if 'masculino' in it_low else "No especificado"),
                    'diagnostico': motivo,
                    'hospital': hosp,
                    'ubicacion': ub,
                    'personal': "Guardia PC Medellín",
                    'resumen': res_clean,
                    'wa_id': wa_id,
                    'timestamp': dt
                })

# MODO B: FRAPs Formales
for mid, dt, push, mtype, txt, tiene_img, wa_id, remitente in todos_2026:
    if not txt or len(txt) < 15: continue
    t_low = txt.lower()
    
    if any(k in t_low for k in ['novedades del d', 'reportes de servicios pc', 'novedades de pc', 'reporte de guardia']) and any(s in t_low for s in ['sale unidad', 'sale el compañero', 'salida', '1_']):
        continue

    if any(k in t_low for k in ['*nombre del paciente*', '*paciente*:', '*hospital de traslado*', 'hospital de traslado:']):
        f_rep_raw = extraer_multilinea(r'FECHA', txt)
        f_rep = normalizar_fecha(f_rep_raw, dt.strftime('%Y-%m-%d'))
        h_rep = extraer_multilinea(r'HORA', txt) or dt.strftime('%H:%M')
        u = detectar_unidad(extraer_multilinea(r'AMBULANCIA', txt) or txt)
        
        es_muerte = bool(PAT_MUERTE.search(txt))
        raw_pac = extraer_multilinea(r'NOMBRE\s*DEL\s*PACIENTE', txt)
        pac = limpiar_nombre_paciente(raw_pac, txt)
        if not pac:
            pac = "Persona fallecida en sitio" if es_muerte else "Paciente en valoración prehospitalaria"

        edad = extraer_edad_valida(txt)
        sexo = "Femenino" if 'femenin' in t_low else ("Masculino" if 'masculin' in t_low else "No especificado")
        
        dx = extraer_multilinea(r'DIAGN[OÓ]STICO', txt) or extraer_multilinea(r'MOTIVO\s*DE\s*LA\s*ATENCI[OÓ]N', txt) or "Atención Médica Prehospitalaria"
        raw_hosp = extraer_multilinea(r'HOSPITAL\s*DE\s*TRASLADO', txt)
        hosp = estandarizar_hospital(raw_hosp, t_low, es_fallecido=es_muerte)
        ub = estandarizar_ubicacion(txt)
        
        rem_alias = MAPA_REMITENTES.get(remitente, MAPA_REMITENTES.get(push, push or "Guardia Operativa PC Medellín"))
        
        res_clean = re.sub(r'\*+', '', txt)
        res_clean = re.sub(r'[\r\n]+', ' | ', res_clean)
        res_clean = re.sub(r'\s+', ' ', res_clean).strip(' |')

        eventos.append({
            'origen': 'frap_formal',
            'fecha': f_rep,
            'hora': h_rep[:5],
            'unidad': u,
            'paciente': pac,
            'edad': edad,
            'sexo': sexo,
            'diagnostico': dx,
            'hospital': hosp,
            'ubicacion': ub,
            'personal': f"Personal a Bordo ({rem_alias})",
            'resumen': res_clean,
            'wa_id': wa_id,
            'timestamp': dt
        })

# MODO C: Notas Clínicas
for mid, dt, push, mtype, txt, tiene_img, wa_id, remitente in todos_2026:
    if not txt or len(txt) < 15: continue
    t_low = txt.lower()
    
    if any(k in t_low for k in ['novedades del d', 'reportes de servicios pc', 'novedades de pc', 'reporte de guardia']) and any(s in t_low for s in ['sale unidad', 'sale el compañero', 'salida', '1_']):
        continue
    if any(k in t_low for k in ['*nombre del paciente*', '*paciente*:', '*hospital de traslado*', 'hospital de traslado:']):
        continue
        
    es_nota = False
    if any(k in t_low for k in ['signos vitales', 'a la ef', 'glasgow', 'toma de sv', 't/a 1', 't/a 9', 't/a 8', 't/a 12', 't/a: 1', 'fc1', 'fc8', 'fc9', 'spo2']):
        es_nota = True
    elif (re.search(r'\b(?:femenina|masculino)\s+de\s+\d{1,2}\s*a[ñn]os', t_low)):
        es_nota = True
    elif 'nombre del px' in t_low or 'nombre de la femenina' in t_low or 'nombre del masculino' in t_low or 'px de nombre' in t_low or 'px gediatrica' in t_low or 'segundo px' in t_low or 'px:' in t_low:
        es_nota = True

    if es_nota:
        es_muerte = bool(PAT_MUERTE.search(txt))
        pac = limpiar_nombre_paciente("", txt)
        if not pac:
            pac = "Persona fallecida en sitio" if es_muerte else "Paciente valorado en sitio"

        edad = extraer_edad_valida(txt)
        sexo = "Femenino" if 'femenin' in t_low else ("Masculino" if 'masculin' in t_low else "No especificado")
        hosp = estandarizar_hospital("", t_low, es_fallecido=es_muerte)
        ub = estandarizar_ubicacion(txt)
        u = detectar_unidad(txt)
        
        m_dx = re.search(r'\b(?:dx|diagn[oó]stico|posible|presentando|refiere|por|acude por)\s*:?\s*([^\n\.,;]+)', txt, re.IGNORECASE)
        dx = m_dx.group(1).strip().capitalize() if m_dx else "Evaluación Médica Prehospitalaria"
        if len(dx) > 60: dx = dx[:60] + "..."

        rem_alias = MAPA_REMITENTES.get(remitente, MAPA_REMITENTES.get(push, push or "Paramédico Evaluador PC"))

        res_clean = re.sub(r'\*+', '', txt)
        res_clean = re.sub(r'[\r\n]+', ' | ', res_clean)
        res_clean = re.sub(r'\s+', ' ', res_clean).strip(' |')

        eventos.append({
            'origen': 'nota_clinica',
            'fecha': dt.strftime('%Y-%m-%d'),
            'hora': dt.strftime('%H:%M'),
            'unidad': u,
            'paciente': pac,
            'edad': edad,
            'sexo': sexo,
            'diagnostico': dx,
            'hospital': hosp,
            'ubicacion': ub,
            'personal': f"Paramédico Evaluador ({rem_alias})",
            'resumen': res_clean,
            'wa_id': wa_id,
            'timestamp': dt
        })

# ======================================================================
# FASE 4: CORRELACIÓN Y FUSIÓN ESTRICTA (MULTI-PACIENTE)
# ======================================================================
print("Ejecutando motor de correlación estricta (respetando pacientes individuales)...")

def se_deben_fusionar(e1, e2):
    # Regla 1: Dos salidas de bitácora nunca se fusionan (cada ítem es una salida individual)
    if e1['origen'] == 'resumen_guardia' and e2['origen'] == 'resumen_guardia':
        return False
        
    # Regla 2: Mismo día obligatorio
    if e1['fecha'] != e2['fecha']:
        return False

    nombres_genericos = [
        'paciente no especificado en bitácora', 'paciente valorado en sitio',
        'paciente en valoración prehospitalaria', 'persona de la tercera edad',
        'menor de edad', 'femenina en valoración', 'masculino en valoración',
        'persona fallecida en sitio', 'no especificado'
    ]
    
    p1 = e1['paciente'].strip().lower()
    p2 = e2['paciente'].strip().lower()
    tiene_nom1 = p1 not in nombres_genericos
    tiene_nom2 = p2 not in nombres_genericos
    
    # Regla 3: Si ambos tienen nombres identificados pero son distintos -> NUNCA fusionar (son 2 pacientes distintos)
    if tiene_nom1 and tiene_nom2:
        tokens1 = set(p1.split())
        tokens2 = set(p2.split())
        inter = tokens1.intersection(tokens2)
        if len(inter) >= 2:
            return True # Mismo paciente reportado dos veces
        else:
            return False # Pacientes distintos (ej. Jesús Antonino vs Miguel Ángel Cruz)

    # Regla 4: Dos notas clínicas distintas nunca se fusionan si una dice "segundo px", "2do", etc.
    if e1['origen'] == 'nota_clinica' and e2['origen'] == 'nota_clinica':
        if any(w in e1['resumen'].lower() or w in e2['resumen'].lower() for w in ['segundo px', '2do px', 'segundo paciente', 'tercer px']):
            return False
        return False

    # Regla 5: Fusionar resumen de guardia con nota clínica SOLO si son el mismo servicio
    t_diff_min = 999
    try:
        h1 = [int(x) for x in e1['hora'].split(':')]
        h2 = [int(x) for x in e2['hora'].split(':')]
        t_diff_min = abs((h1[0]*60 + h1[1]) - (h2[0]*60 + h2[1]))
    except:
        pass

    if t_diff_min > 45:
        return False

    # Coincidencia por hospital receptor específico
    hosp1 = e1['hospital']
    hosp2 = e2['hospital']
    sitios = ['Atención en Sitio', 'Atención en Sitio (No ameritó traslado)', 'Atención en Sitio / Domicilio']
    if hosp1 not in sitios and hosp2 not in sitios and hosp1 == hosp2:
        return True

    # Coincidencia por palabras clave únicas del servicio
    kw_coincide = False
    for kw in ['policía municipal', 'policia municipal', 'hemodiálisis', 'hemodialisis', 'agustín mendoza', 'agustin mendoza', 'embarazo', 'parto', 'dengue', 'cristal']:
        if kw in e1['resumen'].lower() and kw in e2['resumen'].lower():
            kw_coincide = True
            break
            
    if kw_coincide and t_diff_min <= 60:
        return True

    return False

eventos_ordenados = sorted(eventos, key=lambda x: (x['fecha'], x['hora']))
eventos_finales = []
usados = set()

for i, e1 in enumerate(eventos_ordenados):
    if i in usados: continue
    merged_with = None
    for j in range(i + 1, len(eventos_ordenados)):
        if j in usados: continue
        e2 = eventos_ordenados[j]
        if se_deben_fusionar(e1, e2):
            merged_with = j
            break
            
    if merged_with is not None:
        e2 = eventos_ordenados[merged_with]
        usados.add(merged_with)
        
        nombres_genericos = [
            'paciente no especificado en bitácora', 'paciente valorado en sitio',
            'paciente en valoración prehospitalaria', 'persona de la tercera edad',
            'menor de edad', 'femenina en valoración', 'masculino en valoración',
            'persona fallecida en sitio', 'no especificado'
        ]
        pac_final = e1['paciente'] if e1['paciente'].strip().lower() not in nombres_genericos else e2['paciente']
        edad_final = e1['edad'] if e1['edad'] != 'N/D' else e2['edad']
        sexo_final = e1['sexo'] if e1['sexo'] != 'No especificado' else e2['sexo']
        hosp_final = e1['hospital'] if 'Sitio' not in e1['hospital'] else e2['hospital']
        ub_final = e1['ubicacion'] if e1['ubicacion'] != 'Medellín de Bravo, Ver.' else e2['ubicacion']
        
        res_combinado = f"{e1['resumen']} || [NOTA CLÍNICA VINCULADA]: {e2['resumen']}"
        
        eventos_finales.append({
            'fecha': e1['fecha'],
            'hora': e1['hora'],
            'unidad': e1['unidad'],
            'paciente': pac_final,
            'edad': edad_final,
            'sexo': sexo_final,
            'diagnostico': e1['diagnostico'] if len(e1['diagnostico']) > len(e2['diagnostico']) else e2['diagnostico'],
            'hospital': hosp_final,
            'ubicacion': ub_final,
            'personal': e1['personal'] if 'Personal' in e1['personal'] else e2['personal'],
            'resumen': res_combinado,
            'wa_id': e1['wa_id'],
            'timestamp': e1['timestamp']
        })
    else:
        eventos_finales.append(e1)

# Corregir posible año tipográfico 2025 en guardias de 2026
for e in eventos_finales:
    if e['fecha'].startswith('2025-') and e['timestamp'].year == 2026:
        e['fecha'] = '2026' + e['fecha'][4:]

# Estructuración
filas_excel = []

for idx, e in enumerate(eventos_finales, 1):
    hosp = e['hospital']
    t_full = (e['diagnostico'] + " " + hosp + " " + e['resumen']).lower()
    es_muerte = bool(PAT_MUERTE.search(e['diagnostico']) or PAT_MUERTE.search(e['resumen']))
    
    if es_muerte and hosp not in ['Hospital General de Boca del Río', 'Hospital General de Zona 71 IMSS Díaz Mirón', 'Hospital Regional de Alta Especialidad (HAEV)', 'Clínica Gaínza']:
        estatus = "Persona Fallecida en Sitio (Código 14)"
        recibe = "Policía Municipal / Periciales (FGE)"
        triage = "Negro (Fallecido)"
        hosp = "No aplica (Persona Fallecida en Sitio)"
    elif any(k in hosp.lower() for k in ['negativa', 'se niega', 'niega traslado', 'deslinde', 'responsiva']):
        estatus = "Negativa de Traslado (Responsiva firmada)"
        recibe = "N/A (Firma de responsiva)"
        triage = "Verde (Estable)"
    elif any(k in hosp.lower() for k in ['no ameritó', 'no amerit', 'no amerita']):
        estatus = "Atención Prehospitalaria en Sitio (No ameritó)"
        recibe = "N/A (Atención en sitio)"
        triage = "Verde (Estable)"
    elif any(k in hosp.lower() for k in ['alta médica', 'a su domicilio', 'alta a domicilio']):
        estatus = "Traslado a Domicilio (Alta hospitalaria)"
        recibe = "Familiar en Domicilio"
        triage = "Verde (Estable)"
    elif any(h in hosp for h in ['Hospital', 'Clínica', 'UMF', 'HGZ', 'IMSS', 'ISSSTE', 'PEMEX', 'Covadonga', 'Gaínza', 'Pediátrica', 'Naval', 'Militar', 'Star Médica', 'CRIVER', 'Sanatorio']):
        estatus = "Traslado Efectivo a Hospital"
        recibe = "Personal Médico de Guardia Hospitalaria"
        triage = "Rojo (Prioridad Alta)" if any(k in t_full for k in ['iam', 'infart', 'inconsciente', 'fractura', 'intoxic', 'acb', 'múltiples', 'choque', 'tce severo']) else "Amarillo (Moderado)"
    else:
        estatus = "Atención Prehospitalaria en Sitio"
        recibe = "N/A (Atención en sitio)"
        triage = "Amarillo (Moderado)" if any(k in t_full for k in ['violencia', 'golpea', 'caída', 'derrape', 'ketorolaco']) else "Verde (Estable)"

    folio_fmt = f"PC-2026-{idx:04d}"
    turno = clasificar_turno(e['hora'])
    cat_clinica = clasificar_categoria_clinica(e['diagnostico'], e['resumen'])
    sector_salud = clasificar_sector_salud(hosp)

    filas_excel.append({
        'No_Folio': folio_fmt,
        'Fecha': e['fecha'],
        'Hora': e['hora'],
        'Turno_Operativo': turno,
        'Unidad_Atiende': e['unidad'],
        'Personal_Que_Realizo_Traslado': e['personal'],
        'Persona_Trasladada_Paciente': e['paciente'],
        'Edad': e['edad'],
        'Sexo': e['sexo'],
        'Categoria_Clinica': cat_clinica,
        'Diagnostico_Motivo': e['diagnostico'],
        'Estatus_Traslado': estatus,
        'Hospital_Destino': hosp,
        'Sector_Salud': sector_salud,
        'Medico_Recibe': recibe,
        'Direccion_Ubicacion': e['ubicacion'],
        'Codigo_Triage': triage,
        'Resumen_Operativo': e['resumen']
    })

df_anual = pd.DataFrame(filas_excel)
df_anual = df_anual.sort_values(by=['Fecha', 'Hora']).reset_index(drop=True)
df_anual['No_Folio'] = [f"PC-2026-{i+1:04d}" for i in range(len(df_anual))]
df_anual['Mes'] = df_anual['Fecha'].astype(str).str[:7]

print(f"Total registros consolidados: {len(df_anual)}")

# ======================================================================
# FASE 6: CONSTRUCCIÓN DEL LIBRO EXCEL MULTI-HOJA
# ======================================================================
print("Generando libro Excel corregido con nombres explícitos y sin traslape de gráficas...")

wb = openpyxl.Workbook()

NAVY_HEADER = "1B365D"
WHITE = "FFFFFF"
GRAY_BORDER = "CBD5E1"
ZEBRA_FILL = "F8FAFC"
TITLE_COLOR = "0F294A"

font_title = Font(name='Arial', size=13, bold=True, color=TITLE_COLOR)
font_subtitle = Font(name='Arial', size=10, italic=True, color="4B5563")
font_sec_header = Font(name='Arial', size=11, bold=True, color=NAVY_HEADER)
font_table_hdr = Font(name='Arial', size=9, bold=True, color=WHITE)
fill_table_hdr = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
fill_zebra = PatternFill(start_color=ZEBRA_FILL, end_color=ZEBRA_FILL, fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
font_data = Font(name='Arial', size=9)
font_data_bold = Font(name='Arial', size=9, bold=True)

thin_border = Border(
    left=Side(style='thin', color=GRAY_BORDER),
    right=Side(style='thin', color=GRAY_BORDER),
    top=Side(style='thin', color=GRAY_BORDER),
    bottom=Side(style='thin', color=GRAY_BORDER)
)

def crear_kpi_card(ws, r_start, c_start, r_end, c_end, titulo, valor, subtitulo, bg_color, text_color):
    ws.merge_cells(start_row=r_start, start_column=c_start, end_row=r_start, end_column=c_end)
    ws.merge_cells(start_row=r_start+1, start_column=c_start, end_row=r_start+1, end_column=c_end)
    ws.merge_cells(start_row=r_start+2, start_column=c_start, end_row=r_start+2, end_column=c_end)
    
    c_tit = ws.cell(r_start, c_start, titulo)
    c_tit.font = Font(name='Arial', size=8, bold=True, color=text_color)
    c_tit.alignment = Alignment(horizontal='center', vertical='center')
    
    c_val = ws.cell(r_start+1, c_start, valor)
    c_val.font = Font(name='Arial', size=16, bold=True, color=text_color)
    c_val.alignment = Alignment(horizontal='center', vertical='center')
    
    c_sub = ws.cell(r_start+2, c_start, subtitulo)
    c_sub.font = Font(name='Arial', size=7, italic=True, color=text_color)
    c_sub.alignment = Alignment(horizontal='center', vertical='center')
    
    fill_card = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
    for r in range(r_start, r_start+3):
        for c in range(c_start, c_end+1):
            ws.cell(r, c).fill = fill_card
            ws.cell(r, c).border = thin_border

# ----------------------------------------------------------------------
# HOJA 1: 📊 DASHBOARD EJECUTIVO
# ----------------------------------------------------------------------
ws1 = wb.active
ws1.title = "📊 Dashboard Ejecutivo"
ws1.views.sheetView[0].showGridLines = True

ws1['A1'] = "H. AYUNTAMIENTO DE MEDELLÍN DE BRAVO, VERACRUZ"
ws1['A1'].font = font_title
ws1['A2'] = "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS"
ws1['A2'].font = Font(name='Arial', size=11, bold=True, color="374151")
ws1['A3'] = "DASHBOARD EJECUTIVO ANUAL DE SERVICIOS PREHOSPITALARIOS Y TRASLADOS (ENERO - SEPTIEMBRE 2026)"
ws1['A3'].font = font_subtitle

tot_serv = len(df_anual)
tras_hosp = len(df_anual[df_anual['Estatus_Traslado'] == 'Traslado Efectivo a Hospital'])
tras_dom = len(df_anual[df_anual['Estatus_Traslado'] == 'Traslado a Domicilio (Alta hospitalaria)'])
tot_tras = tras_hosp + tras_dom
aten_sitio = tot_serv - tot_tras
pct_tras = (tot_tras / tot_serv) * 100 if tot_serv > 0 else 0

# KPIs en filas 5-7
crear_kpi_card(ws1, 5, 1, 7, 2, "TOTAL DE SERVICIOS", str(tot_serv), "Período Ene - Sep 2026", "F1F5F9", "0F172A")
crear_kpi_card(ws1, 5, 3, 7, 4, "TRASLADOS A HOSPITAL", str(tras_hosp), f"{(tras_hosp/tot_serv)*100:.1f}% Urgencias canalizadas", "DEF7EC", "03543F")
crear_kpi_card(ws1, 5, 5, 7, 6, "ALTAS A DOMICILIO", str(tras_dom), f"{(tras_dom/tot_serv)*100:.1f}% Apoyo hospitalario", "E0F2FE", "0E7490")
crear_kpi_card(ws1, 5, 7, 7, 8, "ATENCIONES EN SITIO", str(aten_sitio), f"{(aten_sitio/tot_serv)*100:.1f}% Estabilizados en punto", "DBEAFE", "1E40AF")
turno_counts = df_anual['Turno_Operativo'].value_counts()
top_turno = turno_counts.index[0] if len(turno_counts) > 0 else "Matutino"
top_turno_cnt = turno_counts.iloc[0] if len(turno_counts) > 0 else 0
top_turno_nom = top_turno.split()[0].upper()
top_turno_pct = (top_turno_cnt / tot_serv) * 100 if tot_serv > 0 else 0
subtxt_turno = f"{top_turno_cnt} servicios ({top_turno_pct:.1f}%)"
crear_kpi_card(ws1, 5, 11, 7, 12, "TURNO PICO", top_turno_nom, subtxt_turno, "FEF3C7", "92400E")

ws1.row_dimensions[5].height = 16
ws1.row_dimensions[6].height = 24
ws1.row_dimensions[7].height = 14

# 1. TABLA DE EVOLUCIÓN MENSUAL (Filas 9 a 20)
ws1['A9'] = "1. EVOLUCIÓN MENSUAL DE SERVICIOS (ENERO A SEPTIEMBRE 2026)"
ws1['A9'].font = font_sec_header

cols_t1 = [("Mes de Registro", 18), ("Traslados a Hospital", 15), ("Atenciones en Sitio", 15), ("Total de Servicios", 15)]
for idx, (nom, w) in enumerate(cols_t1, 1):
    c = ws1.cell(10, idx, nom)
    c.fill = fill_table_hdr
    c.font = font_table_hdr
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = thin_border
    ws1.column_dimensions[get_column_letter(idx)].width = w

nombres_meses = [
    ('2026-01', 'Enero 2026'),
    ('2026-02', 'Febrero 2026'),
    ('2026-03', 'Marzo 2026'),
    ('2026-04', 'Abril 2026'),
    ('2026-05', 'Mayo 2026'),
    ('2026-06', 'Junio 2026'),
    ('2026-07', 'Julio 2026'),
    ('2026-08', 'Agosto 2026'),
    ('2026-09', 'Septiembre 2026')
]

r_t1 = 11
for ym, nom_mes in nombres_meses:
    df_m = df_anual[df_anual['Fecha'].astype(str).str.startswith(ym)]
    tot_m = len(df_m)
    tras_m = len(df_m[df_m['Estatus_Traslado'].isin(['Traslado Efectivo a Hospital', 'Traslado a Domicilio (Alta hospitalaria)'])])
    aten_m = tot_m - tras_m
    
    ws1.cell(r_t1, 1, nom_mes).alignment = Alignment(horizontal='left')
    ws1.cell(r_t1, 2, tras_m).alignment = Alignment(horizontal='center')
    ws1.cell(r_t1, 3, aten_m).alignment = Alignment(horizontal='center')
    ws1.cell(r_t1, 4, tot_m).alignment = Alignment(horizontal='center')
    for c_idx in range(1, 5):
        cell_obj = ws1.cell(r_t1, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_t1 % 2 == 0 else fill_zebra
    r_t1 += 1

# Fila Total
c1 = ws1.cell(r_t1, 1, "TOTAL ANUAL CONSOLIDADO"); c1.font = font_data_bold; c1.alignment = Alignment(horizontal='left'); c1.border = thin_border
c2 = ws1.cell(r_t1, 2, tot_tras); c2.font = font_data_bold; c2.alignment = Alignment(horizontal='center'); c2.border = thin_border
c3 = ws1.cell(r_t1, 3, aten_sitio); c3.font = font_data_bold; c3.alignment = Alignment(horizontal='center'); c3.border = thin_border
c4 = ws1.cell(r_t1, 4, tot_serv); c4.font = font_data_bold; c4.alignment = Alignment(horizontal='center'); c4.border = thin_border
for c_idx in range(1, 5): ws1.cell(r_t1, c_idx).fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")

# Gráfica 1: Evolución Mensual
chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "Evolución Mensual: Traslados a Hospital vs Atenciones en Sitio (2026)"
chart1.y_axis.title = "Número de Servicios"
chart1.x_axis.title = "Mes del Año 2026"
chart1.width = 18
chart1.height = 10
chart1.dataLabels = DataLabelList()
chart1.dataLabels.showVal = True
data1 = Reference(ws1, min_col=2, min_row=10, max_col=3, max_row=19)
cats1 = Reference(ws1, min_col=1, min_row=11, max_row=19)
chart1.add_data(data1, titles_from_data=True)
chart1.set_categories(cats1)
ws1.add_chart(chart1, "F9")

# 2. MATRIZ MENSUAL POR UNIDAD MÓVIL (Fila 23 en adelante)
ws1['A23'] = "2. MATRIZ MENSUAL POR UNIDAD MÓVIL OPERATIVA (ENERO - SEPTIEMBRE 2026)"
ws1['A23'].font = font_sec_header

cols_matriz = [
    ("Unidad Móvil", 26), ("Ene", 8), ("Feb", 8), ("Mar", 8), ("Abr", 8),
    ("May", 8), ("Jun", 8), ("Jul", 8), ("Ago", 8), ("Sep", 8), ("Total 2026", 12), ("% Part.", 10)
]
for idx, (nom, w) in enumerate(cols_matriz, 1):
    c = ws1.cell(24, idx, nom)
    c.fill = fill_table_hdr
    c.font = font_table_hdr
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = thin_border
    if idx > 4:
        ws1.column_dimensions[get_column_letter(idx)].width = w

r_mat = 25
unidades_orden = [
    "Ambulancia PC Medellín",
    "Unidad 097 (Ambulancia)",
    "Unidad 208 (Ambulancia)",
    "Unidad 098 (Ambulancia)",
    "Unidad Motorizada PC",
    "Unidad 096 (Rescate / Apoyo)"
]

for u_nom in unidades_orden:
    df_u = df_anual[df_anual['Unidad_Atiende'] == u_nom]
    ws1.cell(r_mat, 1, u_nom).alignment = Alignment(horizontal='left')
    u_tot = len(df_u)
    for m_idx, (ym, _) in enumerate(nombres_meses, 2):
        cant_m = len(df_u[df_u['Mes'] == ym])
        c_m = ws1.cell(r_mat, m_idx, cant_m)
        c_m.alignment = Alignment(horizontal='center')
    c_tot = ws1.cell(r_mat, 11, u_tot)
    c_tot.alignment = Alignment(horizontal='center')
    c_tot.font = font_data_bold
    c_pct = ws1.cell(r_mat, 12, f"{(u_tot/tot_serv)*100:.1f}%")
    c_pct.alignment = Alignment(horizontal='center')
    
    for c_idx in range(1, 13):
        cell_obj = ws1.cell(r_mat, c_idx)
        cell_obj.font = font_data if c_idx != 11 else font_data_bold
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_mat % 2 == 0 else fill_zebra
    r_mat += 1

# Fila Total Matriz
c_tot_lbl = ws1.cell(r_mat, 1, "TOTAL MENSUAL CONSOLIDADO")
c_tot_lbl.font = font_data_bold
c_tot_lbl.alignment = Alignment(horizontal='left')
for m_idx, (ym, _) in enumerate(nombres_meses, 2):
    tot_col = len(df_anual[df_anual['Mes'] == ym])
    c_col = ws1.cell(r_mat, m_idx, tot_col)
    c_col.alignment = Alignment(horizontal='center')
    c_col.font = font_data_bold
c_gran_tot = ws1.cell(r_mat, 11, tot_serv)
c_gran_tot.alignment = Alignment(horizontal='center')
c_gran_tot.font = font_data_bold
c_100 = ws1.cell(r_mat, 12, "100.0%")
c_100.alignment = Alignment(horizontal='center')
c_100.font = font_data_bold
for c_idx in range(1, 13):
    ws1.cell(r_mat, c_idx).border = thin_border
    ws1.cell(r_mat, c_idx).fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")

# 3. GRÁFICAS DE UNIDADES Y ESTATUS (Colocadas abajo sin traslape a partir de fila 34)
# Gráfica 2: Barras Horizontales con NOMBRES EXPLÍCITOS DE CADA UNIDAD y ETIQUETAS DE DATOS
chart2 = BarChart()
chart2.type = "bar"
chart2.style = 13
chart2.title = "Servicios Totales por Unidad Móvil (Enero - Septiembre 2026)"
chart2.width = 17
chart2.height = 10
chart2.y_axis.title = "Unidad Móvil de PC Medellín"
chart2.x_axis.title = "Total de Servicios Atendidos"
chart2.legend = None
chart2.dataLabels = DataLabelList()
chart2.dataLabels.showVal = True
# Referencia a la columna 'Unidad Móvil' (col 1, filas 25-30) y 'Total 2026' (col 11, filas 24-30)
data2 = Reference(ws1, min_col=11, min_row=24, max_row=r_mat-1)
cats2 = Reference(ws1, min_col=1, min_row=25, max_row=r_mat-1)
chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)
ws1.add_chart(chart2, "A34")

# Gráfica 3: Participación Porcentual por Unidad Móvil
chart3 = DoughnutChart()
chart3.title = "Participación Operativa por Unidad Móvil (%)"
chart3.width = 15
chart3.height = 10
chart3.dataLabels = DataLabelList()
chart3.dataLabels.showPercent = True
chart3.dataLabels.showVal = False
data3 = Reference(ws1, min_col=11, min_row=24, max_row=r_mat-1)
cats3 = Reference(ws1, min_col=1, min_row=25, max_row=r_mat-1)
chart3.add_data(data3, titles_from_data=True)
chart3.set_categories(cats3)
ws1.add_chart(chart3, "H34")

# ----------------------------------------------------------------------
# HOJA 2: 📈 PADECIMIENTOS Y DEMOGRAFÍA
# ----------------------------------------------------------------------
ws2 = wb.create_sheet(title="📈 Padecimientos y Demografía")
ws2.views.sheetView[0].showGridLines = True

ws2['A1'] = "PERFIL EPIDEMIOLÓGICO Y DEMOGRAFÍA DE PACIENTES (ENERO - SEPTIEMBRE 2026)"
ws2['A1'].font = font_title
ws2['A2'] = "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS DE MEDELLÍN DE BRAVO"
ws2['A2'].font = font_subtitle

# Sección 1: Categorías Clínicas
ws2['A4'] = "1. DISTRIBUCIÓN POR CATEGORÍA CLÍNICA Y PADECIMIENTOS (AÑO 2026)"
ws2['A4'].font = font_sec_header
ws2.cell(5, 1, "Categoría Clínica").fill = fill_table_hdr
ws2.cell(5, 1).font = font_table_hdr
ws2.cell(5, 2, "Casos").fill = fill_table_hdr
ws2.cell(5, 2).font = font_table_hdr
ws2.cell(5, 3, "%").fill = fill_table_hdr
ws2.cell(5, 3).font = font_table_hdr
for c_idx in range(1, 4): ws2.cell(5, c_idx).border = thin_border
ws2.column_dimensions['A'].width = 38
ws2.column_dimensions['B'].width = 12
ws2.column_dimensions['C'].width = 12

r_p1 = 6
for cat_nom, cant in df_anual['Categoria_Clinica'].value_counts().items():
    ws2.cell(r_p1, 1, cat_nom).alignment = Alignment(horizontal='left')
    ws2.cell(r_p1, 2, cant).alignment = Alignment(horizontal='center')
    ws2.cell(r_p1, 3, f"{(cant/tot_serv)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 4):
        cell_obj = ws2.cell(r_p1, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_p1 % 2 == 0 else fill_zebra
    r_p1 += 1

chart_p1 = BarChart()
chart_p1.type = "bar"
chart_p1.style = 10
chart_p1.title = "Urgencias por Categoría Clínica (Enero - Septiembre 2026)"
chart_p1.width = 17
chart_p1.height = 10
chart_p1.legend = None
chart_p1.dataLabels = DataLabelList()
chart_p1.dataLabels.showVal = True
chart_p1.x_axis.title = "Número de Casos"
chart_p1.y_axis.title = "Categoría Clínica"
data_p1 = Reference(ws2, min_col=2, min_row=5, max_row=r_p1-1)
cats_p1 = Reference(ws2, min_col=1, min_row=6, max_row=r_p1-1)
chart_p1.add_data(data_p1, titles_from_data=True)
chart_p1.set_categories(cats_p1)
ws2.add_chart(chart_p1, "E4")

# Sección 2: Grupos Etarios (Fila 17)
ws2['A17'] = "2. DISTRIBUCIÓN POR GRUPOS DE EDAD (AÑO 2026)"
ws2['A17'].font = font_sec_header
ws2.cell(18, 1, "Grupo Etario").fill = fill_table_hdr
ws2.cell(18, 1).font = font_table_hdr
ws2.cell(18, 2, "Rango de Edad").fill = fill_table_hdr
ws2.cell(18, 2).font = font_table_hdr
ws2.cell(18, 3, "Pacientes").fill = fill_table_hdr
ws2.cell(18, 3).font = font_table_hdr
ws2.cell(18, 4, "%").fill = fill_table_hdr
ws2.cell(18, 4).font = font_table_hdr
for c_idx in range(1, 5): ws2.cell(18, c_idx).border = thin_border

def clasificar_grupo_edad(edad_str):
    if not edad_str or edad_str in ['N/D', 'Desconocida']: return 'Sin registro de edad'
    m = re.search(r'\b(\d+)\b', str(edad_str))
    if not m: return 'Sin registro de edad'
    val = int(m.group(1))
    if 'mes' in str(edad_str).lower() or 'dia' in str(edad_str).lower() or val <= 12: return 'Pediátrico (0-12 años)'
    elif 13 <= val <= 25: return 'Jóvenes / Adolescentes (13-25 años)'
    elif 26 <= val <= 59: return 'Adultos (26-59 años)'
    else: return 'Adultos Mayores (60+ años)'

df_anual['Grupo_Edad'] = df_anual['Edad'].apply(clasificar_grupo_edad)

grupos_meta = [
    ('Pediátrico (0-12 años)', '0 - 12 años'),
    ('Jóvenes / Adolescentes (13-25 años)', '13 - 25 años'),
    ('Adultos (26-59 años)', '26 - 59 años'),
    ('Adultos Mayores (60+ años)', '60 años o más'),
    ('Sin registro de edad', 'No documentado en bitácora')
]

r_p2 = 19
for g_nom, g_rango in grupos_meta:
    cant = len(df_anual[df_anual['Grupo_Edad'] == g_nom])
    ws2.cell(r_p2, 1, g_nom).alignment = Alignment(horizontal='left')
    ws2.cell(r_p2, 2, g_rango).alignment = Alignment(horizontal='center')
    ws2.cell(r_p2, 3, cant).alignment = Alignment(horizontal='center')
    ws2.cell(r_p2, 4, f"{(cant/tot_serv)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 5):
        cell_obj = ws2.cell(r_p2, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_p2 % 2 == 0 else fill_zebra
    r_p2 += 1

chart_p2 = BarChart()
chart_p2.type = "col"
chart_p2.style = 11
chart_p2.title = "Pacientes por Grupo Etario (Enero - Septiembre 2026)"
chart_p2.width = 17
chart_p2.height = 9
chart_p2.legend = None
chart_p2.dataLabels = DataLabelList()
chart_p2.dataLabels.showVal = True
chart_p2.x_axis.title = "Grupo de Edad"
chart_p2.y_axis.title = "Número de Pacientes"
data_p2 = Reference(ws2, min_col=3, min_row=18, max_row=r_p2-1)
cats_p2 = Reference(ws2, min_col=1, min_row=19, max_row=r_p2-1)
chart_p2.add_data(data_p2, titles_from_data=True)
chart_p2.set_categories(cats_p2)
ws2.add_chart(chart_p2, "E17")

# Sección 3: Distribución por Género (Fila 27)
ws2['A27'] = "3. DISTRIBUCIÓN POR GÉNERO / SEXO (AÑO 2026)"
ws2['A27'].font = font_sec_header
ws2.cell(28, 1, "Género").fill = fill_table_hdr
ws2.cell(28, 1).font = font_table_hdr
ws2.cell(28, 2, "Pacientes").fill = fill_table_hdr
ws2.cell(28, 2).font = font_table_hdr
ws2.cell(28, 3, "%").fill = fill_table_hdr
ws2.cell(28, 3).font = font_table_hdr
for c_idx in range(1, 4): ws2.cell(28, c_idx).border = thin_border

r_p3 = 29
for sex_nom, cant in df_anual['Sexo'].value_counts().items():
    ws2.cell(r_p3, 1, sex_nom).alignment = Alignment(horizontal='left')
    ws2.cell(r_p3, 2, cant).alignment = Alignment(horizontal='center')
    ws2.cell(r_p3, 3, f"{(cant/tot_serv)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 4):
        cell_obj = ws2.cell(r_p3, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_p3 % 2 == 0 else fill_zebra
    r_p3 += 1

chart_p3 = PieChart()
chart_p3.title = "Proporción por Sexo (Enero - Septiembre 2026)"
chart_p3.width = 14
chart_p3.height = 8.5
chart_p3.dataLabels = DataLabelList()
chart_p3.dataLabels.showPercent = True
chart_p3.dataLabels.showVal = False
data_p3 = Reference(ws2, min_col=2, min_row=28, max_row=r_p3-1)
cats_p3 = Reference(ws2, min_col=1, min_row=29, max_row=r_p3-1)
chart_p3.add_data(data_p3, titles_from_data=True)
chart_p3.set_categories(cats_p3)
ws2.add_chart(chart_p3, "E27")

# ----------------------------------------------------------------------
# HOJA 3: 📍 COBERTURA Y TURNOS
# ----------------------------------------------------------------------
ws3 = wb.create_sheet(title="📍 Cobertura y Turnos")
ws3.views.sheetView[0].showGridLines = True

ws3['A1'] = "ANÁLISIS DE COBERTURA GEOGRÁFICA Y CARGA HORARIA (ENERO - SEPTIEMBRE 2026)"
ws3['A1'].font = font_title
ws3['A2'] = "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS DE MEDELLÍN DE BRAVO"
ws3['A2'].font = font_subtitle

# Sección 1: Localidades Top
ws3['A4'] = "1. COBERTURA POR LOCALIDAD Y FRACCIONAMIENTO (TOP 10 ZONAS 2026)"
ws3['A4'].font = font_sec_header
ws3.cell(5, 1, "Localidad / Fraccionamiento").fill = fill_table_hdr
ws3.cell(5, 1).font = font_table_hdr
ws3.cell(5, 2, "Servicios").fill = fill_table_hdr
ws3.cell(5, 2).font = font_table_hdr
ws3.cell(5, 3, "%").fill = fill_table_hdr
ws3.cell(5, 3).font = font_table_hdr
for c_idx in range(1, 4): ws3.cell(5, c_idx).border = thin_border
ws3.column_dimensions['A'].width = 38
ws3.column_dimensions['B'].width = 12
ws3.column_dimensions['C'].width = 12

r_g1 = 6
top_locs = df_anual['Direccion_Ubicacion'].value_counts().head(10)
for loc_nom, cant in top_locs.items():
    loc_clean = loc_nom.replace(', Medellín de Bravo, Ver.', '')
    ws3.cell(r_g1, 1, loc_clean).alignment = Alignment(horizontal='left')
    ws3.cell(r_g1, 2, cant).alignment = Alignment(horizontal='center')
    ws3.cell(r_g1, 3, f"{(cant/tot_serv)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 4):
        cell_obj = ws3.cell(r_g1, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_g1 % 2 == 0 else fill_zebra
    r_g1 += 1

chart_g1 = BarChart()
chart_g1.type = "bar"
chart_g1.style = 10
chart_g1.title = "Top 10 Zonas con Mayor Demanda Prehospitalaria (2026)"
chart_g1.width = 17
chart_g1.height = 11
chart_g1.legend = None
chart_g1.dataLabels = DataLabelList()
chart_g1.dataLabels.showVal = True
chart_g1.x_axis.title = "Servicios de Ambulancia"
chart_g1.y_axis.title = "Zona / Fraccionamiento"
data_g1 = Reference(ws3, min_col=2, min_row=5, max_row=r_g1-1)
cats_g1 = Reference(ws3, min_col=1, min_row=6, max_row=r_g1-1)
chart_g1.add_data(data_g1, titles_from_data=True)
chart_g1.set_categories(cats_g1)
ws3.add_chart(chart_g1, "E4")

# Sección 2: Carga Horaria por Turnos (Fila 19)
ws3['A19'] = "2. CARGA OPERATIVA POR TURNO DE GUARDIA (ENERO - SEPTIEMBRE 2026)"
ws3['A19'].font = font_sec_header
ws3.cell(20, 1, "Turno de Guardia").fill = fill_table_hdr
ws3.cell(20, 1).font = font_table_hdr
ws3.cell(20, 2, "Horario Operativo").fill = fill_table_hdr
ws3.cell(20, 2).font = font_table_hdr
ws3.cell(20, 3, "Servicios").fill = fill_table_hdr
ws3.cell(20, 3).font = font_table_hdr
ws3.cell(20, 4, "%").fill = fill_table_hdr
ws3.cell(20, 4).font = font_table_hdr
for c_idx in range(1, 5): ws3.cell(20, c_idx).border = thin_border

turnos_meta = [
    ('Matutino (07:00 - 14:59)', '07:00 a 14:59 hrs'),
    ('Vespertino (15:00 - 22:59)', '15:00 a 22:59 hrs'),
    ('Nocturno (23:00 - 06:59)', '23:00 a 06:59 hrs')
]

r_g2 = 21
for t_nom, t_horario in turnos_meta:
    cant = len(df_anual[df_anual['Turno_Operativo'] == t_nom])
    ws3.cell(r_g2, 1, t_nom.split(' (')[0]).alignment = Alignment(horizontal='left')
    ws3.cell(r_g2, 2, t_horario).alignment = Alignment(horizontal='center')
    ws3.cell(r_g2, 3, cant).alignment = Alignment(horizontal='center')
    ws3.cell(r_g2, 4, f"{(cant/tot_serv)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 5):
        cell_obj = ws3.cell(r_g2, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_g2 % 2 == 0 else fill_zebra
    r_g2 += 1

chart_g2 = DoughnutChart()
chart_g2.title = "Carga de Servicios por Turno de Guardia (Ene - Sep 2026)"
chart_g2.width = 14
chart_g2.height = 9
chart_g2.dataLabels = DataLabelList()
chart_g2.dataLabels.showPercent = True
chart_g2.dataLabels.showVal = False
data_g2 = Reference(ws3, min_col=3, min_row=20, max_row=r_g2-1)
cats_g2 = Reference(ws3, min_col=1, min_row=21, max_row=r_g2-1)
chart_g2.add_data(data_g2, titles_from_data=True)
chart_g2.set_categories(cats_g2)
ws3.add_chart(chart_g2, "E19")

# ----------------------------------------------------------------------
# HOJA 4: 🏥 RED HOSPITALARIA Y TRIAGE
# ----------------------------------------------------------------------
ws4 = wb.create_sheet(title="🏥 Red Hospitalaria y Triage")
ws4.views.sheetView[0].showGridLines = True

ws4['A1'] = "ARTICULACIÓN CON EL SECTOR SALUD Y CLASIFICACIÓN DE TRIAGE (ENERO - SEPTIEMBRE 2026)"
ws4['A1'].font = font_title
ws4['A2'] = "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS DE MEDELLÍN DE BRAVO"
ws4['A2'].font = font_subtitle

# Sección 1: Absorción por Sector Salud (Filas 4 a 12)
ws4['A4'] = "1. ABSORCIÓN POR SECTOR DE SALUD RECEPTOR (2026)"
ws4['A4'].font = font_sec_header
ws4.cell(5, 1, "Sector de Salud").fill = fill_table_hdr
ws4.cell(5, 1).font = font_table_hdr
ws4.cell(5, 2, "Pacientes").fill = fill_table_hdr
ws4.cell(5, 2).font = font_table_hdr
ws4.cell(5, 3, "%").fill = fill_table_hdr
ws4.cell(5, 3).font = font_table_hdr
for c_idx in range(1, 4): ws4.cell(5, c_idx).border = thin_border
ws4.column_dimensions['A'].width = 38
ws4.column_dimensions['B'].width = 14
ws4.column_dimensions['C'].width = 12

r_h1 = 6
for sec_nom, cant in df_anual['Sector_Salud'].value_counts().items():
    ws4.cell(r_h1, 1, sec_nom).alignment = Alignment(horizontal='left')
    ws4.cell(r_h1, 2, cant).alignment = Alignment(horizontal='center')
    ws4.cell(r_h1, 3, f"{(cant/tot_serv)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 4):
        cell_obj = ws4.cell(r_h1, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_h1 % 2 == 0 else fill_zebra
    r_h1 += 1

chart_h1 = PieChart()
chart_h1.title = "Absorción de Pacientes por Sector Salud (2026)"
chart_h1.width = 15
chart_h1.height = 9.5
chart_h1.dataLabels = DataLabelList()
chart_h1.dataLabels.showPercent = True
chart_h1.dataLabels.showVal = False
data_h1 = Reference(ws4, min_col=2, min_row=5, max_row=r_h1-1)
cats_h1 = Reference(ws4, min_col=1, min_row=6, max_row=r_h1-1)
chart_h1.add_data(data_h1, titles_from_data=True)
chart_h1.set_categories(cats_h1)
ws4.add_chart(chart_h1, "E4")

# Sección 2: Top Hospitales Receptores (Fila 16)
ws4['A16'] = "2. RANKING DE HOSPITALES Y CLÍNICAS RECEPTORAS (TOP 8)"
ws4['A16'].font = font_sec_header
ws4.cell(17, 1, "Hospital / Clínica de Recepción").fill = fill_table_hdr
ws4.cell(17, 1).font = font_table_hdr
ws4.cell(17, 2, "Pacientes").fill = fill_table_hdr
ws4.cell(17, 2).font = font_table_hdr
ws4.cell(17, 3, "%").fill = fill_table_hdr
ws4.cell(17, 3).font = font_table_hdr
for c_idx in range(1, 4): ws4.cell(17, c_idx).border = thin_border

hosp_efectivos = df_anual[df_anual['Estatus_Traslado'] == 'Traslado Efectivo a Hospital']
top_hosp = hosp_efectivos['Hospital_Destino'].value_counts().head(8)
r_h2 = 18
for h_nom, cant in top_hosp.items():
    ws4.cell(r_h2, 1, h_nom).alignment = Alignment(horizontal='left')
    ws4.cell(r_h2, 2, cant).alignment = Alignment(horizontal='center')
    ws4.cell(r_h2, 3, f"{(cant/tras_hosp)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 4):
        cell_obj = ws4.cell(r_h2, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_h2 % 2 == 0 else fill_zebra
    r_h2 += 1

chart_h2 = BarChart()
chart_h2.type = "bar"
chart_h2.style = 13
chart_h2.title = "Pacientes Recibidos por Hospital (Ene - Sep 2026)"
chart_h2.width = 16
chart_h2.height = 9.5
chart_h2.legend = None
chart_h2.dataLabels = DataLabelList()
chart_h2.dataLabels.showVal = True
chart_h2.x_axis.title = "Pacientes Recibidos"
chart_h2.y_axis.title = "Hospital / Nosocomio"
data_h2 = Reference(ws4, min_col=2, min_row=17, max_row=r_h2-1)
cats_h2 = Reference(ws4, min_col=1, min_row=18, max_row=r_h2-1)
chart_h2.add_data(data_h2, titles_from_data=True)
chart_h2.set_categories(cats_h2)
ws4.add_chart(chart_h2, "E16")

# Sección 3: Distribución de Triage (Fila 28)
ws4['A28'] = "3. CLASIFICACIÓN POR CÓDIGO TRIAGE DE URGENCIA (2026)"
ws4['A28'].font = font_sec_header
ws4.cell(29, 1, "Código Triage").fill = fill_table_hdr
ws4.cell(29, 1).font = font_table_hdr
ws4.cell(29, 2, "Significado Clínico").fill = fill_table_hdr
ws4.cell(29, 2).font = font_table_hdr
ws4.cell(29, 3, "Pacientes").fill = fill_table_hdr
ws4.cell(29, 3).font = font_table_hdr
ws4.cell(29, 4, "%").fill = fill_table_hdr
ws4.cell(29, 4).font = font_table_hdr
for c_idx in range(1, 5): ws4.cell(29, c_idx).border = thin_border

triage_meta = [
    ('Rojo (Prioridad Alta)', 'Emergencia Vital Crítica Inmediata'),
    ('Amarillo (Moderado)', 'Urgencia Médica Estable / Trauma Moderado'),
    ('Verde (Estable)', 'Urgencia Menor / Atención en Sitio'),
    ('Negro (Fallecido)', 'Persona Fallecida en Sitio (Código 14)')
]

r_h3 = 30
for tr_nom, tr_sig in triage_meta:
    cant = len(df_anual[df_anual['Codigo_Triage'] == tr_nom])
    ws4.cell(r_h3, 1, tr_nom).alignment = Alignment(horizontal='left')
    ws4.cell(r_h3, 2, tr_sig).alignment = Alignment(horizontal='left')
    ws4.cell(r_h3, 3, cant).alignment = Alignment(horizontal='center')
    ws4.cell(r_h3, 4, f"{(cant/tot_serv)*100:.1f}%").alignment = Alignment(horizontal='center')
    for c_idx in range(1, 5):
        cell_obj = ws4.cell(r_h3, c_idx)
        cell_obj.font = font_data
        cell_obj.border = thin_border
        cell_obj.fill = fill_white if r_h3 % 2 == 0 else fill_zebra
    r_h3 += 1

chart_h3 = BarChart()
chart_h3.type = "col"
chart_h3.style = 10
chart_h3.title = "Distribución de Triage Prehospitalario (2026)"
chart_h3.width = 16
chart_h3.height = 8.5
chart_h3.legend = None
chart_h3.dataLabels = DataLabelList()
chart_h3.dataLabels.showVal = True
chart_h3.x_axis.title = "Código Triage"
chart_h3.y_axis.title = "Número de Pacientes"
data_h3 = Reference(ws4, min_col=3, min_row=29, max_row=r_h3-1)
cats_h3 = Reference(ws4, min_col=1, min_row=30, max_row=r_h3-1)
chart_h3.add_data(data_h3, titles_from_data=True)
chart_h3.set_categories(cats_h3)
ws4.add_chart(chart_h3, "E28")

# ----------------------------------------------------------------------
# HOJA 5: 📋 CATÁLOGO MAESTRO 2026
# ----------------------------------------------------------------------
ws5 = wb.create_sheet(title="📋 Catálogo Maestro 2026")
ws5.views.sheetView[0].showGridLines = True

ws5.merge_cells('A1:R1')
ws5['A1'] = "H. AYUNTAMIENTO DE MEDELLÍN DE BRAVO, VERACRUZ"
ws5['A1'].font = Font(name='Arial', size=14, bold=True, color=TITLE_COLOR)
ws5['A1'].alignment = Alignment(horizontal='center', vertical='center')

ws5.merge_cells('A2:R2')
ws5['A2'] = "DIRECCIÓN MUNICIPAL DE PROTECCIÓN CIVIL Y BOMBEROS"
ws5['A2'].font = Font(name='Arial', size=11, bold=True, color="374151")
ws5['A2'].alignment = Alignment(horizontal='center', vertical='center')

ws5.merge_cells('A3:R3')
ws5['A3'] = "CATÁLOGO MAESTRO Y CONTROL OFICIAL DE SERVICIOS PREHOSPITALARIOS Y TRASLADOS DE AMBULANCIA 2026"
ws5['A3'].font = Font(name='Arial', size=10, italic=True, color="4B5563")
ws5['A3'].alignment = Alignment(horizontal='center', vertical='center')

ws5.row_dimensions[1].height = 24
ws5.row_dimensions[2].height = 18
ws5.row_dimensions[3].height = 18
ws5.row_dimensions[4].height = 8
ws5.row_dimensions[5].height = 28

columnas_maestro = [
    ("No. Folio", 14, 'center'),
    ("Fecha", 13, 'center'),
    ("Hora", 10, 'center'),
    ("Turno Operativo", 24, 'center'),
    ("Unidad que Atiende", 24, 'center'),
    ("Personal a Cargo del Traslado", 36, 'left'),
    ("Nombre de la Persona Trasladada (Paciente)", 36, 'left'),
    ("Edad", 12, 'center'),
    ("Sexo", 12, 'center'),
    ("Categoría Clínica", 30, 'left'),
    ("Diagnóstico / Motivo del Traslado", 36, 'left'),
    ("Estatus del Traslado", 32, 'center'),
    ("Hospital de Destino / Receptor", 38, 'left'),
    ("Sector de Salud", 26, 'left'),
    ("Médico / Personal que Recibe", 28, 'left'),
    ("Dirección / Ubicación del Servicio", 36, 'left'),
    ("Código Triage", 20, 'center'),
    ("Resumen del Reporte Oficial", 85, 'left')
]

for col_idx, (col_name, col_width, col_align) in enumerate(columnas_maestro, 1):
    cell = ws5.cell(row=5, column=col_idx)
    cell.value = col_name
    cell.fill = fill_table_hdr
    cell.font = Font(name='Arial', size=10, bold=True, color=WHITE)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border
    col_letter = get_column_letter(col_idx)
    ws5.column_dimensions[col_letter].width = col_width

fill_verde_traslado = PatternFill(start_color="DEF7EC", end_color="DEF7EC", fill_type="solid")
font_verde_traslado = Font(name='Arial', size=9, bold=True, color="03543F")

fill_naranja_negada = PatternFill(start_color="FEEBC8", end_color="FEEBC8", fill_type="solid")
font_naranja_negada = Font(name='Arial', size=9, bold=True, color="C05621")

fill_azul_sitio = PatternFill(start_color="EBF8FF", end_color="EBF8FF", fill_type="solid")
font_azul_sitio = Font(name='Arial', size=9, color="2B6CB0")

fill_rojo_triage = PatternFill(start_color="FED7D7", end_color="FED7D7", fill_type="solid")
font_rojo_triage = Font(name='Arial', size=9, bold=True, color="9B2C2C")

fill_amarillo_triage = PatternFill(start_color="FEFCBF", end_color="FEFCBF", fill_type="solid")
font_amarillo_triage = Font(name='Arial', size=9, bold=True, color="975A16")

fill_verde_triage = PatternFill(start_color="C6F6D5", end_color="C6F6D5", fill_type="solid")
font_verde_triage = Font(name='Arial', size=9, bold=True, color="22543D")

fill_negro_triage = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
font_negro_triage = Font(name='Arial', size=9, bold=True, color="1E293B")

for row_idx, r in df_anual.iterrows():
    current_row = row_idx + 6
    ws5.row_dimensions[current_row].height = 20
    is_even = (row_idx % 2 == 0)
    row_fill = fill_white if is_even else fill_zebra
    
    valores = [
        r['No_Folio'],
        r['Fecha'],
        r['Hora'],
        r['Turno_Operativo'],
        r['Unidad_Atiende'],
        r['Personal_Que_Realizo_Traslado'],
        r['Persona_Trasladada_Paciente'],
        r['Edad'],
        r['Sexo'],
        r['Categoria_Clinica'],
        r['Diagnostico_Motivo'],
        r['Estatus_Traslado'],
        r['Hospital_Destino'],
        r['Sector_Salud'],
        r['Medico_Recibe'],
        r['Direccion_Ubicacion'],
        r['Codigo_Triage'],
        r['Resumen_Operativo']
    ]
    
    for col_idx, val in enumerate(valores, 1):
        cell = ws5.cell(row=current_row, column=col_idx)
        cell.value = val
        cell.font = font_data
        cell.border = thin_border
        cell.fill = row_fill
        
        _, _, col_align = columnas_maestro[col_idx - 1]
        cell.alignment = Alignment(horizontal=col_align, vertical='center', wrap_text=(col_idx == 18))
        
        if col_idx == 12:
            if "Traslado Efectivo" in str(val):
                cell.fill = fill_verde_traslado
                cell.font = font_verde_traslado
            elif "Negativa" in str(val):
                cell.fill = fill_naranja_negada
                cell.font = font_naranja_negada
            elif "Atención Prehospitalaria" in str(val):
                cell.fill = fill_azul_sitio
                cell.font = font_azul_sitio
                
        elif col_idx == 17:
            if "Rojo" in str(val):
                cell.fill = fill_rojo_triage
                cell.font = font_rojo_triage
            elif "Amarillo" in str(val):
                cell.fill = fill_amarillo_triage
                cell.font = font_amarillo_triage
            elif "Verde" in str(val):
                cell.fill = fill_verde_triage
                cell.font = font_verde_triage
            elif "Negro" in str(val):
                cell.fill = fill_negro_triage
                cell.font = font_negro_triage

ws5.auto_filter.ref = f"A5:R{len(df_anual) + 5}"
ws5.freeze_panes = "A6"

# Guardado
output_path = "c:/Users/burel/OneDrive/Documentos/PCivil Digital/Catalogo_Anual_Completo_Traslados_Medellin_2026.xlsx"

try:
    wb.save(output_path)
    print(f"\n¡Éxito! Libro Ejecutivo Anual guardado en: {output_path}")
except PermissionError:
    alt_path = "c:/Users/burel/OneDrive/Documentos/PCivil Digital/Catalogo_Anual_Completo_Traslados_Medellin_2026_V2.xlsx"
    wb.save(alt_path)
    print(f"\nArchivo en uso. Guardado exitosamente en ruta alternativa: {alt_path}")
    output_path = alt_path

conn.close()
print("=" * 80)
print(f"PROCESO CONCLUIDO CON ÉXITO: 5 HOJAS Y 11 GRÁFICAS GENERADAS")
print("=" * 80)
