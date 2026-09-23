import openpyxl

wb = openpyxl.load_workbook('Catalogo_Oficial_Traslados_Medellin_2026.xlsx')
ws = wb['Catálogo General de Traslados']

print("Verificando consistencia de columnas...")

# Verificar Hospital Destino
hospitales = set()
for r in range(6, ws.max_row + 1):
    h = ws.cell(row=r, column=10).value
    hospitales.add(str(h))

print("\nMuestra de Hospitales:")
for h in sorted(list(hospitales))[:15]:
    print("  -", h)

# Verificar Pacientes
pacientes = set()
for r in range(6, ws.max_row + 1):
    p = ws.cell(row=r, column=6).value
    pacientes.add(str(p))

print(f"\nTotal pacientes únicos: {len(pacientes)}")
for p in sorted(list(pacientes))[:15]:
    print("  -", p)
