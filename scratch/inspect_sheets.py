import openpyxl
import pandas as pd
import sqlite3

print("=== HOJA 020926 EN 100 dias.xlsx ===")
wb = openpyxl.load_workbook('Documentacion/ORIGEN/100 dias.xlsx', data_only=True)
ws = wb['020926']
headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
print("Columnas en Fila 4:", [(idx+1, h) for idx, h in enumerate(headers) if h is not None])

# Ver los datos en esa hoja
data_rows = []
for r in range(5, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, len(headers)+1)]
    if any(v is not None for v in vals):
        data_rows.append(vals)
df_0209 = pd.DataFrame(data_rows, columns=headers)
print("Dimensiones df_0209:", df_0209.shape)
for col in df_0209.columns:
    if col and any(w in str(col).lower() for w in ['traslado', 'emerg', 'mdica', 'medica']):
        print(f"Col: {col}")
        print(df_0209[col].value_counts(dropna=False).head(10))

print("\n=== HOJAS EN Reportes_Agosto_Completo_Proteccion_Civil.xlsx ===")
wb_agosto = openpyxl.load_workbook('Reportes_Agosto_Completo_Proteccion_Civil.xlsx', data_only=True)
print("Hojas agosto:", wb_agosto.sheetnames)
for sname in wb_agosto.sheetnames:
    ws_a = wb_agosto[sname]
    print(f"Hoja {sname}: {ws_a.max_row} filas, {ws_a.max_column} cols")
    # encabezados
    r1 = [str(ws_a.cell(row=1, column=c).value) for c in range(1, min(ws_a.max_column+1, 15))]
    print("  R1:", r1)
