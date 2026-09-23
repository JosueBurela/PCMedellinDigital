import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('Documentacion/ORIGEN/informe de resultados.docx')
print("=== CONTENIDO COMPLETO DE 'informe de resultados.docx' ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")

for t_idx, t in enumerate(doc.tables):
    print(f"\n--- TABLA {t_idx} ---")
    for r in t.rows:
        print([c.text.strip().replace('\n', ' ') for c in r.cells])
