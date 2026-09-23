import openpyxl
import docx
import pandas as pd
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("=== 1. DETALLE DE HOJAS EN 100 dias.xlsx ===")
wb = openpyxl.load_workbook('Documentacion/ORIGEN/100 dias.xlsx', data_only=True)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"\n--- Hoja: {sheet} (max_row: {ws.max_row}, max_col: {ws.max_column}) ---")
    # Buscar palabras clave en todas las celdas de la hoja
    for r in range(1, ws.max_row + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        row_str = " | ".join(str(v) for v in row_vals if v is not None)
        if any(w in row_str.lower() for w in ['traslado', 'traslados', 'total', 'médic', 'medic', 'u-097', 'u-208', 'u-098']):
            print(f"  Fila {r:02d}: {row_str[:160]}")

print("\n=== 2. DETALLE EN R Antonio Rosas-Anel.docx ===")
doc2 = docx.Document('Documentacion/ORIGEN/R Antonio Rosas-Anel.docx')
for i, p in enumerate(doc2.paragraphs):
    if p.text.strip():
        print(f"P{i:02d}: {p.text[:140]}")
for t_i, t in enumerate(doc2.tables):
    print(f"\n--- Tabla {t_i} en R Antonio Rosas-Anel.docx ---")
    for r in t.rows:
        print([c.text.strip().replace('\n', ' ') for c in r.cells])

print("\n=== 3. DETALLE EN Reportes_Agosto_Completo_Proteccion_Civil.xlsx ===")
wb_ag = openpyxl.load_workbook('Reportes_Agosto_Completo_Proteccion_Civil.xlsx', data_only=True)
print("Hojas agosto completo:", wb_ag.sheetnames)
ws_ag = wb_ag.active
print(f"Filas: {ws_ag.max_row}, Cols: {ws_ag.max_column}")
for r in range(1, 15):
    print([ws_ag.cell(row=r, column=c).value for c in range(1, ws_ag.max_column + 1)])
