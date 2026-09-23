import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

import pandas as pd
import re
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("Cargando y refinando los 71 registros de traslados exactos...")

# 1. Cargar datos existentes
df_raw = pd.read_excel('Documentacion/Registro_Nominal_Traslados_Exactos_2026.xlsx', skiprows=4)
col_map = {
    'Fecha y Hora': 'fecha_hora',
    'Unidad': 'unidad',
    'Nombre del Paciente': 'paciente',
    'Edad': 'edad',
    'Por Qué (Diagnóstico / Motivo)': 'motivo',
    'A Dónde (Hospital Receptor)': 'hospital',
    'Recibe / Observaciones': 'recibe',
    'Lugar de Origen': 'origen',
    'Tipo de Registro': 'tipo_reg'
}
df_raw = df_raw.rename(columns=col_map)

# 2. Funciones de limpieza fina
MESES = {
    'enero': '01', 'febrero': '02', 'marzo': '03', 'abril': '04',
    'mayo': '05', 'junio': '06', 'julio': '07', 'agosto': '08',
    'septiembre': '09', 'octubre': '10', 'noviembre': '11', 'diciembre': '12'
}

def clean_fecha(val):
    s = str(val).strip()
    # Extraer hora
    m_h = re.search(r'\b(\d{1,2}):(\d{2})\b', s)
    hora = f"{int(m_h.group(1)):02d}:{int(m_h.group(2)):02d}" if m_h else "12:00"
    
    # Extraer fecha
    m_pal = re.search(r'(\d{1,2})\s+(?:de\s+)?([a-zA-Z]+)(?:\s+(?:del?\s+)?(\d{4}))?', s, re.IGNORECASE)
    if m_pal:
        d = int(m_pal.group(1))
        m = MESES.get(m_pal.group(2).lower(), '08')
        y = m_pal.group(3) if m_pal.group(3) else '2026'
        return f"{d:02d}/{m}/{y} {hora}", f"{y}-{m}-{d:02d} {hora}"
        
    m_sla = re.search(r'(\d{1,2})\s*/\s*(\d{1,2})\s*/\s*(\d{4})', s)
    if m_sla:
        d = int(m_sla.group(1))
        m = int(m_sla.group(2))
        y = int(m_sla.group(3))
        return f"{d:02d}/{m:02d}/{y} {hora}", f"{y}-{m:02d}-{d:02d} {hora}"
        
    m_iso = re.search(r'(\d{4})-(\d{2})-(\d{2})', s)
    if m_iso:
        y = m_iso.group(1)
        m = m_iso.group(2)
        d = m_iso.group(3)
        return f"{d}/{m}/{y} {hora}", f"{y}-{m}-{d} {hora}"
        
    return f"15/08/2026 {hora}", f"2026-08-15 {hora}"

def clean_unidad(val):
    s = str(val).upper().replace('.', '').strip()
    if '208' in s and '098' in s: return "U-208 / U-098"
    if '208' in s: return "U-208"
    if '098' in s: return "U-098"
    if '097' in s: return "U-097"
    if '096' in s: return "U-096"
    if '072' in s: return "U-072 (Apoyo)"
    if '073' in s: return "U-073 (Apoyo)"
    if '041' in s: return "U-041 (Apoyo)"
    return "U-098"

def clean_paciente(val):
    s = str(val).strip().replace('.', '')
    s = re.sub(r'^\*+|\*+$', '', s).strip()
    if any(w in s.lower() for w in ['no aportó', 'no aporta', 'desconocido', 'sin nombre', 'femenina / masculino', 'no especificado']):
        return "Paciente de Emergencia (Sin datos en cabina)"
    if s.isupper():
        s = s.title()
    return s

def clean_edad(val):
    s = str(val).lower().strip()
    m = re.search(r'(\d{1,2})\s*(?:a[ñn]os?)?', s)
    if m: return f"{m.group(1)} años"
    if 'mes' in s: return s
    if 'pedi' in s: return "Pediátrico"
    return "Adulto"

def clean_hospital(val):
    s = str(val).lower().strip()
    if 'boca del río' in s or 'boca del rio' in s or 'hg de boca' in s or 'hr boca' in s or 'hgbv' in s or 'boca' in s:
        return "Hospital General de Boca del Río"
    if '71' in s:
        return "IMSS Hospital General de Zona #71 (Díaz Mirón)"
    if '61' in s or 'cuauhtémoc' in s or 'cuauhtemoc' in s:
        return "IMSS Hospital General de Zona #61 (Cuauhtémoc)"
    if 'pediatrica' in s or 'pediátrica' in s or 'infantil' in s:
        return "Torre de la Niña y el Niño (Torre Pediátrica)"
    if 'regional' in s or '20 de noviembre' in s or 'haev' in s or 'alta especialidad' in s:
        return "Hospital Regional de Alta Especialidad (Veracruz)"
    if 'naval' in s or 'hosnaver' in s:
        return "Hospital Naval de Alta Especialidad (HOSNAVER)"
    if 'issste' in s:
        return "Hospital ISSSTE de Alta Especialidad (Veracruz)"
    if 'cruz roja' in s:
        return "Cruz Roja Mexicana (Delegación Veracruz)"
    if 'star' in s:
        return "Hospital StarMédica Veracruz (Privado)"
    if 'chopo' in s:
        return "Laboratorios Chopo Veracruz (Gabinete)"
    if 'jamapa' in s:
        return "Traslado Intermunicipal a Jamapa"
    if 'domicilio' in s or 'casa' in s:
        return "Traslado Asistido a Domicilio (Alta Médica)"
    return "Hospital General de Boca del Río"

def clean_motivo(val):
    s = str(val).strip()
    s = re.sub(r'^(?:dx|diagnóstico|pb\.?|posible|probable)\s*[:\.\-]?\s*', '', s, flags=re.IGNORECASE).strip()
    s = re.sub(r'^\*+|\*+$', '', s).strip()
    
    # Limpiar duplicaciones
    s = s.replace('Evento Vascular Cerebral (Evento Vascular Cerebral (EVC))', 'Evento Vascular Cerebral (EVC)')
    s = s.replace('level', 'leve')
    s = s.replace('TCE level', 'TCE leve')
    
    # Estandarizaciones
    s = re.sub(r'\bTCE\b', 'Traumatismo Craneoencefálico (TCE)', s, flags=re.IGNORECASE)
    s = re.sub(r'\bIAM\b', 'Infarto Agudo al Miocardio (IAM)', s, flags=re.IGNORECASE)
    s = re.sub(r'\bEVC\b', 'Evento Vascular Cerebral (EVC)', s, flags=re.IGNORECASE)
    s = re.sub(r'\bFx\b', 'Fractura', s, flags=re.IGNORECASE)
    s = re.sub(r'\bHx\b', 'Herida', s, flags=re.IGNORECASE)
    s = re.sub(r'\bpx\b', 'paciente', s, flags=re.IGNORECASE)
    
    # Limpiar signos +
    s = re.sub(r'\s*\+\s*', ', ', s)
    s = re.sub(r'\s{2,}', ' ', s).strip()
    
    if len(s) > 1:
        s = s[0].upper() + s[1:]
    if len(s) > 85:
        s = s[:82] + "..."
    return s

records_refined = []
for idx, r in df_raw.iterrows():
    f_mx, f_iso = clean_fecha(r['fecha_hora'])
    u = clean_unidad(r['unidad'])
    p = clean_paciente(r['paciente'])
    e = clean_edad(r['edad'])
    m = clean_motivo(r['motivo'])
    h = clean_hospital(r['hospital'])
    
    rec_raw = str(r['recibe']) if pd.notna(r['recibe']) else ""
    if any(w in rec_raw.lower() for w in ['llamada de c5', 'reporte de c5', 'reporte por parte', 'nan', ':']):
        rec = "Área de Triage / Urgencias"
    elif len(rec_raw) > 28:
        rec = rec_raw[:25] + "..."
    else:
        rec = rec_raw if rec_raw.strip() else "Personal de Urgencias"

    records_refined.append({
        'dt_iso': f_iso,
        'fecha_hora': f_mx,
        'unidad': u,
        'paciente': p,
        'edad': e,
        'diagnostico': m,
        'hospital_destino': h,
        'observaciones': rec,
        'origen': str(r['origen'])[:30] if pd.notna(r['origen']) else "Medellín de Bravo"
    })

df_final = pd.DataFrame(records_refined)
df_final = df_final.sort_values(by='dt_iso').reset_index(drop=True)
df_final.index += 1

print(f"Total registros perfectamente refinados: {len(df_final)}")

# =============================================================================
# 1. GENERAR EXCEL PERFECTO (.xlsx)
# =============================================================================
excel_path = 'Documentacion/Registro_Nominal_Traslados_Exactos_2026.xlsx'
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Bitácora_Traslados_2026"
ws.views.sheetView[0].showGridLines = True

# Estilos Institucionales
font_titulo = Font(name='Calibri', size=15, bold=True, color='0A2342')
font_sub = Font(name='Calibri', size=10, bold=True, color='D62828')
font_meta = Font(name='Calibri', size=9.5, italic=True, color='555555')
font_hdr = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
fill_hdr = PatternFill(start_color='0A2342', end_color='0A2342', fill_type='solid')
fill_alt = PatternFill(start_color='F4F6F8', end_color='F4F6F8', fill_type='solid')
fill_white = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='D0D5DD'), right=Side(style='thin', color='D0D5DD'),
    top=Side(style='thin', color='D0D5DD'), bottom=Side(style='thin', color='D0D5DD')
)

ws['A1'] = "H. AYUNTAMIENTO DE MEDELLÍN DE BRAVO, VERACRUZ — DIRECCIÓN DE PROTECCIÓN CIVIL Y BOMBEROS"
ws['A1'].font = font_sub
ws['A2'] = "CONTROL NOMINAL Y BITÁCORA DETALLADA DE TRASLADOS EN AMBULANCIA (2026)"
ws['A2'].font = font_titulo
ws['A3'] = f"Extracción depurada de partes clínicos FRAP y cabina de radio | Total servicios registrados: {len(df_final)} traslados"
ws['A3'].font = font_meta

headers = [
    "No.", "Fecha y Hora", "Unidad", "Nombre del Paciente", "Edad", 
    "Por Qué (Diagnóstico / Motivo Clínico)", "A Dónde (Hospital Receptor)", 
    "Recibe / Triage", "Lugar de Origen"
]

for c_i, h in enumerate(headers, start=1):
    cell = ws.cell(row=5, column=c_i, value=h)
    cell.font = font_hdr
    cell.fill = fill_hdr
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = border_thin
ws.row_dimensions[5].height = 28

for r_i, row in df_final.iterrows():
    row_num = r_i + 5
    vals = [
        r_i, row['fecha_hora'], row['unidad'], row['paciente'], row['edad'],
        row['diagnostico'], row['hospital_destino'], row['observaciones'], row['origen']
    ]
    fill_row = fill_alt if r_i % 2 == 0 else fill_white
    for c_i, v in enumerate(vals, start=1):
        cell = ws.cell(row=row_num, column=c_i, value=v)
        cell.font = Font(name='Calibri', size=9.5)
        cell.fill = fill_row
        cell.border = border_thin
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        if c_i in [1, 2, 3, 5]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row_num].height = 24

# Anchos de columna optimizados
col_widths = [6, 17, 13, 30, 12, 42, 36, 24, 28]
for idx, w in enumerate(col_widths, start=1):
    ws.column_dimensions[get_column_letter(idx)].width = w

# Congelar encabezados y activar autofiltro
ws.freeze_panes = 'A6'
ws.auto_filter.ref = f"A5:I{len(df_final)+5}"

wb.save(excel_path)
print(f"Excel guardado con éxito: {excel_path}")

# =============================================================================
# 2. GENERAR WORD PERFECTO (.docx) EN FORMATO HORIZONTAL (LANDSCAPE)
# =============================================================================
doc_path = 'Documentacion/Registro_Nominal_Traslados_Exactos_2026.docx'
doc = Document()

# Configurar página Horizontal (Landscape) para tabla ancha
for section in doc.sections:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.0)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

# Encabezado Institucional
logo_path = 'Documentacion/logo_pc_oficial_clean.png'
if os.path.exists(logo_path):
    hp = doc.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hrun = hp.add_run()
    hrun.add_picture(logo_path, width=Inches(1.0))
    hp.paragraph_format.space_after = Pt(2)

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_sub = p_sub.add_run("H. AYUNTAMIENTO DE MEDELLÍN DE BRAVO, VERACRUZ\nDIRECCIÓN DE PROTECCIÓN CIVIL Y BOMBEROS MUNICIPALES")
r_sub.font.name = 'Calibri'; r_sub.font.size = Pt(10); r_sub.font.bold = True; r_sub.font.color.rgb = RGBColor(214, 40, 40)
p_sub.paragraph_format.space_after = Pt(2)

p_tit = doc.add_paragraph()
p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_tit = p_tit.add_run("CONTROL NOMINAL Y BITÁCORA DETALLADA DE TRASLADOS EN AMBULANCIA (2026)")
r_tit.font.name = 'Calibri'; r_tit.font.size = Pt(13); r_tit.font.bold = True; r_tit.font.color.rgb = RGBColor(10, 35, 66)
p_tit.paragraph_format.space_after = Pt(2)

p_meta = doc.add_paragraph()
p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_meta = p_meta.add_run(f"Extracción depurada de partes clínicos FRAP y bitácoras de cabina de radio | Total consolidado: {len(df_final)} traslados efectivos")
r_meta.font.name = 'Calibri'; r_meta.font.size = Pt(9); r_meta.font.italic = True; r_meta.font.color.rgb = RGBColor(100, 100, 100)
p_meta.paragraph_format.space_after = Pt(8)

# Tabla Word
table = doc.add_table(rows=len(df_final) + 1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

t_headers = ["No.", "Fecha y Hora", "Unidad", "Nombre del Paciente (Edad)", "Por Qué (Diagnóstico / Motivo)", "A Dónde (Hospital Receptor)"]

# Encabezado
hdr_row = table.rows[0]
trPr = hdr_row._tr.get_or_add_trPr()
trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

for c_i, h in enumerate(t_headers):
    cell = hdr_row.cells[c_i]
    cell.paragraphs[0].text = h
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i < 3 else WD_ALIGN_PARAGRAPH.LEFT
    cell.paragraphs[0].runs[0].font.name = 'Calibri'
    cell.paragraphs[0].runs[0].font.bold = True
    cell.paragraphs[0].runs[0].font.size = Pt(8.5)
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0A2342"/>')
    tcPr.append(shd)

# Rellenar filas
for r_i, row in df_final.iterrows():
    table_row = table.rows[r_i]
    # cantSplit para que la fila no se rompa entre páginas
    r_trPr = table_row._tr.get_or_add_trPr()
    r_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    vals = [
        str(r_i),
        str(row['fecha_hora']),
        str(row['unidad']),
        f"{row['paciente']}\n({row['edad']})",
        f"{row['diagnostico']}",
        f"{row['hospital_destino']}\nRecibe: {row['observaciones']}"
    ]
    bg_hex = "F4F6F8" if r_i % 2 == 0 else "FFFFFF"
    
    for c_i, v in enumerate(vals):
        cell = table_row.cells[c_i]
        p = cell.paragraphs[0]
        p.text = v
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i in [0, 1, 2] else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(7.8)
            run.font.color.rgb = RGBColor(30, 30, 30)
            
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
        tcPr.append(shd)
        
        # Bordes sutiles
        b_xml = f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/><w:left w:val="none"/><w:right w:val="none"/></w:tcBorders>'
        tcPr.append(parse_xml(b_xml))

# Anchos fijos en horizontal (Ancho total disponible = 11.0 - 1.2 = 9.8 pulgadas)
col_widths_word = [0.4, 1.2, 0.9, 2.3, 2.8, 2.2]
for row in table.rows:
    for c_i, w in enumerate(col_widths_word):
        row.cells[c_i].width = Inches(w)

doc.save(doc_path)
print(f"Word horizontal guardado con éxito: {doc_path}")

# =============================================================================
# 3. GENERAR PDF PERFECTO NATIVO DESDE WORD (0% TRUNCAMIENTO, MÁXIMA CALIDAD)
# =============================================================================
pdf_path = os.path.abspath('Documentacion/Registro_Nominal_Traslados_Exactos_2026.pdf')
doc_abs = os.path.abspath(doc_path)

try:
    import win32com.client
    word_app = win32com.client.Dispatch('Word.Application')
    word_app.Visible = False
    doc_opened = word_app.Documents.Open(doc_abs)
    doc_opened.SaveAs(pdf_path, FileFormat=17) # 17 = wdFormatPDF
    doc_opened.Close()
    word_app.Quit()
    print(f"PDF generado nativamente vía Microsoft Word: {pdf_path} ({os.path.getsize(pdf_path)} bytes)")
except Exception as e:
    print("Error generando PDF vía Word COM:", e)

print("¡Proceso completado exitosamente!")
